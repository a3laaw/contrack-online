# -*- coding: utf-8 -*-
"""صفحة «ربط مراحل ودفعات نوع العقد» — نقطة الدخول الرئيسية للربط.

تُفتح من: الإعدادات ← أنواع العقود ← «ربط مراحل ودفعات هذا النوع».
داخلها: شجرة مراحل النوع (للعرض — إدارتها من تبويب مراحل المتابعة)
+ قالب دفعاته: كل دفعة برقمها وقيمتها (نسبة أو مبلغ ثابت) ومراحلها
(واحدة أو أكثر أو بلا) مع تحديد **المرحلة المُحفِّزة** التي إنهاؤها
فعليًا يستحق الدفعة — وكل حفظ فوري برسالة تأكيد.
"""
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (QCheckBox, QComboBox, QDialog, QDialogButtonBox,
                               QDoubleSpinBox, QFormLayout, QHBoxLayout,
                               QHeaderView, QLabel, QLineEdit,
                               QPushButton, QSpinBox, QTableWidget,
                               QTableWidgetItem, QVBoxLayout, QWidget)

from services import workflow_service as wf
from services.customer_service import ValidationError
from services.db_facade_service import DatabaseError
from ui import style
from ui.widgets import (flow_host, sync_header_alignment,
                        attach_row_menu, apply_column_weights, make_searchable)


_AR_ORDINAL_UNITS = {
    1: "الأولى", 2: "الثانية", 3: "الثالثة", 4: "الرابعة", 5: "الخامسة",
    6: "السادسة", 7: "السابعة", 8: "الثامنة", 9: "التاسعة",
}
_AR_ORDINAL_TEENS = {
    10: "العاشرة", 11: "الحادية عشرة", 12: "الثانية عشرة", 13: "الثالثة عشرة",
    14: "الرابعة عشرة", 15: "الخامسة عشرة", 16: "السادسة عشرة",
    17: "السابعة عشرة", 18: "الثامنة عشرة", 19: "التاسعة عشرة",
}
_AR_ORDINAL_TENS = {
    20: "العشرون", 30: "الثلاثون", 40: "الأربعون", 50: "الخمسون",
    60: "الستون", 70: "السبعون", 80: "الثمانون", 90: "التسعون",
}
_AR_ORDINAL_COMPOUND_UNITS = dict(_AR_ORDINAL_UNITS)
_AR_ORDINAL_COMPOUND_UNITS[1] = "الحادية"


def payment_title(pay_no: int) -> str:
    """اسم الدفعة المعروض تلقائيًا من رقمها."""
    n = int(pay_no or 1)
    if n in _AR_ORDINAL_UNITS:
        word = _AR_ORDINAL_UNITS[n]
    elif n in _AR_ORDINAL_TEENS:
        word = _AR_ORDINAL_TEENS[n]
    elif n in _AR_ORDINAL_TENS:
        word = _AR_ORDINAL_TENS[n]
    elif 20 < n < 100:
        word = f"{_AR_ORDINAL_COMPOUND_UNITS.get(n % 10, str(n % 10))} و{_AR_ORDINAL_TENS.get((n // 10) * 10, str((n // 10) * 10))}"
    else:
        word = str(n)
    return f"الدفعة {word}"


class PaymentEditDialog(QDialog):
    """نافذة إضافة/تعديل دفعة في قالب النوع."""

    def __init__(self, parent, type_name: str, payment: dict | None = None):
        super().__init__(parent)
        self.type_name = type_name
        self.payment = payment
        self.stage_ids: list[int] = ([s["id"] for s in payment["stages"]]
                                     if payment else [])
        self.setWindowTitle("تعديل الدفعة" if payment else "دفعة جديدة")
        self.setLayoutDirection(Qt.RightToLeft)
        self.setMinimumWidth(470)
        # تُعبَّأ بعد بناء الصفوف؛ تبقى None أثناء إشارات التهيئة المبكرة.
        self.stages_host = None
        self._stages_label = None
        self._trigger_help_label = None
        self._trigger_label = None
        self._days_label = None
        self._pct_label = None
        self._fixed_label = None
        form = QFormLayout(self)
        form.setSpacing(10)

        self.no = QSpinBox(self)
        self.no.setRange(1, 99)
        self.title = QLineEdit(self)
        self.title.setToolTip("اكتب اسم الدفعة كما سيظهر داخل العقد؛ يمكن مثلًا: كامل قيمة العقد")
        self.full_payment = QCheckBox("دفعة كاملة / بدون ترتيب", self)
        self.full_payment.setToolTip("عند التفعيل تُسجل الدفعة باسم ثابت «دفعة كاملة» ولا ترتبط بترتيب مراحل أو رقم يدوي.")
        self.full_payment.toggled.connect(self._toggle_full_payment)
        self.value_mode = QComboBox(self)  # داخلي فقط؛ العرض للمستخدم بزرّي اختيار واضحين
        make_searchable(self.value_mode)
        self.value_mode.addItem("نسبة من قيمة العقد", "pct")
        self.value_mode.addItem("مبلغ يدوي", "fixed")
        self.value_mode.currentIndexChanged.connect(self._value_mode_changed)
        self.value_mode.setVisible(False)  # داخلي فقط — لا يظهر أبدًا

        self.use_pct = QCheckBox("نسبة من قيمة العقد", self)
        self.use_fixed = QCheckBox("مبلغ يدوي", self)
        self.use_pct.setToolTip("تُحسب الدفعة كنسبة من قيمة عقد العميل.")
        self.use_fixed.setToolTip("اكتب مبلغ الدفعة يدويًا في القالب؛ لا يلزم أن يطابق أي مبلغ موجود في الإعدادات.")
        self.use_pct.toggled.connect(lambda on: self._amount_choice("pct", on))
        self.use_fixed.toggled.connect(lambda on: self._amount_choice("fixed", on))
        self.amount_choice_host = QWidget(self)
        self.amount_choice_host.setObjectName("transparentHost")
        ac = QHBoxLayout(self.amount_choice_host)
        ac.setContentsMargins(0, 0, 0, 0)
        ac.addWidget(self.use_pct)
        ac.addWidget(self.use_fixed)
        ac.addStretch()

        self.pct = QDoubleSpinBox(self)
        self.pct.setRange(0, 100)
        self.pct.setDecimals(1)
        self.pct.setToolTip("نسبة من قيمة العقد؛ تتغير قيمة الدفعة تلقائيًا حسب مبلغ كل عميل")
        self.fixed = QDoubleSpinBox(self)
        self.fixed.setDecimals(3)
        self.fixed.setRange(0, 999_999_999)
        self.fixed.setGroupSeparatorShown(True)
        self.fixed.setToolTip("مبلغ يدوي داخل قالب الدفعات؛ يمكنك تغييره لاحقًا داخل العقد.")

        self.link_stage = QCheckBox("ربط الدفعة بمرحلة", self)
        self.link_manual = QCheckBox("استحقاق يدوي/بالأيام", self)
        self.link_stage.setToolTip("عند إنهاء المرحلة المحفزة تصبح الدفعة مستحقة فعليًا.")
        self.link_manual.setToolTip("دفعة لا تنتظر مرحلة؛ تستحق حسب عدد الأيام من تاريخ التوقيع.")
        self.link_stage.toggled.connect(lambda on: self._link_choice("stage", on))
        self.link_manual.toggled.connect(lambda on: self._link_choice("manual", on))
        self.link_choice_host = QWidget(self)
        self.link_choice_host.setObjectName("transparentHost")
        lc = QHBoxLayout(self.link_choice_host)
        lc.setContentsMargins(0, 0, 0, 0)
        lc.addWidget(self.link_stage)
        lc.addWidget(self.link_manual)
        lc.addStretch()

        self.stages_lbl = QLabel("بلا مراحل (دفعة بالأيام)", self)
        self.stages_lbl.setProperty("hint", True)
        self.pick_btn = QPushButton("اختيار المراحل...", self)
        self.pick_btn.setProperty("flat", True)
        self.pick_btn.setToolTip("دفعة تجمع مرحلة أو أكثر — والمرحلة الواحدة تتبع "
                                 "دفعة واحدة فقط")
        self.pick_btn.clicked.connect(self._pick_stages)
        self.trigger = QComboBox(self)
        make_searchable(self.trigger, "")
        self.trigger.setToolTip("عند إنهاء هذه المرحلة تستحق الدفعة.")
        self.trigger_help = QLabel("المرحلة المُحفِّزة: عند إنهائها تستحق الدفعة.", self)
        self.trigger_help.setProperty("hint", True)
        self.trigger_help.setWordWrap(True)
        self.days = QSpinBox(self)
        self.days.setRange(0, 9999)
        self.days.setToolTip("عدد أيام تقديري أو يدوي من تاريخ العقد.")

        if payment:
            self.no.setValue(int(payment["pay_no"]))
            self.title.setText(payment.get("title") or payment_title(self.no.value()))
            self.pct.setValue(float(payment["pct"] or 0))
            self.fixed.setValue(float(payment.get("amount_fixed") or 0))
            if self.fixed.value() > 0:
                self.use_fixed.setChecked(True)
            elif self.pct.value() > 0:
                self.use_pct.setChecked(True)
            self.days.setValue(int(payment["days_offset"] or 0))
            names = " + ".join(s["stage_name"] for s in payment["stages"])
            if names:
                self.stages_lbl.setText(names)
            if (payment.get("title") or "").strip() == "دفعة كاملة" and not self.stage_ids:
                self.full_payment.setChecked(True)
        else:
            nxt = 1
            used = {int(p["pay_no"]) for p in
                    wf.type_payments(type_name)}
            while nxt in used:
                nxt += 1
            self.no.setValue(nxt)
            self.title.setText(payment_title(nxt))

        if not self.use_pct.isChecked() and not self.use_fixed.isChecked():
            self.use_fixed.setChecked(True)
        if self.full_payment.isChecked():
            self.link_manual.setChecked(True)
        else:
            self.link_stage.setChecked(True)
        self._value_mode_changed()
        form.addRow("دفعة كاملة:", self.full_payment)
        form.addRow("رقم الدفعة:", self.no)
        form.addRow("اسم الدفعة:", self.title)
        form.addRow("طريقة قيمة الدفعة:", self.amount_choice_host)
        form.addRow("النسبة %:", self.pct)
        form.addRow("المبلغ اليدوي:", self.fixed)
        form.addRow("طريقة الاستحقاق:", self.link_choice_host)
        st_host, st_row = flow_host(spacing=8)
        self.stages_host = st_host
        st_row.addWidget(self.stages_lbl, 1)
        st_row.addWidget(self.pick_btn)
        form.addRow("المراحل:", st_host)
        form.addRow("شرح المُحفِّزة:", self.trigger_help)
        form.addRow("المرحلة المُحفِّزة:", self.trigger)
        form.addRow("الاستحقاق بعد (يوم):", self.days)
        self._stages_label = form.labelForField(st_host)
        self._trigger_help_label = form.labelForField(self.trigger_help)
        self._trigger_label = form.labelForField(self.trigger)
        self._days_label = form.labelForField(self.days)
        self._pct_label = form.labelForField(self.pct)
        self._fixed_label = form.labelForField(self.fixed)
        self._value_mode_changed()  # مزامنة أخيرة بعد التقاط التسميات
        self._reload_trigger()
        self.trigger.currentIndexChanged.connect(self._trigger_changed)
        self._link_choice("manual" if self.full_payment.isChecked() else "stage", True)
        self._toggle_full_payment(self.full_payment.isChecked())

        bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel, self)
        bb.button(QDialogButtonBox.Save).setText("حفظ")
        bb.button(QDialogButtonBox.Cancel).setText("إلغاء")
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        form.addRow(bb)

    def _refit(self):
        try:
            if self.isVisible():
                from ui import style as _style
                _style.refit_dialog(self)
        except Exception:
            pass

    def _toggle_full_payment(self, on: bool):
        self.no.setEnabled(not on)
        self.title.setEnabled(not on)
        if on:
            self.title.setText("دفعة كاملة")
            self.use_pct.setChecked(True)
            self.pct.setValue(100.0)
            self.link_manual.setChecked(True)
            self.stage_ids = []
            self.stages_lbl.setText("بدون ترتيب مراحل")
            self._reload_trigger(None)
            self.days.setValue(0)
            self.link_choice_host.setEnabled(False)
            self.pick_btn.setEnabled(False)
            self.trigger.setEnabled(False)
            self.days.setEnabled(False)
            self._show_row(self.days, self._days_label, False)
        else:
            self.link_choice_host.setEnabled(True)
            self.pick_btn.setEnabled(True)
            self.trigger.setEnabled(True)
            self.days.setEnabled(True)
            self._show_row(self.days, self._days_label,
                            self.link_manual.isChecked())
            if self.title.text().strip() == "دفعة كاملة":
                self.title.setText(payment_title(self.no.value()))

    def _sync_title(self):
        if not self.title.text().strip():
            self.title.setText(payment_title(self.no.value()))

    def _amount_choice(self, mode: str, on: bool):
        # V4.5.14: لا يُسمح بتفريغ الخيارين — النقر على الوضع المُعلَّم يُبقيه
        # مُعلَّمًا (سلوك رزّي راديو)، فلا تُفقد قيمة مكتوبة بصمت بعد الآن.
        mine = self.use_fixed if mode == "fixed" else self.use_pct
        other = self.use_pct if mode == "fixed" else self.use_fixed
        if not on:
            if not other.isChecked():
                mine.blockSignals(True)
                mine.setChecked(True)
                mine.blockSignals(False)
            self._value_mode_changed()
            return
        if mode == "fixed":
            self.use_pct.blockSignals(True)
            self.use_pct.setChecked(False)
            self.use_pct.blockSignals(False)
            self.value_mode.setCurrentIndex(self.value_mode.findData("fixed"))
        else:
            self.use_fixed.blockSignals(True)
            self.use_fixed.setChecked(False)
            self.use_fixed.blockSignals(False)
            self.value_mode.setCurrentIndex(self.value_mode.findData("pct"))
        self._value_mode_changed()
        self._refit()

    def _value_mode_changed(self):
        fixed = self.use_fixed.isChecked()
        pct = self.use_pct.isChecked()
        self._show_row(self.pct, self._pct_label, pct)
        self._show_row(self.fixed, self._fixed_label, fixed)
        self.pct.setEnabled(pct)
        self.fixed.setEnabled(fixed)
        if fixed:
            self.pct.setValue(0)
        elif pct:
            self.fixed.setValue(0)

    def _link_choice(self, mode: str, on: bool):
        if not on:
            return
        if mode == "stage":
            self.link_manual.blockSignals(True)
            self.link_manual.setChecked(False)
            self.link_manual.blockSignals(False)
            show = True
        else:
            self.link_stage.blockSignals(True)
            self.link_stage.setChecked(False)
            self.link_stage.blockSignals(False)
            show = False
            self.stage_ids = []
            self.stages_lbl.setText("بلا مراحل (دفعة بالأيام)")
            self._reload_trigger(None)
        self._show_row(self.stages_host, self._stages_label, show)
        self._show_row(self.trigger_help, self._trigger_help_label, show)
        self._show_row(self.trigger, self._trigger_label, show)
        self._show_row(self.days, self._days_label, not show)
        self._refit()

    @staticmethod
    def _show_row(field, label, show: bool):
        """إظهار/إخفاء صف كامل (حقل + تسمية) — يتحمل None المبكر."""
        if field is not None:
            field.setVisible(show)
        if label is not None:
            label.setVisible(show)

    def pct_value(self) -> float:
        return self.pct.value() if self.use_pct.isChecked() else 0.0

    def fixed_value(self) -> float:
        return self.fixed.value() if self.use_fixed.isChecked() else 0.0

    def selected_stage_ids(self) -> list[int]:
        return list(self.stage_ids) if self.link_stage.isChecked() else []

    def trigger_value(self):
        if self.full_payment.isChecked():
            return 0  # مهم في التعديل: يمسح أي مرحلة محفزة قديمة
        return self.trigger.currentData() if self.link_stage.isChecked() else None

    def is_fixed_mode(self) -> bool:
        return self.use_fixed.isChecked()

    def accept(self):
        """V4.5.14: حفظ دفعة بلا قيمة يحتاج تأكيدًا صريحًا.

        بلا قيمة (لا نسبة ولا مبلغ) تُعتبر الدفعة «مرنة (تُحدد لاحقًا)»
        ويُكتب مبلغها داخل كل عقد على حدة.
        """
        if self.pct_value() <= 0 and self.fixed_value() <= 0:
            from ui import style as _style
            _msg = ("لم تحدد قيمة للدفعة.\n"
                    "ستُعتبر «مرنة (تُحدد لاحقًا)» — ويُكتب مبلغها داخل "
                    "كل عقد على حدة.\n\nهل تريد الحفظ؟")
            try:
                if not _style.confirm(self, _msg):
                    return
            except Exception:
                pass
        super().accept()

    def _reload_trigger(self, keep: int | None = None):
        self.trigger.blockSignals(True)
        self.trigger.clear()
        self.trigger.addItem("— دفعة بالأيام (بلا مُحفِّز) —", None)
        path = wf.stage_path(wf.TX_CONTRACT, self.type_name)
        for s in path:
            self.trigger.addItem(s["stage_name"], s["id"])
            if keep and s["id"] == keep:
                self.trigger.setCurrentIndex(self.trigger.count() - 1)
        self.trigger.blockSignals(False)

    def _pick_stages(self):
        """اختيار مراحل الدفعة (متعدد) — ثم إعادة تعبئة المُحفِّز."""
        from PySide6.QtWidgets import QCheckBox, QScrollArea, QVBoxLayout
        stages = wf.stage_path(wf.TX_CONTRACT, self.type_name)
        if not stages:
            style.info(self, "لا مراحل في قالب هذا النوع — أضفها أولًا "
                             "من: الإعدادات ← مراحل المتابعة.")
            return
        linked = {s["id"] for p in wf.type_payments(self.type_name)
                  if not self.payment or p["id"] != self.payment["id"]
                  for s in p["stages"]}
        dlg = QDialog(self)
        dlg.setWindowTitle("مراحل هذه الدفعة")
        dlg.setLayoutDirection(Qt.RightToLeft)
        dlg.setMinimumWidth(380)
        v = QVBoxLayout(dlg)
        hint = QLabel("اختر مراحل الدفعة. المراحل المعطلة مربوطة بدفعات أخرى.")
        hint.setProperty("hint", True)
        hint.setWordWrap(True)
        v.addWidget(hint)
        area = QScrollArea()
        area.setWidgetResizable(True)
        area.setFrameShape(QScrollArea.NoFrame)
        host = QWidget()
        host.setObjectName("transparentHost")
        hl = QVBoxLayout(host)
        boxes = []
        for s in stages:
            cb = QCheckBox(f"{s['stage_name']}"
                           + (f" · متوقع {s['expected_days']} يوم"
                              if s.get("expected_days") else ""))
            if s["id"] in self.stage_ids:
                cb.setChecked(True)
            elif s["id"] in linked:
                cb.setEnabled(False)
            boxes.append((cb, s))
            hl.addWidget(cb)
        hl.addStretch()
        area.setWidget(host)
        v.addWidget(area, 1)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("تحديد")
        bb.button(QDialogButtonBox.Cancel).setText("إلغاء")
        bb.accepted.connect(dlg.accept)
        bb.rejected.connect(dlg.reject)
        v.addWidget(bb)
        if dlg.exec() != QDialog.Accepted:
            return
        self.stage_ids = [s["id"] for cb, s in boxes if cb.isChecked()]
        names = " + ".join(s["stage_name"] for cb, s in boxes
                           if cb.isChecked())
        self.stages_lbl.setText(names or "بلا مراحل (دفعة بالأيام)")
        # إعادة تعبئة المُحفِّز من المختار فقط
        keep = self.trigger.currentData()
        self.trigger.blockSignals(True)
        self.trigger.clear()
        self.trigger.addItem("— دفعة بالأيام (بلا مُحفِّز) —", None)
        for cb, s in boxes:
            if cb.isChecked():
                self.trigger.addItem(s["stage_name"], s["id"])
                if keep == s["id"]:
                    self.trigger.setCurrentIndex(self.trigger.count() - 1)
        self.trigger.blockSignals(False)
        self._trigger_changed()

    def _trigger_changed(self):
        sid = self.trigger.currentData()
        if sid:
            self.days.setValue(
                wf.stage_cumulative_days(self.type_name, sid))


class TypeWorkflowDialog(QDialog):
    """مراحل ودفعات نوع عقد — نقطة الدخول الرئيسية للربط."""

    def __init__(self, parent, type_name: str):
        super().__init__(parent)
        self.type_name = type_name
        self.setWindowTitle(f"مراحل ودفعات نوع: {type_name}")
        self.setLayoutDirection(Qt.RightToLeft)
        self.resize(1000, 720)
        outer = QVBoxLayout(self)
        outer.setContentsMargins(16, 14, 16, 14)
        outer.setSpacing(10)

        head = QLabel(f"كل عقد جديد من نوع «{type_name}» يبدأ مراحله "
                      "وخطة سداده من هنا فورًا")
        head.setProperty("sectionTitle", True)
        outer.addWidget(head)

        # ===== شجرة مراحل النوع (عرض) =====
        st_card = QWidget()
        st_card.setProperty("card", True)
        stl = QVBoxLayout(st_card)
        stl.setContentsMargins(14, 10, 14, 10)
        st_head = QWidget()
        st_head.setObjectName("transparentHost")
        st_head_row = QHBoxLayout(st_head)
        st_head_row.setContentsMargins(0, 0, 0, 0)
        st_head_row.setSpacing(8)
        st_title = QLabel("شجرة مراحل التنفيذ")
        st_title.setProperty("sectionTitle", True)
        st_head_row.addWidget(st_title)
        st_head_row.addStretch()
        self.download_tree_btn = QPushButton("⬇ تحميل نموذج هذه الشجرة")
        self.download_tree_btn.setProperty("flat", True)
        self.download_tree_btn.setToolTip("نموذج Excel مخصص لنوع العقد الحالي فقط")
        self.download_tree_btn.clicked.connect(self._download_stage_template_for_type)
        st_head_row.addWidget(self.download_tree_btn)
        self.import_tree_btn = QPushButton("رفع شجرة هذا النوع من Excel")
        self.import_tree_btn.setObjectName("btnPrimary")
        self.import_tree_btn.setToolTip("يرفع/يستبدل شجرة ودفعات نوع العقد الحالي فقط، ولا يقبل ملفًا يحتوي أنواعًا أخرى")
        self.import_tree_btn.clicked.connect(self._import_stage_tree_for_this_type)
        st_head_row.addWidget(self.import_tree_btn)
        stl.addWidget(st_head)
        self.stages_table = QTableWidget(0, 3)
        self.stages_table.setHorizontalHeaderLabels(
            ["المرحلة", "الأيام المتوقعة", ""])
        self.stages_table.verticalHeader().setVisible(False)
        self.stages_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.stages_table.setSelectionMode(QTableWidget.NoSelection)
        self.stages_table.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.Stretch)
        self.stages_table.setMaximumHeight(210)
        stl.addWidget(self.stages_table)
        st_hint = QLabel("إدارة المراحل (إضافة/تعديل/ترتيب) من: الإعدادات "
                         "← مراحل المتابعة — هذه الشاشة لربط الدفعات.")
        st_hint.setProperty("hint", True)
        stl.addWidget(st_hint)
        outer.addWidget(st_card)

        # ===== قالب الدفعات =====
        pay_card = QWidget()
        pay_card.setProperty("card", True)
        pl = QVBoxLayout(pay_card)
        pl.setContentsMargins(14, 10, 14, 10)
        pay_title = QLabel("قالب الدفعات")
        pay_title.setProperty("sectionTitle", True)
        pl.addWidget(pay_title)
        pay_hint = QLabel("حدد قيمة الدفعة وطريقة استحقاقها.")
        pay_hint.setProperty("hint", True)
        pay_hint.setWordWrap(True)
        pl.addWidget(pay_hint)

        self.pay_table = QTableWidget(0, 7)
        self.pay_table.setHorizontalHeaderLabels(
            ["رقم", "الدفعة", "القيمة", "النوع", "المراحل",
             "المرحلة المُحفِّزة", "بعد (يوم)"])
        self.pay_table.verticalHeader().setVisible(False)
        self.pay_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.pay_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.pay_table.horizontalHeader().setSectionResizeMode(
            1, QHeaderView.Stretch)
        apply_column_weights(self.pay_table, [6, 30, 14, 9, 23, 12, 6],
                             min_widths=[42, 130, 86, 62, 150, 110, 58])
        self.pay_table.setMinimumHeight(200)
        attach_row_menu(self.pay_table, lambda: [
            ("تعديل الدفعة", self._edit_payment),
            None,
            ("حذف الدفعة", self._del_payment, True),
        ])
        self.pay_table.doubleClicked.connect(lambda *_: self._edit_payment())
        pl.addWidget(self.pay_table, 1)

        add_host, add_row = flow_host(spacing=8)
        add_btn = QPushButton("＋ إضافة دفعة")
        add_btn.setObjectName("btnPrimary")
        add_btn.setToolTip("إضافة دفعة جديدة لقالب هذا النوع")
        add_btn.clicked.connect(self._add_payment)
        add_row.addWidget(add_btn)
        edit_btn = QPushButton("✎ تعديل المحدد")
        edit_btn.setProperty("flat", True)
        edit_btn.setToolTip("تعديل الدفعة المحددة في الجدول")
        edit_btn.clicked.connect(self._edit_payment)
        add_row.addWidget(edit_btn)
        self.total_lbl = QLabel("")
        self.total_lbl.setProperty("hint", True)
        self.total_lbl.setWordWrap(True)
        add_row.addWidget(self.total_lbl)
        pl.addWidget(add_host)
        outer.addWidget(pay_card, 1)

        self._reload()
        style.guard_table_sections(self)


    # ==================== رفع شجرة النوع الحالي ====================
    def _download_stage_template_for_type(self):
        """تحميل نموذج Excel مخصص للشجرة الحالية فقط."""
        from PySide6.QtWidgets import QFileDialog
        import os
        safe = "".join(ch for ch in self.type_name if ch not in r'\/:*?"<>|').strip() or "نوع_عقد"
        default = os.path.join(os.path.expanduser("~"),
                               f"نموذج_شجرة_{safe}.xlsx")
        path, _ = QFileDialog.getSaveFileName(
            self, "حفظ نموذج شجرة هذا النوع", default, "Excel Files (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            wf.generate_stage_tree_template(path)
            # نضيف ملاحظة فقط ولا نملأ الصفوف الفارغة؛ لأن تعبئة نوع العقد
            # وحده دون مرحلة رئيسية ستجعل الصف يُقرأ كخطأ عند الاستيراد.
            try:
                from openpyxl import load_workbook
                wb = load_workbook(path)
                if "اقرأني أولاً" in wb.sheetnames:
                    ws = wb["اقرأني أولاً"]
                    ws.cell(11, 1, f"هذا النموذج مخصص لنوع العقد: {self.type_name}. اكتب هذا الاسم في عمود نوع العقد لكل صف تستخدمه، ولا تضع أنواع عقود أخرى في هذا الملف.")
                wb.save(path)
            except Exception:
                pass
        except Exception as exc:
            style.error(self, f"تعذر إنشاء النموذج.\n{exc}")
            return
        style.success(self, f"تم إنشاء نموذج شجرة النوع الحالي:\n{path}")

    def _import_stage_tree_for_this_type(self):
        """رفع شجرة Excel لنوع العقد الحالي فقط بدل المسار العام."""
        from PySide6.QtWidgets import QFileDialog
        path, _ = QFileDialog.getOpenFileName(
            self, "رفع شجرة هذا النوع من Excel", "", "Excel Files (*.xlsx)")
        if not path:
            return
        try:
            info = wf.read_stage_tree_file(path)
            types = list(info.get("types") or [])
            if types != [self.type_name]:
                raise ValidationError(
                    "هذا الرفع مخصص لنوع العقد الحالي فقط.\n"
                    f"النوع الحالي: {self.type_name}\n"
                    f"الأنواع الموجودة في الملف: {', '.join(types) if types else 'لا يوجد'}\n\n"
                    "حمّل نموذج هذه الشجرة من نفس النافذة أو اجعل عمود نوع العقد مطابقًا للنوع الحالي فقط."
                )
            msg = (
                f"سيتم استبدال شجرة ودفعات نوع العقد «{self.type_name}» فقط.\n\n"
                f"المراحل: {info['row_count']} صف · رئيسية {info['main_count']} · "
                f"فرعية {info['sub_count']} · دفعات {info.get('payment_count', 0)}\n\n"
                "العقود القديمة لا تتأثر لأنها تملك نسخة Snapshot مستقلة. هل تريد المتابعة؟"
            )
            if not style.confirm(self, msg, "تأكيد رفع شجرة هذا النوع"):
                return
            res = wf.import_stage_tree_file(path, replace=True)
        except Exception as exc:
            style.error(self, f"تعذر رفع شجرة هذا النوع.\n{exc}")
            return
        self._reload()
        style.success(
            self,
            f"تم رفع شجرة نوع العقد «{self.type_name}» بنجاح.\n"
            f"تمت إضافة {res.get('added', 0)} مرحلة و{res.get('payments_added', 0)} دفعة."
        )

    # ==================== العرض ====================
    def _reload(self):
        self._reload_stages()
        self._reload_payments()

    def _reload_stages(self):
        from PySide6.QtGui import QFont, QColor
        tree = wf.stage_tree(wf.TX_CONTRACT, self.type_name)
        rows = []
        for m in tree:
            rows.append((m, 0))
            for sub in m["children"]:
                rows.append((sub, 1))
        self.stages_table.setRowCount(len(rows))
        bold = QFont()
        bold.setBold(True)
        for r, (st, lvl) in enumerate(rows):
            name = st["stage_name"] if lvl == 0 else "      └ " + st["stage_name"]
            a = QTableWidgetItem(name)
            a.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            if lvl == 0:
                a.setFont(bold)
            self.stages_table.setItem(r, 0, a)
            if lvl == 1:
                b = QTableWidgetItem(str(st.get("expected_days") or 0))
                b.setTextAlignment(Qt.AlignCenter)
                c = QTableWidgetItem("●")
                c.setTextAlignment(Qt.AlignCenter)
                c.setForeground(QColor(st["color"]))
                self.stages_table.setItem(r, 1, b)
                self.stages_table.setItem(r, 2, c)
            else:
                dash = QTableWidgetItem("—")
                dash.setTextAlignment(Qt.AlignCenter)
                self.stages_table.setItem(r, 1, dash)
        sync_header_alignment(self.stages_table)

    def _reload_payments(self):
        from PySide6.QtGui import QColor
        items = wf.type_payments(self.type_name)
        self._items = items
        self.pay_table.setRowCount(len(items))
        t = style.theme()
        for r, it in enumerate(items):
            fixed = float(it.get("amount_fixed") or 0)
            pct = float(it.get("pct") or 0)
            # V4.5.14: بلا قيمة = «مرنة (تُحدد لاحقًا)» بدل «0%» المضلِّلة
            if fixed > 0:
                val = f"{fixed:,.3f} ثابت"
            elif pct > 0:
                val = f"{pct:g}%"
            else:
                val = "مرنة (تُحدد لاحقًا)"
            cells = [str(it["pay_no"]), it["title"] or "—", val,
                     it.get("kind") or "أيام",
                     it.get("stage_name") or "— بلا —",
                     it.get("trigger_name") or "بالأيام",
                     str(it.get("days_offset") or 0)]
            for col, txt in enumerate(cells):
                item = QTableWidgetItem(txt)
                item.setTextAlignment(Qt.AlignCenter)
                if col == 1:
                    item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                if col == 3 and it.get("kind") == "إنجاز":
                    item.setForeground(QColor(t["SUCCESS_700"]))
                if col == 5 and it.get("trigger_name"):
                    item.setForeground(QColor(t["BRAND_600"]))
                    item.setToolTip("إنهاء هذه المرحلة فعليًا يستحق الدفعة "
                                    "بتاريخه")
                self.pay_table.setItem(r, col, item)
        sync_header_alignment(self.pay_table)
        # إجمالي النسب (بلا المرنة: لا قيمة لها بعد)
        # V4.5.14: المرنة (بلا نسبة وبلا مبلغ) تُعدّ وتُذكر صراحةً.
        pcts = [p for p in items if float(p.get("amount_fixed") or 0) <= 0
                and float(p["pct"] or 0) > 0]
        total = round(sum(float(p["pct"] or 0) for p in pcts), 2)
        fixed_sum = sum(float(p.get("amount_fixed") or 0) for p in items)
        flex_count = sum(1 for p in items
                         if float(p.get("amount_fixed") or 0) <= 0
                         and float(p["pct"] or 0) <= 0)
        flex_txt = (f" · {flex_count} مرنة (تُحدد لاحقًا)"
                    if flex_count else "")
        if not items:
            txt = ("لا دفعات بعد — عقود هذا النوع تبدأ مراحلها بلا خطة "
                   "سداد تلقائية.")
        elif fixed_sum > 0 and not pcts and not flex_count:
            txt = (f"كل الدفعات بمبالغ ثابتة (إجماليها "
                   f"{fixed_sum:,.3f}).")
        elif abs(total - 100) < 0.01 and not fixed_sum and not flex_count:
            txt = (f"✅ إجمالي النسب {total:g}% — قيمة العقد الجديد "
                   "تتغطى كاملة.")
        else:
            extra = (f" + مبالغ ثابتة {fixed_sum:,.3f}"
                     if fixed_sum > 0 else "")
            txt = (f"⚠ إجمالي النسب {total:g}%{extra}{flex_txt} — "
                   "يُفضَّل 100% حتى تتغطى قيمة العقد كاملة.")
        self.total_lbl.setText(txt)

    # ==================== الإجراءات ====================
    def _add_payment(self):
        dlg = PaymentEditDialog(self, self.type_name)
        if dlg.exec() != QDialog.Accepted:
            return
        try:
            wf.add_type_payment(
                self.type_name, dlg.no.value(), dlg.title.text().strip() or payment_title(dlg.no.value()),
                dlg.pct_value(), dlg.selected_stage_ids(),
                dlg.days.value(), dlg.trigger_value(),
                dlg.fixed_value())
        except (ValidationError, DatabaseError) as exc:
            style.error(self, str(exc))
            return
        self._reload_payments()
        style.toast(self, "تم حفظ الدفعة في قالب النوع")

    def _edit_payment(self):
        items = getattr(self, "_items", [])
        row = self.pay_table.currentRow()
        if row < 0 or row >= len(items):
            style.info(self, "اختر دفعة من الجدول أولاً.")
            return
        it = items[row]
        dlg = PaymentEditDialog(self, self.type_name, it)
        if dlg.exec() != QDialog.Accepted:
            return
        try:
            wf.update_type_payment(
                it["id"], pay_no=dlg.no.value(), title=dlg.title.text().strip() or payment_title(dlg.no.value()),
                pct=dlg.pct_value(),
                stage_id=dlg.selected_stage_ids(),
                days_offset=dlg.days.value(),
                trigger_stage_id=dlg.trigger_value(),
                amount_fixed=dlg.fixed_value())
        except (ValidationError, DatabaseError) as exc:
            style.error(self, str(exc))
            return
        self._reload_payments()
        style.toast(self, "تم حفظ تعديل الدفعة")

    def _del_payment(self):
        items = getattr(self, "_items", [])
        row = self.pay_table.currentRow()
        if row < 0 or row >= len(items):
            style.info(self, "اختر دفعة من الجدول أولاً.")
            return
        it = items[row]
        if not style.confirm(self, f"حذف الدفعة رقم {it['pay_no']} "
                              f"«{it['title']}» من قالب النوع؟"):
            return
        wf.delete_type_payment(it["id"])
        self._reload_payments()
        style.toast(self, "تم حذف الدفعة")
