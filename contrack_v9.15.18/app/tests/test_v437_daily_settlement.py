# -*- coding: utf-8 -*-
"""V4.37: كشف التسوية اليومي للأخرى (تجميع بتاريخ التسوية).

§1 الخدمة (تجميع/صافي/جزئي/استبعادات/تحويلات) · §2 الشاشة (وضع
التسوية + بحث + تفاصيل اليوم) · §3 المصنف (أوراق + ملف فعلي).
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import database.database as db

_tmpdir = tempfile.mkdtemp(prefix="v437_")
db.DATA_DIR = _tmpdir
db.DB_PATH = os.path.join(_tmpdir, "test.sqlite")

from database import database  # noqa: E402

database.init_db()

from services import settlement_service as st  # noqa: E402

PASS = FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print("  PASS " + name)
    else:
        FAIL += 1
        print("  FAIL " + name + (" | " + str(detail) if detail else ""))


G = "KNET"
st.save_profile(None, G, "بنك K", "KNET", 1, 0.150)
st.ensure_ops_table()
with database.get_connection() as conn:
    conn.executemany(
        "INSERT INTO gateway_ops (gateway, op_date, phone, amount, ref,"
        " status, customer, batch, gw_paid, distributed,"
        " classified_manual, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
        [(G, "2026-09-10", "91110001", 500.0, "A1", "other", "أحمد",
          "B1.xlsx", 1, 1, 1, "2026-09-11 10:00:00"),
         (G, "2026-09-10", "91110002", 200.0, "A2", "other", "سارة",
          "B1.xlsx", 1, 1, 1, "2026-09-11 10:00:00"),
         (G, "2026-09-10", "91110003", 1000.0, "M1", "mine", "نورة",
          "B1.xlsx", 1, 1, 1, "2026-09-11 10:00:00"),
         (G, "2026-09-10", "91110004", 300.0, "X1", "other", "مستبعد",
          "B1.xlsx", 0, 1, 1, "2026-09-11 10:00:00"),
         (G, "2026-09-10", "91110005", 400.0, "U1", "unmarked", "مجهول",
          "B1.xlsx", 1, 0, 0, "2026-09-11 10:00:00"),
         (G, "2026-09-12", "91110006", 100.0, "B1", "other", "خالد",
          "B2.xlsx", None, 1, 1, "2026-09-13 10:00:00")])
    _ids = {r["ref"]: r["id"] for r in conn.execute(
        "SELECT id, ref FROM gateway_ops").fetchall()}
t1 = st.add_other_transfer("2026-09-13", 400.0, "TR1", "", "تحويل بنكي",
                           [_ids["A1"]])
t2 = st.add_other_transfer("2026-09-14", 199.85, "TR2", "", "نقدي",
                           [_ids["A2"]])
check("تحويلان برقمين", t1 > 0 and t2 > 0, (t1, t2))

print("=== 1. الخدمة ===")
rep = st.other_daily_settlement("2026-09-01", "2026-09-30")
check("يومان بالترتيب", [d["day"] for d in rep["days"]]
      == ["2026-09-10", "2026-09-12"],
      [d["day"] for d in rep["days"]])
d10 = rep["days"][0]
check("يوم 10: عمليتان مستحقتان فقط",
      d10["n"] == 2
      and [o["ref"] for o in d10["ops"]] == ["A1", "A2"]
      and all("مستبعد" not in o["customer"] for o in d10["ops"]))
check("يوم 10: كامل/عمولة/صافي",
      (d10["gross"], d10["fees"], d10["net"])
      == (700.0, 0.3, 699.7),
      (d10["gross"], d10["fees"], d10["net"]))
check("جزئي A1 وكامل A2",
      d10["ops"][0]["paid"] == 400.0
      and d10["ops"][0]["remaining"] == 99.85
      and d10["ops"][0]["pay_status"] == "partial"
      and d10["ops"][1]["paid"] == 199.85
      and d10["ops"][1]["remaining"] == 0.0
      and d10["ops"][1]["pay_status"] == "paid")
check("يوم 10: مدفوع/متبقي/جزئي",
      (d10["paid"], d10["remaining"], d10["status"])
      == (599.85, 99.85, "partial"),
      (d10["paid"], d10["remaining"], d10["status"]))
d12 = rep["days"][1]
check("يوم 12: مستحق كامل",
      (d12["n"], d12["net"], d12["paid"], d12["remaining"],
       d12["status"]) == (1, 99.85, 0.0, 99.85, "unpaid"),
      (d12["n"], d12["net"], d12["paid"], d12["remaining"]))
check("إجماليات الفترة",
      rep["totals"] == {"n": 3, "gross": 800.0, "fees": 0.45,
                        "net": 799.55, "paid": 599.85,
                        "remaining": 199.7},
      rep["totals"])
check("التحويلات بتواريخها",
      [(t["tdate"], t["amount"]) for t in rep["transfers"]]
      == [("2026-09-13", 400.0), ("2026-09-14", 199.85)]
      and rep["transfers_total"] == 599.85,
      rep["transfers_total"])
_nar = st.other_daily_settlement("2026-09-12", "2026-09-12")
check("الفترة تضيّق الأيام والتحويلات",
      [d["day"] for d in _nar["days"]] == ["2026-09-12"]
      and _nar["transfers"] == []
      and _nar["totals"]["n"] == 1)
_emp = st.other_daily_settlement("2026-01-01", "2026-01-31")
check("فارغ بلا كسر",
      _emp["days"] == [] and _emp["totals"]["remaining"] == 0.0
      and _emp["transfers_total"] == 0.0)

print("=== 2. الشاشة ===")
from PySide6.QtCore import QDate  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402
from ui.other_company_page import OtherCompanyPage  # noqa: E402

app = QApplication.instance() or QApplication([])
oc = OtherCompanyPage()
oc.d1.setDate(QDate(2026, 9, 1))
oc.d2.setDate(QDate(2026, 9, 30))
oc.refresh()
oc.mode.setCurrentData("daily")
app.processEvents()
check("الوضع والقائمة",
      [b.text() for b in oc.mode._buttons]
      == ["📌 المستحق", "✅ المسدد", "📊 التقرير", "الكشف المالي",
          "📅 التسوية اليومية"]
      and not hasattr(oc, "back_btn")
      and oc.kind.isHidden())
check("يومان + إجمالي بـ 8 أعمدة",
      oc.table.columnCount() == 8
      and oc.table.rowCount() == 3
      and oc.table.item(0, 0).text() == "2026-09-10"
      and oc.table.item(0, 1).text() == "2"
      and oc.table.item(0, 2).text() == "700.000"
      and oc.table.item(0, 4).text() == "699.700"
      and oc.table.item(0, 5).text() == "599.850"
      and oc.table.item(0, 6).text() == "99.850"
      and oc.table.item(0, 7).text() == "🟡 جزئي"
      and oc.table.item(2, 1).text() == "⚖️ الإجمالي"
      and oc.table.item(2, 4).text() == "799.550"
      and oc.table.item(2, 6).text() == "199.700",
      oc.table.rowCount())
oc.search.setText("أحمد")
app.processEvents()
check("البحث بالاسم + إجمالي الظاهر",
      oc.table.rowCount() == 2
      and oc.table.item(0, 0).text() == "2026-09-10"
      and oc.table.item(1, 1).text() == "⚖️ إجمالي الظاهر"
      and oc.table.item(1, 4).text() == "699.700")
oc.search.setText("")
app.processEvents()
oc._on_table_click(0, 1)
app.processEvents()
check("تفاصيل اليوم بالأسماء",
      not oc.detail_box.isHidden()
      and oc.ops_table.columnCount() == 9
      and oc.ops_table.rowCount() == 2
      and oc.ops_table.item(0, 0).text() == "أحمد"
      and oc.ops_table.item(0, 5).text() == "499.850"
      and oc.ops_table.item(0, 6).text() == "400.000"
      and oc.ops_table.item(0, 7).text() == "99.850"
      and oc.ops_table.item(0, 8).text() == "🟡 جزئي"
      and oc.ops_table.item(1, 0).text() == "سارة"
      and oc.ops_table.item(1, 8).text() == "✅ مسدد"
      and "المتبقي: 99.850" in oc.detail_info.text())
oc.mode.setCurrentData("due")
app.processEvents()
check("التبويب للمستحق", oc.mode.currentData() == "due")

print("=== 3. المصنف ===")
oc.mode.setCurrentData("daily")
app.processEvents()
sheets = oc._daily_workbook_data()
check("4 أوراق مسماة",
      [sh["name"] for sh in sheets]
      == ["التسوية اليومية", "يوم 2026-09-10", "يوم 2026-09-12",
          "التحويلات"],
      [sh["name"] for sh in sheets])
check("ورقة الأيام + إجماليها",
      sheets[0]["rows"][-1][4] == 799.55
      and sheets[0]["rows"][-1][6] == 199.7
      and sheets[0]["rows"][0][0] == "2026-09-10")
_dsh = sheets[1]
check("ورقة اليوم: عمليات + إجمالي",
      len(_dsh["rows"]) == 3
      and _dsh["rows"][0][0] == "أحمد"
      and _dsh["rows"][0][5] == 499.85
      and _dsh["rows"][-1][6] == 599.85
      and len(_dsh["heads"]) == 9)
_tsh = sheets[-1]
check("ورقة التحويلات + إجماليها",
      len(_tsh["rows"]) == 3
      and _tsh["rows"][0][2] == 400.0
      and _tsh["rows"][-1][2] == 599.85)
from reports.tabular_pdf import generate_report_workbook  # noqa: E402

_xp = os.path.join(_tmpdir, "daily.xlsx")
generate_report_workbook("التسوية اليومية", sheets, _xp)
import openpyxl  # noqa: E402

_wb = openpyxl.load_workbook(_xp)
_vals = [c.value for _ws in _wb.worksheets
         for _r in _ws.iter_rows() for c in _r]
check("ملف فعلي يُقرأ",
      _wb.sheetnames[0] == "التسوية اليومية"
      and 799.55 in _vals and "أحمد" in _vals,
      _wb.sheetnames)

print("")
print(f"V4.37 daily-settlement: {PASS} passed, {FAIL} failed")
sys.exit(1 if FAIL else 0)
