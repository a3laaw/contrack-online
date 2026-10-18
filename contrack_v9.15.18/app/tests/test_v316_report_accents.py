# -*- coding: utf-8 -*-
"""V3.16: صفوف الإجماليات مميزة في PDF وExcel والمعاينة (التقرير المالي)."""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import database.database as db

_tmpdir = tempfile.mkdtemp()
db.DATA_DIR = _tmpdir
db.DB_PATH = os.path.join(_tmpdir, "test.sqlite")
db.init_db()

from database.models import Contract, Customer  # noqa: E402
from services import contract_service as ks  # noqa: E402
from services import customer_service as cs  # noqa: E402
from services import expense_account_service as ea  # noqa: E402
from services import lifecycle_service as lc  # noqa: E402

PASS = FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  [PASS] {name}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name} {detail}")


cid = cs.add_customer(Customer(customer_code="C-16", name="عميل 16", phone="1"))
k1 = ks.add_contract(Contract(contract_number="K-16", customer_id=cid,
                              contract_date="2026-01-05", description="عقد 16",
                              amount=10000.0))
mat = ea.add_account("مواد", None, allow_posting=True)
lc.add_voucher("receipt", "2026-02-01", "", "نقدًا", "", 900.0, "دفعة",
               contract_id=k1, party_type="customer", party_id=cid)
lc.add_voucher("payment", "2026-02-02", "مورد", "نقدًا", "", 100.0, "شراء",
               contract_id=k1, expense_account_id=mat)
lc.add_voucher("transfer", "2026-02-03", "", "الصندوق", "البنك", 50.0,
               "تغذية بنك")

print("=== 1. المواصفة تحمل التمييز ===")
from ui.reports_page import _data_cash_flow  # noqa: E402

spec = _data_cash_flow("detail", "", "")
accents = spec.get("accent_rows") or {}
check("10 صفوف مميزة (V4.54: + المسترد)", len(accents) == 10, list(accents))
second_cells = [str(r[1]) for r in spec["rows"] if len(r) > 1]
check("كل المفاتيح لها صف",
      all(k in second_cells for k in accents),
      [k for k in accents if k not in second_cells])

print("=== 2. PDF ملوّن فعلًا (فحص بكسل) ===")
from reports.tabular_pdf import (generate_table_excel,  # noqa: E402
                                 generate_table_pdf)

pdf_path = os.path.join(_tmpdir, "cashflow.pdf")
generate_table_pdf("التقرير المالي", spec["heads"], spec["widths"],
                   spec["rows"], pdf_path, totals_row=spec["totals"],
                   footer_note=spec["note"], color_cols=spec["color_cols"],
                   row_kinds=spec.get("row_kinds"),
                   right_align_cols=tuple(spec.get("right_align_cols") or ()),
                   number_cols=spec.get("number_cols"),
                   accent_rows=spec.get("accent_rows"))
check("PDF وُلد", os.path.getsize(pdf_path) > 5000,
      os.path.getsize(pdf_path))
import pymupdf  # noqa: E402

doc = pymupdf.open(pdf_path)
img = b"".join(p.samples for p in (doc[i].get_pixmap(dpi=80) for i in range(len(doc))))
n = doc[0].get_pixmap(dpi=80).n
doc.close()


def _count_near(hex_color, tol=24):
    tr = int(hex_color[1:3], 16)
    tg = int(hex_color[3:5], 16)
    tb = int(hex_color[5:7], 16)
    c = 0
    for i in range(0, len(img), n):
        r, g, b = img[i], img[i + 1], img[i + 2]
        if abs(r - tr) <= tol and abs(g - tg) <= tol and abs(b - tb) <= tol:
            c += 1
    return c


for key, hx in (("الإجمالي · الوارد", "#DCFCE7"),
                ("الإجمالي · المنصرف", "#FEE2E2"),
                ("صافي الحركة", "#DBEAFE"),
                ("رصيد البنك", "#FEF3C7"),
                ("رصيد الصندوق", "#EDE9FE")):
    c = _count_near(hx)
    check(f"لون {key} ({hx})", c >= 400, f"pixels={c}")

print("=== 3. Excel مميز ===")
xls_path = os.path.join(_tmpdir, "cashflow.xlsx")
generate_table_excel("التقرير المالي", spec["heads"], spec["rows"],
                     xls_path, totals_row=spec["totals"],
                     number_cols=spec["number_cols"],
                     accent_rows=spec.get("accent_rows"))
from openpyxl import load_workbook  # noqa: E402

wb = load_workbook(xls_path)
ws = wb.active
found = {}
for row in ws.iter_rows(min_row=5):
    label = str(row[1].value or "")
    if label in accents:
        found[label] = row[1]
check("صفوف التمييز موجودة", len(found) == 10, list(found))
ok_fill = True
for label, cell in found.items():
    rgb = (cell.fill.fgColor.rgb or "")[-6:].upper()
    if rgb != accents[label].lstrip("#").upper():
        ok_fill = False
    if not cell.font.bold:
        ok_fill = False
check("تعبئة + عريض في Excel", ok_fill)

print("=== 4. المعاينة مميزة (offscreen) ===")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PySide6.QtWidgets import QApplication  # noqa: E402

app = QApplication.instance() or QApplication([])
from services import auth_service  # noqa: E402
from ui.reports_page import ReportsPage  # noqa: E402

auth_service.set_current_user("admin")
pg = ReportsPage()
pg.show()
app.processEvents()
idx = next(i for i, x in enumerate(pg._visible) if x["key"] == "rep_cashflow")
pg.open_report(idx)
app.processEvents()
from PySide6.QtCore import Qt as _Qt  # noqa: E402
items = [it for it in pg.view.table.findItems("صافي الحركة", _Qt.MatchExactly)
         if it.column() == 1]
hits = []
if items:
    r0 = items[0].row()
    hits = [pg.view.table.item(r0, c) for c in range(pg.view.table.columnCount())]
bg = hits[5].background().color().name().upper() if hits else ""
check("صف الصافي مميز أزرق", bg == "#DBEAFE", bg)
bold = hits[5].font().bold() if hits else False
check("صف الصافي عريض", bold)

print(f"\nالنتيجة: {PASS} ناجح / {FAIL} فاشل")
sys.exit(1 if FAIL else 0)
