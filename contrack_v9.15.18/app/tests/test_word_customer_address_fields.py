# -*- coding: utf-8 -*-
"""ربط نماذج Word يجب أن يرى حقول عنوان العميل المفصلة."""
import os
import sys
import tempfile
from datetime import date

APP_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, APP_ROOT)
os.chdir(APP_ROOT)
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from database import database  # noqa: E402
database.init_db()

from database.models import Customer, Contract  # noqa: E402
from services import customer_service, contract_service, form_service, word_forms_service, word_service  # noqa: E402


def main():
    cid = customer_service.add_customer(Customer(
        customer_code="CUS-TADDR", name="عميل العنوان", phone="55500000",
        address_governorate="حولي", address_area="السالمية",
        address_block="10", address_plot="25"))
    con_id = contract_service.add_contract(Contract(
        contract_number="CON-TADDR", customer_id=cid,
        contract_date=date.today().isoformat(), amount=1000,
        contract_type="تصميم", description="عقد فحص العنوان"))

    tree = form_service.data_source_tree()
    customer_fields = tree["customers"]["fields"]
    assert customer_fields["address_governorate"] == "محافظة العميل"
    assert customer_fields["address_area"] == "منطقة العميل"
    assert customer_fields["address_block"] == "قطعة العميل"
    assert customer_fields["address_plot"] == "قسيمة العميل"

    for placeholder, source_key in {
        "محافظة_العميل": "address_governorate",
        "منطقة_العميل": "address_area",
        "قطعة_العميل": "address_block",
        "قسيمة_العميل": "address_plot",
        "المحافظة": "address_governorate",
        "المنطقة": "address_area",
        "القطعة": "address_block",
        "القسيمة": "address_plot",
    }.items():
        guessed = word_forms_service.guess_field(placeholder)
        assert guessed["source_type"] == "system", placeholder
        assert guessed["source_entity"] == "customers", placeholder
        assert guessed["source_key"] == source_key, placeholder

    ctx = form_service.build_context(contract_id=con_id)
    assert form_service.resolve_field({"source_entity": "customers", "source_key": "address_governorate"}, ctx) == "حولي"
    assert form_service.resolve_field({"source_entity": "customers", "source_key": "address_area"}, ctx) == "السالمية"
    assert form_service.resolve_field({"source_entity": "customers", "source_key": "address_block"}, ctx) == "10"
    assert form_service.resolve_field({"source_entity": "customers", "source_key": "address_plot"}, ctx) == "25"

    # قالب قديم: كان الحقل محفوظًا Manual قبل إضافة العنوان المفصل، يجب ترقيته تلقائيًا.
    with database.get_connection() as conn:
        cur = conn.execute(
            "INSERT INTO form_templates (base_name, version, status, source_type, word_path, pdf_path, pages_count) VALUES ('قالب عنوان قديم',1,'معتمد','word','dummy.docx','',1)")
        tpl_id = int(cur.lastrowid)
    form_service.add_field(tpl_id, "المحافظة", source_type="manual", placeholder="المحافظة")
    upgraded = word_forms_service.template_fields(tpl_id)[0]
    assert upgraded["source_type"] == "system"
    assert upgraded["source_entity"] == "customers"
    assert upgraded["source_key"] == "address_governorate"

    values = word_service.contract_field_values(con_id)
    assert values["محافظة_العميل"] == "حولي"
    assert values["منطقة_العميل"] == "السالمية"
    assert values["قطعة_العميل"] == "10"
    assert values["قسيمة_العميل"] == "25"

    from docx import Document
    doc = Document()
    doc.add_paragraph("{{محافظة_العميل}} {{منطقة_العميل}} {{قطعة_العميل}} {{قسيمة_العميل}}")
    tmpdir = tempfile.mkdtemp()
    tpl = os.path.join(tmpdir, "address_template.docx")
    out = os.path.join(tmpdir, "out.docx")
    doc.save(tpl)
    word_service.fill_template(tpl, con_id, out)
    filled = "\n".join(p.text for p in Document(out).paragraphs)
    assert "حولي السالمية 10 25" in filled
    print("OK: ربط النماذج يدعم حقول عنوان العميل المفصلة")


if __name__ == "__main__":
    main()
