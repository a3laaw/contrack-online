# -*- coding: utf-8 -*-
"""تصفية بوابات الدفع مع البنك — ربط ثلاثي (سيستم × بوابة × بنك).

القواعد المعتمدة (خطة_تصفية_البوابات_والبنك):
  • حساب بنكي واحد يستقبل الكل → كل تسوية تُقسّم (نصيبي/الأخرى).
  • العمولة مبلغ ثابت لكل عملية.
  • التأخير T+1 بأيام العمل (الجمعة والسبت عطلة).
  • مطابقة البنك: مبلغ + تاريخ فقط (± يوم عمل + تسامح قروش).
"""
import os
import shutil
from datetime import date, datetime, timedelta

from database import database as _db
from database.database import get_connection

# الجمعة (4) والسبت (5) عطلة — الكويت
WEEKEND = (4, 5)

BANK_DATE_KEYS = ("تاريخ", "date", "التاريخ", "يوم", "القيد", "posting")
BANK_CREDIT_KEYS = ("دائن", "credit", "له", "إيداع", "deposit", "وارد")
BANK_DEBIT_KEYS = ("مدين", "debit", "عليه", "مسحوب")
BANK_AMOUNT_KEYS = ("مبلغ", "المبلغ", "amount", "value", "القيمة")
BANK_DESCR_KEYS = ("بيان", "وصف", "description", "narrative", "details",
                   "المرجع", "ref")
BANK_TYPE_KEYS = ("نوع", "type", "حركة", "dc", "علامة")


# ---------- أيام العمل ----------
def is_workday(d: date) -> bool:
    return d.weekday() not in WEEKEND


def add_business_days(d: date, n: int) -> date:
    """D+N بأيام العمل (N موجب/سالب)."""
    step = 1 if n >= 0 else -1
    cur = d
    for _ in range(abs(n)):
        cur += timedelta(days=step)
        while not is_workday(cur):
            cur += timedelta(days=step)
    return cur


def effective_date(payment_date: str, method: str) -> str:
    """تاريخ الاستحقاق الفعلي للدفعة.

    • بوابة / تحويل بنكي → اليوم التالي (T+1) مع تخطي الجمعة والسبت.
    • نقدي / شيك / KNET → فوري (نفس اليوم).
    """
    if not payment_date:
        return payment_date
    _m = (method or "").strip()
    if _m not in ("بوابة", "تحويل بنكي", "gateway"):
        return payment_date  # نقدي / شيك / KNET = فوري
    try:
        d = date.fromisoformat(payment_date)
    except (ValueError, TypeError):
        return payment_date
    return add_business_days(d, 1).isoformat()


# ---------- ملفات البوابات ----------
def ensure_profile_table() -> None:
    with get_connection() as conn:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS gateway_profiles ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, "
            "name TEXT NOT NULL UNIQUE, bank_account TEXT DEFAULT '', "
            "methods TEXT DEFAULT 'KNET', lag_days INTEGER DEFAULT 1, "
            "fee_fixed REAL DEFAULT 0, tolerance REAL DEFAULT 0.005, "
            "late_days INTEGER DEFAULT 3, is_active INTEGER DEFAULT 1)")


def list_profiles(active_only: bool = True) -> list[dict]:
    ensure_profile_table()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT * FROM gateway_profiles "
            + ("WHERE is_active=1 " if active_only else "")
            + "ORDER BY name").fetchall()
    out = [dict(r) for r in rows]
    if not out:
        _create_default()
        return list_profiles(active_only)
    return out


def _create_default() -> None:
    with get_connection() as conn:
        conn.execute(
            "INSERT OR IGNORE INTO gateway_profiles "
            "(name, bank_account, methods, lag_days, fee_fixed) "
            "VALUES ('البوابة الافتراضية', '', 'KNET', 1, 0)")


def save_profile(pid: int | None, name: str, bank_account: str,
                 methods: str, lag_days: int, fee_fixed: float,
                 tolerance: float = 0.005, late_days: int = 3) -> int:
    ensure_profile_table()
    name = (name or "").strip() or "بوابة"
    with get_connection() as conn:
        if pid:
            conn.execute(
                "UPDATE gateway_profiles SET name=?, bank_account=?, "
                "methods=?, lag_days=?, fee_fixed=?, tolerance=?, "
                "late_days=? WHERE id=?",
                (name, bank_account.strip(), methods.strip(),
                 int(lag_days), float(fee_fixed), float(tolerance),
                 int(late_days), pid))
            return pid
        cur = conn.execute(
            "INSERT INTO gateway_profiles (name, bank_account, methods, "
            "lag_days, fee_fixed, tolerance, late_days) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (name, bank_account.strip(), methods.strip() or "KNET",
             int(lag_days), float(fee_fixed), float(tolerance),
             int(late_days)))
        return cur.lastrowid


def delete_profile(pid: int) -> None:
    with get_connection() as conn:
        conn.execute("DELETE FROM gateway_profiles WHERE id=?", (pid,))


def get_profile(pid: int) -> dict | None:
    ensure_profile_table()
    with get_connection() as conn:
        r = conn.execute(
            "SELECT * FROM gateway_profiles WHERE id=?", (pid,)).fetchone()
    return dict(r) if r else None


# ---------- عمليات البوابة (مثبتة من شاشة المطابقة) ----------
def ensure_ops_table() -> None:
    with get_connection() as conn:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS gateway_ops ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, gateway TEXT DEFAULT '', "
            "op_date TEXT DEFAULT '', phone TEXT DEFAULT '', "
            "amount REAL DEFAULT 0, ref TEXT DEFAULT '', "
            "status TEXT DEFAULT '', receipt_no TEXT DEFAULT '', "
            "customer TEXT DEFAULT '', batch TEXT DEFAULT '', "
            "created_at TEXT DEFAULT (datetime('now','localtime')))")
        conn.execute(
            "CREATE UNIQUE INDEX IF NOT EXISTS uq_gateway_ops ON "
            "gateway_ops(gateway, op_date, phone, amount, ref)")
        _cols = {r[1] for r in conn.execute("PRAGMA table_info(gateway_ops);")}
        if "paid_transfer_id" not in _cols:
            conn.execute("ALTER TABLE gateway_ops ADD COLUMN "
                         "paid_transfer_id INTEGER DEFAULT 0")
        if "classified_manual" not in _cols:
            conn.execute("ALTER TABLE gateway_ops ADD COLUMN "
                         "classified_manual INTEGER DEFAULT 0")
        if "distributed" not in _cols:
            conn.execute("ALTER TABLE gateway_ops ADD COLUMN "
                         "distributed INTEGER DEFAULT 0")
        if "gw_paid" not in _cols:
            conn.execute("ALTER TABLE gateway_ops ADD COLUMN "
                         "gw_paid INTEGER")
        if "gw_paid_date" not in _cols:
            conn.execute("ALTER TABLE gateway_ops ADD COLUMN "
                         "gw_paid_date TEXT DEFAULT ''")
        if "gw_token" not in _cols:
            conn.execute("ALTER TABLE gateway_ops ADD COLUMN "
                         "gw_token TEXT DEFAULT ''")
        if "distributed_at" not in _cols:
            conn.execute("ALTER TABLE gateway_ops ADD COLUMN "
                         "distributed_at TEXT DEFAULT ''")
        # V8.11: تاريخ الاستحقاق الفعلي (T+1 للبوابة والتحويل البنكي)
        if "effective_date" not in _cols:
            conn.execute("ALTER TABLE gateway_ops ADD COLUMN "
                         "effective_date TEXT DEFAULT ''")
        conn.execute("CREATE TABLE IF NOT EXISTS schema_flags ("
                     "name TEXT PRIMARY KEY)")
        _mig = conn.execute("SELECT 1 FROM schema_flags "
                            "WHERE name='dist_migrated_v416'").fetchone()
        if not _mig:
            # ترحيل لمرة واحدة: المصنف قبل V4.16 يُعتبر موزَّعًا
            # (يحفظ الأرقام) · المعلّق/diff القديم يبقى مرحلة مؤقتة.
            conn.execute("UPDATE gateway_ops SET distributed=1, "
                         "classified_manual=1 WHERE status IN ('mine','other')"
                         " AND COALESCE(distributed,0)=0")
            conn.execute("INSERT OR IGNORE INTO schema_flags (name) "
                         "VALUES ('dist_migrated_v416')")


def save_gateway_ops(gateway: str, matched: list[dict],
                     batch: str) -> int:
    """تثبيت نتيجة التصنيف (مرحلة مؤقتة) — يرجع الجديد فقط.

    الموجود غير المسدد وغير الموزَّع تُحدَّث حالته (التلقائي الجديد) ·
    الموزَّع والمسدد لا يُمسّان · المكرر يُتجاهل · لا أثر حسابي هنا.
    """
    ensure_ops_table()
    n = 0
    with get_connection() as conn:
        for r in matched:
            _dt = r.get("date") or ""
            _ph = (r.get("phone") or "").strip()
            _am = round(float(r.get("amount") or 0), 3)
            _rf = (r.get("ref") or "").strip()
            _ed = effective_date(_dt, "بوابة")  # V8.11: T+1
            if _rf:
                conn.execute(
                    "UPDATE gateway_ops SET op_date=?, effective_date=? "
                    "WHERE gateway=? AND amount=? AND ref=? "
                    "AND COALESCE(paid_transfer_id,0)=0",
                    (_dt, _ed, gateway, _am, _rf))
            conn.execute(
                "UPDATE gateway_ops SET status=?, receipt_no=?, "
                "customer=?, batch=?, effective_date=? WHERE gateway=? AND op_date=? AND "
                "phone=? AND amount=? AND ref=? AND "
                "COALESCE(paid_transfer_id,0)=0 AND "
                "COALESCE(classified_manual,0)=0 AND "
                "COALESCE(distributed,0)=0",
                (r.get("status") or "", r.get("receipt_no") or "",
                 (r.get("customer") or "").strip(), batch, _ed,
                 gateway, _dt, _ph, _am, _rf))
            cur = conn.execute(
                "INSERT OR IGNORE INTO gateway_ops (gateway, op_date, phone, "
                "amount, ref, status, receipt_no, customer, batch, "
                "gw_paid, gw_paid_date, gw_token, effective_date) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (gateway, _dt, _ph, _am, _rf,
                 r.get("status") or "", r.get("receipt_no") or "",
                 (r.get("customer") or "").strip(), batch,
                 r.get("gw_paid"), r.get("gw_paid_date") or "",
                 r.get("gw_token") or "", _ed))
            n += cur.rowcount
    return n


def refresh_gw_paid(gateway: str, rows: list[dict]) -> int:
    """تحديث حالة تحصيل البوابة من صفوف الملف (غير مقيد بالأقفال).

    المطابقة بالمرجع+المبلغ (التاريخ يتغير عند التحصيل) · التاريخ
    يُحدَّث من الملف · تُستثنى المسددة للشركة الأخرى فقط · الصف
    بلا معلومة يُتجاهل. يرجع عدد المحدَّثة.
    """
    ensure_ops_table()
    n = 0
    with get_connection() as conn:
        for r in rows or []:
            _gw = r.get("gw_paid")
            if _gw not in (0, 1, 2):
                continue
            _am = round(float(r.get("amount") or 0), 3)
            _rf = (r.get("ref") or "").strip()
            _dt = r.get("date") or ""
            _tok = (r.get("gw_token") or "").strip()
            if _rf:
                cur = conn.execute(
                    "UPDATE gateway_ops SET gw_paid=?, gw_paid_date=?, "
                    "op_date=?, gw_token=CASE WHEN ?='' THEN gw_token "
                    "ELSE ? END WHERE gateway=? AND amount=? AND ref=? "
                    "AND COALESCE(paid_transfer_id,0)=0",
                    (_gw, r.get("gw_paid_date") or "", _dt, _tok, _tok,
                     gateway, _am, _rf))
            else:
                cur = conn.execute(
                    "UPDATE gateway_ops SET gw_paid=?, gw_paid_date=?, "
                    "op_date=?, gw_token=CASE WHEN ?='' THEN gw_token "
                    "ELSE ? END WHERE gateway=? AND op_date=? AND phone=? "
                    "AND amount=? AND ref=? "
                    "AND COALESCE(paid_transfer_id,0)=0",
                    (_gw, r.get("gw_paid_date") or "", _dt, _tok, _tok,
                     gateway, _dt,
                     (r.get("phone") or "").strip(), _am, _rf))
            n += cur.rowcount
    return n


# ---------- كشف البنك ----------
def ensure_bank_table() -> None:
    with get_connection() as conn:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS bank_lines ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, bank_account TEXT DEFAULT '', "
            "cdate TEXT DEFAULT '', amount REAL DEFAULT 0, "
            "descr TEXT DEFAULT '', batch TEXT DEFAULT '', "
            "created_at TEXT DEFAULT (datetime('now','localtime')))")
        conn.execute(
            "CREATE UNIQUE INDEX IF NOT EXISTS uq_bank_lines ON "
            "bank_lines(bank_account, cdate, amount, descr)")


def _match_key(header: str, keys: tuple) -> bool:
    h = (header or "").strip().lower()
    return any(k in h for k in keys)


def detect_bank_columns(headers: list) -> dict | None:
    found: dict = {}
    for i, h in enumerate(headers):
        hh = str(h or "")
        if "date" not in found and _match_key(hh, BANK_DATE_KEYS):
            found["date"] = i
        elif "credit" not in found and _match_key(hh, BANK_CREDIT_KEYS):
            found["credit"] = i
        elif "debit" not in found and _match_key(hh, BANK_DEBIT_KEYS):
            found["debit"] = i
        elif "amount" not in found and _match_key(hh, BANK_AMOUNT_KEYS):
            found["amount"] = i
        elif "type" not in found and _match_key(hh, BANK_TYPE_KEYS):
            found["type"] = i
        elif "descr" not in found and _match_key(hh, BANK_DESCR_KEYS):
            found["descr"] = i
    if "date" not in found:
        return None
    if "credit" not in found and "amount" not in found:
        return None
    return found


def _bank_amount(row, cols: dict) -> float:
    from services.gateway_service import parse_amount
    if "credit" in cols:
        return parse_amount(row[cols["credit"]])
    if "debit" in cols and "amount" not in cols:
        return 0.0
    val = parse_amount(row[cols["amount"]])
    if "type" in cols:
        t = str(row[cols["type"]] or "").strip().lower()
        if t.startswith(("d", "مدين", "عليه", "debit", "dr")):
            return 0.0
    return val if val > 0 else 0.0


def _extract_bank_lines(matrix: list) -> list[dict]:
    """مصفوفة خام ← إيداعات بنكية [{cdate, amount, descr}]."""
    from services.gateway_service import parse_date
    if not matrix:
        raise ValueError("الملف فارغ.")
    cols = None
    header_idx = 0
    for i, r in enumerate(matrix[:5]):
        cols = detect_bank_columns([c for c in (r or [])])
        if cols:
            header_idx = i
            break
    if not cols:
        raise ValueError("تعذر إيجاد عمودي التاريخ والمبلغ في أول 5 صفوف.")
    out = []
    for r in matrix[header_idx + 1:]:
        if not r:
            continue
        amt = _bank_amount(r, cols)
        if amt <= 0:
            continue
        out.append({
            "cdate": parse_date(r[cols["date"]]),
            "amount": amt,
            "descr": str(r[cols["descr"]] or "").strip() if "descr" in cols else "",
        })
    return [x for x in out if x["cdate"]]


def read_bank_excel(path: str) -> list[dict]:
    """قراءة كشف البنك: الإيداعات فقط [{cdate, amount, descr}]."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return _extract_bank_lines(rows)


def read_bank_csv(path: str) -> list[dict]:
    """قراءة كشف بنك CSV (ترميز عربي/أجنبي)."""
    import csv
    try:
        with open(path, encoding="utf-8-sig", newline="") as f:
            matrix = list(csv.reader(f))
    except UnicodeDecodeError:
        with open(path, encoding="cp1256", newline="") as f:
            matrix = list(csv.reader(f))
    return _extract_bank_lines(matrix)


def read_bank_file(path: str) -> list[dict]:
    """قراءة كشف بنك (xlsx/csv) حسب الامتداد."""
    if str(path or "").lower().endswith(".csv"):
        return read_bank_csv(path)
    return read_bank_excel(path)


def save_bank_lines(bank_account: str, lines: list[dict],
                    batch: str) -> int:
    ensure_bank_table()
    n = 0
    with get_connection() as conn:
        for x in lines:
            cur = conn.execute(
                "INSERT OR IGNORE INTO bank_lines "
                "(bank_account, cdate, amount, descr, batch) "
                "VALUES (?, ?, ?, ?, ?)",
                (bank_account, x["cdate"], x["amount"], x["descr"], batch))
            n += cur.rowcount
    return n


# ---------- محرك التصفية ----------
def _method_clause(methods: str) -> tuple[str, list]:
    kws = [k.strip() for k in (methods or "").split(",") if k.strip()]
    if not kws:
        return "1=0", []
    return ("(" + " OR ".join(["method LIKE ?"] * len(kws)) + ")",
            [f"%{k}%" for k in kws])


def daily_clearing(date_from: str, date_to: str,
                   gateway: str = "") -> list[dict]:
    """صف لكل يوم تشغيل: سيستم × بوابة × بنك + الحالة والتقسيم."""
    ensure_ops_table()
    ensure_bank_table()
    profiles = {p["name"]: p for p in list_profiles()}
    names = [gateway] if gateway else sorted(profiles)
    out = []
    with get_connection() as conn:
        for gname in names:
            prof = profiles.get(gname)
            if not prof:
                continue
            fee = float(prof["fee_fixed"] or 0)
            tol = float(prof["tolerance"] or 0.005)
            lag = int(prof["lag_days"] or 1)
            late_n = int(prof["late_days"] or 3)
            clause, params = _method_clause(prof["methods"])
            # S: سندات السيستم بهذه البوابة
            sys_rows = conn.execute(
                f"SELECT vdate, COALESCE(SUM(amount),0) AS s "
                f"FROM vouchers WHERE vtype='receipt' "
                f"AND COALESCE(cancelled,0)=0 AND {clause} "
                f"AND vdate BETWEEN ? AND ? GROUP BY vdate",
                (*params, date_from, date_to)).fetchall()
            sys_map = {r["vdate"]: float(r["s"]) for r in sys_rows}
            # G: عمليات البوابة المثبتة (شركتي/الأخرى + العدد)
            gw_rows = conn.execute(
                "SELECT op_date, status, COUNT(*) AS n, "
                "COALESCE(SUM(amount),0) AS g FROM gateway_ops "
                "WHERE gateway=? AND COALESCE(distributed,0)=1 "
                "AND (gw_paid IS NULL OR gw_paid=1) "
                "AND op_date BETWEEN ? AND ? GROUP BY op_date, status",
                (gname, date_from, date_to)).fetchall()
            gw: dict = {}
            for r in gw_rows:
                d = gw.setdefault(r["op_date"],
                                  {"mine": 0.0, "other": 0.0, "n": 0,
                                   "n_mine": 0, "n_other": 0})
                if (r["status"] or "") == "mine":
                    d["mine"] += float(r["g"])
                    d["n_mine"] += int(r["n"])
                    d["n"] += int(r["n"])
                elif (r["status"] or "") == "other":
                    d["other"] += float(r["g"])
                    d["n_other"] += int(r["n"])
                    d["n"] += int(r["n"])
                # «غير مصنّف» (والقديم diff) خارج التصفية تمامًا
            # حركات البنك للحساب المرتبط (نافذة موسعة ± أيام)
            bank_rows = conn.execute(
                "SELECT id, cdate, amount, descr FROM bank_lines "
                "WHERE (?='' OR bank_account=?) "
                "AND cdate BETWEEN date(?, '-7 days') AND date(?, '+7 days') "
                "ORDER BY cdate, id",
                (prof["bank_account"] or "", prof["bank_account"] or "",
                 date_from, date_to)).fetchall()
            bank = [dict(r) for r in bank_rows]
            used: set = set()
            days = sorted(set(sys_map) | set(gw))
            today = date.today().isoformat()
            for d in days:
                s = round(sys_map.get(d, 0.0), 3)
                g = gw.get(d, {"mine": 0.0, "other": 0.0, "n": 0,
                               "n_mine": 0, "n_other": 0})
                g_total = round(g["mine"] + g["other"], 3)
                expect = round(g_total - g["n"] * fee, 3)
                my_share = round(g["mine"] - g["n_mine"] * fee, 3)
                other_share = round(g["other"] - g["n_other"] * fee, 3)
                settle = ""
                try:
                    settle = add_business_days(
                        date.fromisoformat(d), lag).isoformat()
                except ValueError:
                    pass
                # ربط حركة البنك: D+N ± يوم عمل + تسامح
                bank_hit = None
                if settle and expect > 0:
                    lo = add_business_days(
                        date.fromisoformat(settle), -1).isoformat()
                    hi = add_business_days(
                        date.fromisoformat(settle), 1).isoformat()
                    best = None
                    for b in bank:
                        if b["id"] in used:
                            continue
                        if not (lo <= b["cdate"] <= hi):
                            continue
                        if abs(float(b["amount"]) - expect) > tol:
                            continue
                        key = (abs((date.fromisoformat(b["cdate"])
                                     - date.fromisoformat(settle)).days),
                               abs(float(b["amount"]) - expect))
                        if best is None or key < best[0]:
                            best = (key, b)
                    if best:
                        bank_hit = best[1]
                        used.add(bank_hit["id"])
                gap = round((g["mine"] - s), 3)
                if abs(gap) > tol:
                    status = "diff"
                elif bank_hit:
                    status = "done"
                elif settle and settle > today:
                    status = "pending"
                elif settle:
                    late_line = add_business_days(
                        date.fromisoformat(settle), late_n).isoformat()
                    status = "pending" if today <= late_line else "late"
                else:
                    status = "pending"
                out.append({
                    "gateway": gname, "day": d, "sys": s,
                    "g_mine": round(g["mine"], 3),
                    "g_other": round(g["other"], 3), "gap": gap,
                    "expect": expect if expect > 0 else 0.0,
                    "settle_date": settle,
                    "bank_amount": (round(float(bank_hit["amount"]), 3)
                                    if bank_hit else 0.0),
                    "bank_date": bank_hit["cdate"] if bank_hit else "",
                    "bank_descr": bank_hit["descr"] if bank_hit else "",
                    "my_share": my_share, "other_share": other_share,
                    "status": status,
                })
    out.sort(key=lambda r: (r["day"], r["gateway"]))
    return out


STATUS_LABEL = {"done": "✅ مكتمل", "pending": "⏳ منتظر",
                "diff": "⚠️ فرق", "late": "❌ متأخر"}


def distributed_mine_rows(date_from: str = "",
                          date_to: str = "") -> list[dict]:
    """صفوف شركتنا الموزَّعة + حالة يومها البنكية (لصفحة شركتنا)."""
    from datetime import date as _date
    ensure_ops_table()
    d1 = date_from or "2000-01-01"
    d2 = date_to or _date.today().isoformat()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT gateway, op_date, customer, phone, amount, ref "
            "FROM gateway_ops WHERE status='mine' "
            "AND COALESCE(distributed,0)=1 "
            "AND (gw_paid IS NULL OR gw_paid=1) "
            "AND op_date BETWEEN ? AND ? ORDER BY op_date, id",
            (d1, d2)).fetchall()
    cmap = {}
    for g in {r["gateway"] for r in rows}:
        for c in daily_clearing(d1, d2, g):
            cmap[(g, c["day"])] = c
    out = []
    for r in rows:
        c = cmap.get((r["gateway"], r["op_date"]), {})
        out.append({"gateway": r["gateway"], "date": r["op_date"],
                    "customer": r["customer"] or "", "phone": r["phone"],
                    "amount": round(float(r["amount"] or 0), 3),
                    "ref": r["ref"] or "",
                    "day_status": STATUS_LABEL.get(c.get("status", ""),
                                                   "—"),
                    "bank_amount": round(float(c.get("bank_amount") or 0),
                                         3)})
    return out


def distributed_other_rows(gateway: str = "", date_from: str = "",
                           date_to: str = "",
                           include_paid: bool = False) -> list[dict]:
    """V4.5.12.1: صفوف «شركة أخرى» الموزَّعة — لتظهر في كارت الموزَّع وتُرجَع.

    كانت `distributed_mine_rows` تجلب صفوف شركتنا فقط، فلا يجد المستخدم
    الصفوف الموزَّعة على «شركة أخرى» لإرجاعها.

    include_paid=False افتراضيًا: المسدَّد لا يُعرض في الموزَّع (قاعدة
    مستقرة: «المسدَّد خارج الموزَّع» — فكّه يعني حذف التسديد)، ويُبقى
    الخيار للمعاينة/التقارير.
    """
    from datetime import date as _date
    ensure_ops_table()
    ensure_alloc_table()
    d1 = date_from or "2000-01-01"
    d2 = date_to or _date.today().isoformat()
    q = ("SELECT id, gateway, op_date, customer, phone, amount, ref, "
         "COALESCE(paid_transfer_id,0) AS p FROM gateway_ops "
         "WHERE status='other' AND COALESCE(distributed,0)=1 "
         "AND op_date BETWEEN ? AND ?")
    args: list = [d1, d2]
    if gateway:
        q += " AND gateway=?"
        args.append(gateway)
    q += " ORDER BY op_date, id"
    out: list = []
    with get_connection() as conn:
        for r in conn.execute(q, tuple(args)).fetchall():
            try:
                paid = bool(r["p"]) or _op_paid_total(conn, r["id"]) > 0.0005
            except Exception:
                paid = bool(r["p"])
            if paid and not include_paid:
                continue
            out.append({"gateway": r["gateway"], "date": r["op_date"],
                        "customer": r["customer"] or "",
                        "phone": r["phone"] or "",
                        "amount": round(float(r["amount"] or 0), 3),
                        "ref": r["ref"] or "", "op_id": int(r["id"]),
                        "paid": paid, "side": "other"})
    return out


def unpaid_other_items() -> list[dict]:
    """بنود الأخرى الموزَّعة غير المسددة (للتحصيل بال☐).

    V4.49: صافي-الأساس — amount = المتبقي من الصافي (بعد الجزئي) ·
    paid = المدفوع · total = الصافي الأصلي (الإجمالي − عمولة البوابة).
    """
    ensure_ops_table()
    ensure_alloc_table()
    fees = gateway_fee_map()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT o.id, o.gateway, o.op_date, o.customer, o.phone, "
            "o.amount, o.ref, COALESCE(a.pd,0) AS pd, o.gw_paid, "
            "COALESCE(o.gw_paid_date,'') AS gw_date, "
            "COALESCE(o.gw_token,'') AS gw_token FROM gateway_ops o "
            "LEFT JOIN (SELECT op_id, SUM(amount) AS pd FROM "
            "transfer_op_allocs GROUP BY op_id) a ON a.op_id=o.id "
            "WHERE o.status='other' AND COALESCE(o.distributed,0)=1 "
            "AND COALESCE(o.paid_transfer_id,0)=0 "
            "AND (o.gw_paid IS NULL OR o.gw_paid=1) "
            "AND o.amount - COALESCE(a.pd,0) > 0.0005 "
            "ORDER BY o.op_date, o.id").fetchall()
    out = []
    for r in rows:
        _gross = round(float(r["amount"] or 0), 3)
        _pd = round(float(r["pd"] or 0), 3)
        _full = round(_gross - fees.get(r["gateway"], 0), 3)
        if _full - _pd <= 0.0005:
            continue
        out.append({"id": r["id"], "gateway": r["gateway"],
                    "date": r["op_date"], "customer": r["customer"] or "",
                    "phone": r["phone"], "amount": round(_full - _pd, 3),
                    "paid": _pd, "total": _full, "ref": r["ref"] or "",
                    "gross": _gross, "gw_paid": r["gw_paid"],
                    "gw_paid_date": r["gw_date"] or "",
                    "gw_token": r["gw_token"] or ""})
    return out


def unmatched_bank(date_from: str, date_to: str,
                   gateway: str = "") -> list[dict]:
    """حركات بنكية لم تربط بأي يوم (تحويلات/إيداعات أخرى)."""
    rows = daily_clearing(date_from, date_to, gateway)
    linked = {(r["bank_date"], r["bank_amount"]) for r in rows
              if r["bank_amount"] > 0}
    ensure_bank_table()
    with get_connection() as conn:
        bank = conn.execute(
            "SELECT cdate, amount, descr FROM bank_lines "
            "WHERE cdate BETWEEN date(?, '-7 days') AND date(?, '+7 days') "
            "ORDER BY cdate",
            (date_from, date_to)).fetchall()
    return [dict(b) for b in bank
            if (b["cdate"], round(float(b["amount"]), 3)) not in linked]


def other_statement(date_from: str, date_to: str) -> list[dict]:
    """مستحق الأخرى = صفوف «شركة أخرى» الموزَّعة (صافي بعد الرسوم).

    ينشأ عند التوزيع الصريح فقط — التصنيف وحده بلا أثر حسابي.
    """
    ensure_ops_table()
    fees = {p["name"]: round(float(p.get("fee_fixed") or 0), 3)
            for p in list_profiles(active_only=False)}
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT op_date, gateway, COUNT(*) AS n, "
            "COALESCE(SUM(amount),0) AS g FROM gateway_ops "
            "WHERE status='other' AND COALESCE(distributed,0)=1 "
            "AND (gw_paid IS NULL OR gw_paid=1) "
            "AND op_date BETWEEN ? AND ? "
            "GROUP BY op_date, gateway ORDER BY op_date, gateway",
            (date_from, date_to)).fetchall()
    run = 0.0
    out = []
    for r in rows:
        share = round(float(r["g"]) - int(r["n"]) * fees.get(r["gateway"], 0),
                      3)
        if share <= 0:
            continue
        run = round(run + share, 3)
        out.append({"day": r["op_date"], "gateway": r["gateway"],
                    "bank_date": "", "bank_amount": 0.0, "my_share": 0.0,
                    "other_share": share, "cumulative": run})
    return out


def other_external_report(date_from: str, date_to: str) -> dict:
    """تقرير المطابقة الخارجي — قراءة فقط (لا ينشئ أي قيد).

    صفوف «شركة أخرى» بالفترة (التاريخ/الاسم/المبلغ/سُدّد-أو-مستحق/
    تحصيل البوابة) + سطر الرسوم + إجمالي المسدد والمتبقي · الملغاة
    خارج الإجماليات · بلا أي تفاصيل داخلية.
    """
    ensure_ops_table()
    ensure_transfer_table()
    fees = {p["name"]: round(float(p.get("fee_fixed") or 0), 3)
            for p in list_profiles(active_only=False)}
    with get_connection() as conn:
        ops = conn.execute(
            "SELECT op_date, gateway, customer, amount, ref, "
            "COALESCE(paid_transfer_id,0) AS p, gw_paid, "
            "COALESCE(gw_paid_date,'') AS gw_date FROM gateway_ops "
            "WHERE status='other' AND COALESCE(distributed,0)=1 "
            "AND op_date BETWEEN ? AND ? ORDER BY op_date, id",
            (date_from, date_to)).fetchall()
    rows = []
    gross = fee_sum = 0.0
    tr = pe = un = 0.0
    n_cancelled = 0
    for o in ops:
        _gw = o["gw_paid"]
        if _gw == 2:
            n_cancelled += 1
            continue
        if _gw == 0:
            pe = round(pe + round(float(o["amount"] or 0), 3), 3)
            continue
        amt = round(float(o["amount"] or 0), 3)
        fee = fees.get(o["gateway"], 0)
        gross = round(gross + amt, 3)
        fee_sum = round(fee_sum + fee, 3)
        if _gw == 1:
            tr = round(tr + amt, 3)
        else:
            un = round(un + amt, 3)
        rows.append({"date": o["op_date"], "customer": o["customer"] or "",
                     "amount": amt, "gw": _gw,
                     "gw_date": o["gw_date"] or "",
                     "paid": bool(o["p"])})
    net = round(gross - fee_sum, 3)
    paid = other_transfers_total(date_from, date_to)
    opening = get_other_opening()["amount"]
    cum_dues = _dues_cumulative(date_to)
    cum_paid = other_transfers_total("", date_to)
    return {"rows": rows, "gross": gross, "fees": fee_sum, "net": net,
            "paid": paid, "opening": opening,
            "transferred": tr, "pending": pe, "unknown": un,
            "cancelled": n_cancelled,
            "remaining": round(opening + cum_dues - cum_paid, 3)}


def late_settlements(limit: int = 25) -> list[dict]:
    """تسويات تجاوزت المهلة بلا بنك (للتنبيهات)."""
    past = (date.today() - timedelta(days=60)).isoformat()
    rows = [r for r in daily_clearing(past, date.today().isoformat())
            if r["status"] == "late"]
    return rows[:limit]


def day_detail(gateway: str, day: str) -> dict:
    """تفاصيل يوم: عمليات البوابة + سندات السيستم + (البنك عبر daily)."""
    ensure_ops_table()
    profiles = {p["name"]: p for p in list_profiles()}
    prof = profiles.get(gateway, {})
    clause, params = _method_clause(prof.get("methods", ""))
    with get_connection() as conn:
        ops = conn.execute(
            "SELECT phone, amount, ref, status, receipt_no, customer "
            "FROM gateway_ops WHERE gateway=? AND op_date=? ORDER BY id",
            (gateway, day)).fetchall()
        vch = conn.execute(
            f"SELECT voucher_number, vdate, method, amount, description "
            f"FROM vouchers WHERE vtype='receipt' "
            f"AND COALESCE(cancelled,0)=0 AND {clause} AND vdate=? "
            f"ORDER BY id",
            (*params, day)).fetchall()
    return {"ops": [dict(r) for r in ops],
            "vouchers": [dict(r) for r in vch]}


# ---------- تحويلات الشركة الأخرى (V4.7/V4.8) ----------
OTHER_EXPENSE_ACCOUNT = "تسويات بوابات الدفع — الشركة الأخرى"


def ensure_transfer_table() -> None:
    with get_connection() as conn:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS other_company_transfers ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, "
            "tdate TEXT DEFAULT '', amount REAL DEFAULT 0, "
            "ref TEXT DEFAULT '', notes TEXT DEFAULT '', "
            "created_at TEXT DEFAULT (datetime('now','localtime')))")
        _cols = {r[1] for r in conn.execute(
            "PRAGMA table_info(other_company_transfers);")}
        for _col, _ddl in (
                ("method", "TEXT DEFAULT ''"),
                ("voucher_id", "INTEGER DEFAULT 0"),
                ("voucher_number", "TEXT DEFAULT ''")):
            if _col not in _cols:
                conn.execute("ALTER TABLE other_company_transfers "
                             f"ADD COLUMN {_col} {_ddl}")


GATEWAY_FEE_MARK = "[عمولة بوابات]"


def ensure_fee_table() -> None:
    """V4.52: ربط عمولة كل عملية بسندها — يمنع ازدواج سند العمولة."""
    with get_connection() as conn:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS gateway_fee_vouchers ("
            "op_id INTEGER PRIMARY KEY, voucher_id INTEGER NOT NULL, "
            "transfer_id INTEGER NOT NULL, fee REAL NOT NULL)")


def transfer_fee_voucher(transfer_id: int) -> dict | None:
    """V4.52: سند عمولة التسوية (الأحدث غير الملغى) أو None.

    العمولة مجمعة (السند الواحد قد يغطي عدة عمليات).
    """
    ensure_fee_table()
    with get_connection() as conn:
        return conn.execute(
            "SELECT f.voucher_id AS id, v.voucher_number, "
            "SUM(f.fee) AS fee "
            "FROM gateway_fee_vouchers f "
            "JOIN vouchers v ON v.id=f.voucher_id "
            "WHERE f.transfer_id=? AND COALESCE(v.cancelled,0)=0 "
            "GROUP BY f.voucher_id, v.voucher_number "
            "ORDER BY f.voucher_id DESC LIMIT 1",
            (int(transfer_id),)).fetchone()


def ensure_alloc_table() -> None:
    """توزيع التسديدات على العمليات (دعم الدفع الجزئي)."""
    with get_connection() as conn:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS transfer_op_allocs ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, transfer_id INTEGER "
            "DEFAULT 0, op_id INTEGER DEFAULT 0, amount REAL DEFAULT 0)")
        conn.execute(
            "CREATE INDEX IF NOT EXISTS ix_allocs_op ON "
            "transfer_op_allocs(op_id)")
        conn.execute(
            "CREATE INDEX IF NOT EXISTS ix_allocs_tr ON "
            "transfer_op_allocs(transfer_id)")


def _op_paid_total(conn, oid: int) -> float:
    r = conn.execute("SELECT COALESCE(SUM(amount),0) AS s FROM "
                     "transfer_op_allocs WHERE op_id=?", (oid,)).fetchone()
    return round(float(r["s"] or 0), 3)


def ensure_other_expense_account() -> tuple[int, str]:
    """حساب مصروف تسويات الأخرى (يُنشأ مرة واحدة) ← (id, الاسم)."""
    from services import expense_account_service as _ea
    for a in _ea.list_accounts():
        if (a.get("name") or "").strip() == OTHER_EXPENSE_ACCOUNT:
            return int(a["id"]), a.get("full_path") or a["name"]
    aid = _ea.add_account(OTHER_EXPENSE_ACCOUNT)
    acc = _ea.get_account(aid) or {}
    return aid, acc.get("full_path") or OTHER_EXPENSE_ACCOUNT


def _create_settlement_fee_voucher(conn, tid: int, tdate: str,
                                   parent_voucher_id: int, method: str,
                                   completed: list[int]) -> int | None:
    """V4.52: سند عمولة واحد للعمليات المكتملة بهذه التسوية.

    يُسجَّل تحت «عمولات بنكية» مربوطًا بسند التسوية الرئيسي. العمليات
    ذات العمولة الصفرية أو المسجلة بسند فعال تُتخطى — بلا ازدواج.
    """
    from services import lifecycle_service as _lc
    ensure_fee_table()
    _fees = gateway_fee_map()
    _items: list[tuple[int, float]] = []
    for oid in completed:
        _o = conn.execute("SELECT gateway FROM gateway_ops WHERE id=?",
                          (oid,)).fetchone()
        _fee = round(_fees.get(_o["gateway"], 0) if _o else 0, 3)
        if _fee <= 0.0005:
            continue
        _old = conn.execute(
            "SELECT COALESCE(v.cancelled,0) AS c FROM gateway_fee_vouchers f "
            "LEFT JOIN vouchers v ON v.id=f.voucher_id "
            "WHERE f.op_id=?", (oid,)).fetchone()
        if _old is not None and not _old["c"]:
            continue
        _items.append((oid, _fee))
    if not _items:
        return None
    _total = round(sum(f for _, f in _items), 3)
    _acc_id, _acc_name = _lc.ensure_bank_commission_account()
    _gws = sorted({conn.execute("SELECT gateway FROM gateway_ops WHERE id=?",
                                (oid,)).fetchone()["gateway"]
                   for oid, _ in _items})
    _party = f"بوابة {_gws[0]}" if len(_gws) == 1 else "بوابات الدفع"
    try:
        _vid = _lc.add_voucher(
            "payment", tdate, _party, method, _acc_name, _total,
            f"عمولة بوابات عن التسوية ت{tid} — {len(_items)} عملية",
            notes=GATEWAY_FEE_MARK, expense_account_id=_acc_id,
            parent_voucher_id=parent_voucher_id, suppress_bank_fee=True)
    except Exception as exc:
        raise ValueError(f"تعذر إنشاء سند العمولة: {exc}")
    for oid, _fee in _items:
        conn.execute("INSERT OR REPLACE INTO gateway_fee_vouchers "
                     "(op_id, voucher_id, transfer_id, fee) VALUES (?,?,?,?)",
                     (oid, _vid, tid, _fee))
    return _vid


def add_other_transfer(tdate: str, amount: float, ref: str = "",
                       notes: str = "", method: str = "",
                       op_ids: list | None = None,
                       create_voucher: bool = False) -> int:
    """تسجيل تسديد للأخرى — يرجع id. يرفع ValueError.

    op_ids: عمليات بوابة تُربط (يجب أن تكون غير مسددة).
    V4.49: التوزيع على الصافي (الإجمالي − عمولة البوابة).
    V4.5.13: الدفع الكامل أصبح طبيعيًا — يُقبل مبلغ أكبر من صافي المحدد،
    والفارق يظهر تلقائيًا في الواجهة كـ«فروقات عمولة» (مدفوع فوق المستحق)،
    ويُحمَّل على رصيد الأخرى للتوازن. وأُلغيت ميزة إنشاء سند الصرف من هذا
    المسار (ليست إيرادًا ولا مصروفًا): تُقبل المعلمة للتوافق مع النداءات
    القديمة وتُتجاهل تمامًا.
    """
    try:
        tdate = date.fromisoformat((tdate or "")[:10]).isoformat()
    except ValueError:
        raise ValueError("تاريخ التحويل غير صالح.")
    amount = round(float(amount or 0), 3)
    if amount <= 0:
        raise ValueError("مبلغ التحويل يجب أن يكون أكبر من صفر.")
    op_ids = [int(x) for x in (op_ids or [])]
    ensure_transfer_table()
    ensure_ops_table()
    ensure_alloc_table()
    _rem: dict[int, float] = {}
    _fees = gateway_fee_map()
    with get_connection() as conn:
        for oid in op_ids:
            r = conn.execute(
                "SELECT amount, gateway, COALESCE(paid_transfer_id,0) AS p, "
                "COALESCE(distributed,0) AS d, gw_paid FROM gateway_ops "
                "WHERE id=?", (oid,)).fetchone()
            if not r:
                raise ValueError("عملية غير موجودة ضمن المحدد.")
            if r["gw_paid"] in (0, 2):
                raise ValueError("عملية غير محصَّلة ضمن المحدد — لا يمكن "
                                 "تسديد ما لم يدفعه العميل.")
            if r["p"]:
                raise ValueError("عملية مسددة مسبقًا ضمن المحدد — أعد التحديد.")
            if not r["d"]:
                raise ValueError("عملية غير موزَّعة ضمن المحدد — وزّعها أولًا.")
            _pd = _op_paid_total(conn, oid)
            _full = round(float(r["amount"] or 0)
                          - _fees.get(r["gateway"], 0), 3)
            if _pd >= _full - 0.0005:
                raise ValueError("عملية مسددة مسبقًا ضمن المحدد — أعد التحديد.")
            _rem[oid] = round(_full - _pd, 3)
    # V4.5.13: لا رفض عند تجاوز الصافي — الدفع الكامل عادي والفارق «فروقات
    # عمولة» يُظهر في الواجهة. التوزيع يقتصر على المتاح لكل عملية.
    voucher_id, voucher_number = 0, ""
    # V4.5.13: `create_voucher` أُلغيت ميزتها — لا يُنشأ سند صرف من هنا أبدًا.
    _ = create_voucher
    with get_connection() as conn:
        cur = conn.execute(
            "INSERT INTO other_company_transfers "
            "(tdate, amount, ref, notes, method, voucher_id, voucher_number)"
            " VALUES (?, ?, ?, ?, ?, ?, ?)",
            (tdate, amount, (ref or "").strip(), (notes or "").strip(),
             (method or "").strip(), voucher_id, voucher_number))
        tid = int(cur.lastrowid)
        _left = amount
        _completed: list[int] = []
        for oid in op_ids:
            _take = min(_rem.get(oid, 0.0), round(_left, 3))
            if _take > 0.0005:
                conn.execute("INSERT INTO transfer_op_allocs "
                             "(transfer_id, op_id, amount) VALUES (?,?,?)",
                             (tid, oid, round(_take, 3)))
                _left = round(_left - _take, 3)
                if _take >= _rem.get(oid, 0.0) - 0.0005:
                    conn.execute("UPDATE gateway_ops SET paid_transfer_id=? "
                                 "WHERE id=?", (tid, oid))
                    _completed.append(oid)
        if voucher_id and _completed:   # مسار قديم فقط — لم يعد يُنشأ سند
            _create_settlement_fee_voucher(conn, tid, tdate, voucher_id,
                                           (method or "").strip()
                                           or "تحويل بنكي", _completed)
        return tid


def list_other_transfers(date_from: str = "",
                         date_to: str = "") -> list[dict]:
    """تحويلات الأخرى (تقرير) — مرتبة بالتاريخ ثم id."""
    ensure_transfer_table()
    q = ("SELECT id, tdate, amount, ref, notes, "
         "COALESCE(method,'') AS method, COALESCE(voucher_id,0) AS voucher_id, "
         "COALESCE(voucher_number,'') AS voucher_number FROM "
         "other_company_transfers WHERE 1=1")
    params: list = []
    if date_from:
        q += " AND tdate >= ?"
        params.append(date_from)
    if date_to:
        q += " AND tdate <= ?"
        params.append(date_to)
    q += " ORDER BY tdate, id"
    with get_connection() as conn:
        return [dict(r) for r in conn.execute(q, params).fetchall()]


def other_transfers_total(date_from: str = "",
                          date_to: str = "") -> float:
    rows = list_other_transfers(date_from, date_to)
    return round(sum(float(r["amount"]) for r in rows), 3)


def other_dues_total() -> float:
    """إجمالي المستحق للأخرى (تراكمي الكشف حتى اليوم)."""
    rows = other_statement("2000-01-01", date.today().isoformat())
    if not rows:
        return 0.0
    return round(float(rows[-1]["cumulative"]), 3)


def other_balance() -> dict:
    """{dues: المستحق + الافتتاحي, paid: المحوَّل, remaining: المتبقي}."""
    dues = round(get_other_opening()["amount"] + other_dues_total(), 3)
    paid = other_transfers_total()
    return {"dues": dues, "paid": paid,
            "remaining": round(dues - paid, 3)}


def delete_other_transfer(tid: int) -> None:
    """حذف تسديد: يحذف توزيعه ويعيد حساب اكتمال عملياته.

    المرتبط بسند صرف يُرفض.
    """
    ensure_transfer_table()
    ensure_ops_table()
    ensure_alloc_table()
    _fees = gateway_fee_map()
    with get_connection() as conn:
        r = conn.execute(
            "SELECT COALESCE(t.voucher_id,0) AS v, v.cancelled AS c "
            "FROM other_company_transfers t "
            "LEFT JOIN vouchers v ON v.id=t.voucher_id "
            "WHERE t.id=?", (int(tid),)).fetchone()
        if r and r["v"] and (r["c"] is None or not r["c"]):
            # V4.52: الرفض للسند النشط فقط — الملغى يسمح بالحذف وإعادة التسوية.
            raise ValueError("مرتبط بسند صرف — ألغِ السند أولًا من السندات.")
        ensure_fee_table()
        _fv = conn.execute(
            "SELECT v.voucher_number FROM gateway_fee_vouchers f "
            "JOIN vouchers v ON v.id=f.voucher_id "
            "WHERE f.transfer_id=? AND COALESCE(v.cancelled,0)=0 LIMIT 1",
            (int(tid),)).fetchone()
        if _fv:
            raise ValueError(
                f"مرتبط بسند عمولة ({_fv['voucher_number']}) — "
                "ألغِ السند أولًا من السندات.")
        touched = [x[0] for x in conn.execute(
            "SELECT DISTINCT op_id FROM transfer_op_allocs "
            "WHERE transfer_id=?", (int(tid),))]
        touched += [x[0] for x in conn.execute(
            "SELECT id FROM gateway_ops WHERE paid_transfer_id=?",
            (int(tid),))]
        conn.execute("DELETE FROM transfer_op_allocs WHERE transfer_id=?",
                     (int(tid),))
        for oid in set(touched):
            o = conn.execute("SELECT amount, gateway FROM gateway_ops "
                             "WHERE id=?", (oid,)).fetchone()
            if not o:
                continue
            _pd = _op_paid_total(conn, oid)
            _net = round(float(o["amount"] or 0)
                         - _fees.get(o["gateway"], 0), 3)
            if _pd >= _net - 0.0005:
                _last = conn.execute(
                    "SELECT COALESCE(MAX(transfer_id),0) AS m FROM "
                    "transfer_op_allocs WHERE op_id=?", (oid,)).fetchone()
                conn.execute("UPDATE gateway_ops SET paid_transfer_id=? "
                             "WHERE id=?", (int(_last["m"]), oid))
            else:
                conn.execute("UPDATE gateway_ops SET paid_transfer_id=0 "
                             "WHERE id=?", (oid,))
        conn.execute("DELETE FROM other_company_transfers WHERE id=?",
                     (int(tid),))


def get_transfer(tid: int) -> dict | None:
    ensure_transfer_table()
    with get_connection() as conn:
        r = conn.execute(
            "SELECT id, tdate, amount, ref, notes, COALESCE(method,'') AS "
            "method, COALESCE(voucher_id,0) AS voucher_id, "
            "COALESCE(voucher_number,'') AS voucher_number "
            "FROM other_company_transfers WHERE id=?",
            (int(tid),)).fetchone()
        return dict(r) if r else None


def transfer_ops(tid: int) -> list[dict]:
    """عمليات البوابة التي غطاها تسديد معين (+ المغطى لكل عملية)."""
    ensure_ops_table()
    ensure_alloc_table()
    with get_connection() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT o.id, o.gateway, o.op_date, o.phone, o.amount, o.ref, "
            "o.status, o.receipt_no, o.customer, "
            "COALESCE(a.alloc, o.amount) AS alloc FROM gateway_ops o "
            "LEFT JOIN (SELECT op_id, SUM(amount) AS alloc FROM "
            "transfer_op_allocs WHERE transfer_id=? GROUP BY op_id) a "
            "ON a.op_id=o.id WHERE o.paid_transfer_id=? "
            "OR a.op_id IS NOT NULL ORDER BY o.op_date, o.id",
            (int(tid), int(tid))).fetchall()]


def transfer_statement(tid: int) -> dict:
    """كشف تسديد جاهز للعرض/التصدير: بياناته + عملياته + الإجمالي."""
    t = get_transfer(tid) or {}
    ops = transfer_ops(tid) if t else []
    total = round(sum(float(o.get("alloc") or 0) for o in ops), 3)
    return {"transfer": t, "ops": ops, "total": total}


TR_STATUS = {"full": "🟢 مربوط كامل", "partial": "🟡 مربوط جزئي",
             "open": "🔵 على الحساب"}


def transfer_coverage(tid: int) -> dict:
    """V4.43: تغطية تسديد — عملياته مجموعة بتاريخ العملية.

    {"transfer", "ops": [{... gross/net/alloc}], "groups": [{date, n,
    gross, net, paid}], "totals": {amount, alloc, unalloc, n_ops,
    dates, dates_short}} — القديم بلا توزيع يُحسب بكامل المبلغ.
    """
    t = get_transfer(tid) or {}
    ops = transfer_ops(tid) if t else []
    fees = gateway_fee_map()
    _ops = []
    for o in ops:
        _g = round(float(o.get("amount") or 0), 3)
        _ops.append({
            "id": o["id"], "date": (o.get("op_date") or "")[:10],
            "gateway": o.get("gateway") or "",
            "customer": o.get("customer") or "",
            "phone": o.get("phone") or "",
            "gross": _g,
            "net": round(_g - fees.get(o.get("gateway") or "", 0), 3),
            "alloc": round(float(o.get("alloc") or 0), 3)})
    _groups: dict[str, dict] = {}
    for o in _ops:
        g = _groups.setdefault(o["date"], {"date": o["date"], "n": 0,
                                           "gross": 0.0, "net": 0.0,
                                           "paid": 0.0})
        g["n"] += 1
        g["gross"] = round(g["gross"] + o["gross"], 3)
        g["net"] = round(g["net"] + o["net"], 3)
        g["paid"] = round(g["paid"] + o["alloc"], 3)
    _dates = sorted(_groups)
    _alloc = round(sum(o["alloc"] for o in _ops), 3)
    _amount = round(float(t.get("amount") or 0), 3)
    return {
        "transfer": t, "ops": _ops,
        "groups": [_groups[d] for d in _dates],
        "totals": {
            "amount": _amount, "alloc": _alloc,
            "unalloc": round(_amount - _alloc, 3),
            "n_ops": len(_ops), "dates": _dates,
            "dates_short": "، ".join(f"{d[8:10]}/{d[5:7]}"
                                     for d in _dates),
        }}


def list_transfers_summary(date_from: str = "",
                           date_to: str = "") -> list[dict]:
    """V4.43: التسديدات مع تغطيتها — لشاشة المجمعات (مرتبة كالأصل)."""
    trs = list_other_transfers(date_from, date_to)
    if not trs:
        return []
    ensure_alloc_table()
    ensure_ops_table()
    with get_connection() as conn:
        _al = {r["transfer_id"]: r for r in conn.execute(
            "SELECT transfer_id, SUM(amount) AS s, "
            "COUNT(DISTINCT op_id) AS n FROM transfer_op_allocs "
            "GROUP BY transfer_id").fetchall()}
        _dt: dict[int, set] = {}
        for r in conn.execute(
                "SELECT x.transfer_id AS tid, o.op_date AS d FROM "
                "transfer_op_allocs x JOIN gateway_ops o ON o.id=x.op_id "
                "GROUP BY x.transfer_id, o.op_date").fetchall():
            if r["d"]:
                _dt.setdefault(r["tid"], set()).add(r["d"][:10])
        _leg: dict[int, list] = {}
        for r in conn.execute(
                "SELECT o.paid_transfer_id AS tid, o.op_date, o.amount FROM "
                "gateway_ops o WHERE o.paid_transfer_id <> 0 AND o.id NOT IN "
                "(SELECT op_id FROM transfer_op_allocs "
                "WHERE transfer_id=o.paid_transfer_id)").fetchall():
            _leg.setdefault(r["tid"], []).append(r)
    out = []
    for t in trs:
        _a = _al.get(t["id"])
        _alloc = round(float((_a["s"] if _a else 0) or 0), 3)
        _n = int(_a["n"]) if _a else 0
        _dates = set(_dt.get(t["id"], set()))
        for r in _leg.get(t["id"], []):
            _alloc = round(_alloc + float(r["amount"] or 0), 3)
            _n += 1
            if r["op_date"]:
                _dates.add(r["op_date"][:10])
        _un = round(float(t["amount"] or 0) - _alloc, 3)
        _st = "open" if _n == 0 else ("full" if _un <= 0.0005
                                      else "partial")
        _dl = sorted(_dates)
        out.append({"id": t["id"], "tdate": (t["tdate"] or "")[:10],
                    "method": t["method"] or "", "ref": t["ref"] or "",
                    "notes": t["notes"] or "",
                    "amount": round(float(t["amount"] or 0), 3),
                    "alloc": _alloc, "unalloc": _un, "n_ops": _n,
                    "dates": _dl,
                    "dates_short": "، ".join(f"{d[8:10]}/{d[5:7]}"
                                             for d in _dl),
                    "status": _st})
    return out


# ---------- الرصيد الافتتاحي للأخرى (V4.8) ----------
def ensure_opening_table() -> None:
    with get_connection() as conn:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS other_company_opening ("
            "id INTEGER PRIMARY KEY CHECK (id=1), amount REAL DEFAULT 0, "
            "note TEXT DEFAULT '', "
            "updated_at TEXT DEFAULT (datetime('now','localtime')))")
        conn.execute("INSERT OR IGNORE INTO other_company_opening (id, amount)"
                     " VALUES (1, 0)")


def get_other_opening() -> dict:
    ensure_opening_table()
    with get_connection() as conn:
        r = conn.execute("SELECT amount, note FROM other_company_opening "
                         "WHERE id=1").fetchone()
    return {"amount": round(float(r["amount"] or 0), 3) if r else 0.0,
            "note": (r["note"] or "") if r else ""}


def set_other_opening(amount: float, note: str = "") -> None:
    amount = round(float(amount or 0), 3)
    if amount < 0:
        raise ValueError("الرصيد الافتتاحي لا يكون سالبًا.")
    ensure_opening_table()
    with get_connection() as conn:
        conn.execute("UPDATE other_company_opening SET amount=?, note=?, "
                     "updated_at=datetime('now','localtime') WHERE id=1",
                     (amount, (note or "").strip()))


# ---------- «مسدد مسبقًا» عند إعادة الرفع (V4.8) ----------
def linked_ops_breakdown(op_ids: list[int]) -> dict:
    """V4.50: إجمالي/عمولة/صافي عمليات مربوطة — لصندوق الشرح.

    {n, gross, fee, net} — الإجمالي من الملف للمعرفة، والصافي هو
    المستحق الفعلي بعد العمولة.
    """
    ensure_ops_table()
    fees = gateway_fee_map()
    gross = fee = 0.0
    n = 0
    if op_ids:
        with get_connection() as conn:
            for oid in {int(x) for x in op_ids}:
                r = conn.execute(
                    "SELECT amount, gateway FROM gateway_ops WHERE id=?",
                    (oid,)).fetchone()
                if not r:
                    continue
                _g = round(float(r["amount"] or 0), 3)
                _f = round(fees.get(r["gateway"], 0), 3)
                gross = round(gross + _g, 3)
                fee = round(fee + _f, 3)
                n += 1
    return {"n": n, "gross": gross, "fee": fee,
            "net": round(gross - fee, 3)}


def paid_lookup(gateway: str) -> dict:
    """{(التاريخ, الهاتف, المبلغ, المرجع): (تسديد, مرجعه)} للمسددة."""
    ensure_ops_table()
    ensure_transfer_table()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT o.op_date, o.phone, o.amount, o.ref, "
            "o.paid_transfer_id, COALESCE(t.ref,'') AS tref "
            "FROM gateway_ops o LEFT JOIN other_company_transfers t "
            "ON t.id=o.paid_transfer_id "
            "WHERE o.gateway=? AND COALESCE(o.paid_transfer_id,0)>0",
            (gateway,)).fetchall()
    return {((r["op_date"] or ""), (r["phone"] or ""),
             round(float(r["amount"] or 0), 3), (r["ref"] or "")):
            (r["paid_transfer_id"], r["tref"]) for r in rows}


def mark_paid_before(matched: list[dict], gateway: str) -> int:
    """يعلّم الصفوف المسددة مسبقًا (paid_before) — يرجع عددها."""
    if not gateway:
        return 0
    look = paid_lookup(gateway)
    n = 0
    for r in matched:
        key = ((r.get("date") or ""), (r.get("phone") or "").strip(),
               round(float(r.get("amount") or 0), 3),
               (r.get("ref") or "").strip())
        if key in look:
            tid, tref = look[key]
            r["paid_before"] = f"تسديد #{tid}" + (f" ({tref})" if tref else "")
            n += 1
    return n


def seen_lookup(gateway: str) -> dict:
    """كل المثبت مسبقًا للبوابة: المفتاح ← (الحالة السابقة).

    المفتاح الموحد للشركتين = (التاريخ، الهاتف، المبلغ، المرجع)
    والمرجع هو Invoice Number من كشف Hesabe.
    """
    ensure_ops_table()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT op_date, phone, amount, ref, status "
            "FROM gateway_ops WHERE gateway=?", (gateway,)).fetchall()
    def _norm(s):
        return s if s in ("mine", "other", "unmarked") else "unmarked"
    return {((r["op_date"] or ""), (r["phone"] or ""),
             round(float(r["amount"] or 0), 3), (r["ref"] or "")):
            _norm(r["status"] or "") for r in rows}


def manual_lookup(gateway: str) -> dict:
    """المصنف يدويًا: المفتاح ← حالته المحفوظة (يُحترم عند إعادة الرفع)."""
    ensure_ops_table()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT op_date, phone, amount, ref, status FROM gateway_ops "
            "WHERE gateway=? AND COALESCE(classified_manual,0)=1",
            (gateway,)).fetchall()
    return {((r["op_date"] or ""), (r["phone"] or ""),
             round(float(r["amount"] or 0), 3), (r["ref"] or "")):
            (r["status"] or "") for r in rows}


def suggest_row_classes(gateway: str, rows: list[dict]) -> dict:
    """اقتراح توزيع جاهز (V4.17) — مرحلة حرة يراجعها المستخدم قبل ◆ توزيع.

    لكل صف غير مقفل (لا موزَّع ولا مسدد ولا يدوي/مثبت):
      1) المطابق الواضح (حالة الدخول شركتنا) ← شركتنا · السبب: مطابق
      2) سابقة الهاتف (الموزَّعة فقط، نفس البوابة): إجماع ← نفس الجهة ·
         السبب: سابقة · تعارض ← يبقى غير مصنّف + علَم سوابق متعارضة
      3) الهاتف في دليل عملائنا ← شركتنا · السبب: عميل معروف
      4) غير ذلك ← شركة أخرى · السبب: بلا سابقة
    يضع status + _suggested + _suggest_reason (أو _conflict) ويرجع عدّادًا.
    """
    from services.gateway_service import canon_phone
    counts = {"mine": 0, "other": 0, "unmarked": 0, "conflict": 0}
    if not gateway or not rows:
        return counts
    ensure_ops_table()
    with get_connection() as conn:
        hist = conn.execute(
            "SELECT phone, status FROM gateway_ops "
            "WHERE gateway=? AND COALESCE(distributed,0)=1 "
            "AND status IN ('mine','other')", (gateway,)).fetchall()
        try:
            custs = conn.execute(
                "SELECT phone FROM customers "
                "WHERE COALESCE(TRIM(phone),'')<>''").fetchall()
        except Exception:
            custs = []
    by_phone: dict[str, set] = {}
    for h in hist:
        nph = canon_phone(h["phone"])
        if nph:
            by_phone.setdefault(nph, set()).add(h["status"])
    our_phones = {canon_phone(c["phone"]) for c in custs}
    our_phones.discard("")
    dist = distributed_lookup(gateway)
    for r in rows:
        key = ((r.get("date") or ""), (r.get("phone") or "").strip(),
               round(float(r.get("amount") or 0), 3),
               (r.get("ref") or "").strip())
        if key in dist or r.get("distributed") or r.get("paid_before") \
                or r.get("manual"):
            continue
        if (r.get("status") or "") == "mine":
            r["_suggested"] = True
            r["_suggest_reason"] = "مطابق"
            counts["mine"] += 1
            continue
        sides = by_phone.get(canon_phone(r.get("phone")), set())
        if len(sides) > 1:
            r["status"] = "unmarked"
            r["_conflict"] = True
            counts["unmarked"] += 1
            counts["conflict"] += 1
        elif len(sides) == 1:
            side = next(iter(sides))
            r["status"] = side
            r["_suggested"] = True
            r["_suggest_reason"] = ("سابقة: شركتنا" if side == "mine"
                                    else "سابقة: شركة أخرى")
            counts[side] += 1
        elif canon_phone(r.get("phone")) in our_phones:
            r["status"] = "mine"
            r["_suggested"] = True
            r["_suggest_reason"] = "عميل معروف"
            counts["mine"] += 1
        else:
            r["status"] = "other"
            r["_suggested"] = True
            r["_suggest_reason"] = "بلا سابقة"
            counts["other"] += 1
    return counts


def restore_manual(matched: list[dict], gateway: str) -> int:
    """يسترجع التصنيف اليدوي المحفوظ فوق نتيجة التلقائي — يرجع عددها."""
    if not gateway:
        return 0
    look = manual_lookup(gateway)
    n = 0
    for r in matched:
        key = ((r.get("date") or ""), (r.get("phone") or "").strip(),
               round(float(r.get("amount") or 0), 3),
               (r.get("ref") or "").strip())
        if key in look:
            r["status"] = look[key] or "unmarked"
            r["manual"] = True
            n += 1
    return n


def set_row_class(gateway: str, key: tuple, status: str) -> bool:
    """تصنيف يدوي حر (مرحلة مؤقتة — بلا أثر حسابي) · يرفع ValueError.

    الموزَّعة مرفوضة (ارجعها للتصنيف أولًا) · المسددة مرفوضة ·
    غير المثبتة مرفوضة · القفل يبدأ من لحظة التوزيع لا التصنيف.
    """
    if status not in ("mine", "other", "unmarked"):
        raise ValueError("تصنيف غير صالح.")
    ensure_ops_table()
    dt, ph, am, rf = key
    with get_connection() as conn:
        r = conn.execute(
            "SELECT id, COALESCE(paid_transfer_id,0) AS p, "
            "COALESCE(distributed,0) AS d FROM gateway_ops "
            "WHERE gateway=? AND op_date=? AND phone=? AND amount=? "
            "AND ref=?", (gateway, dt, ph, am, rf)).fetchone()
        if not r:
            raise ValueError("العملية غير مثبتة — أعد الاستيراد.")
        if r["p"]:
            raise ValueError("العملية مسددة — لا يمكن إعادة تصنيفها.")
        if r["d"]:
            raise ValueError("العملية موزَّعة — مقفلة (ارجعها للتصنيف أولًا).")
        conn.execute("UPDATE gateway_ops SET status=? WHERE id=?",
                     (status, r["id"]))
    return True


def distributed_lookup(gateway: str) -> set:
    """مفاتيح الموزَّعة للبوابة (مقفل — خارج شاشة التصنيف)."""
    ensure_ops_table()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT op_date, phone, amount, ref FROM gateway_ops "
            "WHERE gateway=? AND COALESCE(distributed,0)=1",
            (gateway,)).fetchall()
    return {((r["op_date"] or ""), (r["phone"] or ""),
             round(float(r["amount"] or 0), 3), (r["ref"] or ""))
            for r in rows}


def distribute_rows(gateway: str, keys: list[tuple]) -> dict:
    """◆ التوزيع الصريح: مصنَّف ← معتمد محاسبيًا · يرجع العدّادات.

    (mine, other, skipped): المعلّق/الموزَّع مسبقًا/المسدد/المفقود
    يُتخطى — لا مسار يتجاوز هذه الدالة.
    """
    ensure_ops_table()
    n_mine = n_other = skipped = 0
    with get_connection() as conn:
        for key in keys:
            dt, ph, am, rf = key
            r = conn.execute(
                "SELECT id, status, COALESCE(paid_transfer_id,0) AS p, "
                "COALESCE(distributed,0) AS d, gw_paid FROM gateway_ops "
                "WHERE gateway=? AND op_date=? AND phone=? AND amount=? "
                "AND ref=?", (gateway, dt, ph, float(am), rf)).fetchone()
            if (not r or r["d"] or r["p"]
                    or (r["status"] or "") not in ("mine", "other")
                    or r["gw_paid"] in (0, 2)):
                skipped += 1
                continue
            conn.execute("UPDATE gateway_ops SET distributed=1, "
                         "classified_manual=1, "
                         "distributed_at=datetime('now','localtime') "
                         "WHERE id=?", (r["id"],))
            if r["status"] == "mine":
                n_mine += 1
            else:
                n_other += 1
    return {"mine": n_mine, "other": n_other, "skipped": skipped}


def undistribute_rows(gateway: str, keys: list[tuple]) -> dict:
    """↩ إرجاع للتصنيف: موزَّع ← مرحلة مؤقتة (تصنيفه يبقى، حر التغيير).

    يرجع (ok, paid_skipped): المسدد يُتخطى (فكّه حذف التسديد).
    """
    ensure_ops_table()
    ensure_alloc_table()
    ok = paid = 0
    with get_connection() as conn:
        for key in keys:
            dt, ph, am, rf = key
            r = conn.execute(
                "SELECT id, COALESCE(paid_transfer_id,0) AS p, "
                "COALESCE(distributed,0) AS d FROM gateway_ops "
                "WHERE gateway=? AND op_date=? AND phone=? AND amount=? "
                "AND ref=?", (gateway, dt, ph, float(am), rf)).fetchone()
            if not r or not r["d"]:
                continue
            if r["p"] or _op_paid_total(conn, r["id"]) > 0.0005:
                paid += 1
                continue
            conn.execute("UPDATE gateway_ops SET distributed=0, "
                         "classified_manual=0, distributed_at='' "
                         "WHERE id=?", (r["id"],))
            ok += 1
    return {"ok": ok, "paid_skipped": paid}


def mark_seen_before(matched: list[dict], gateway: str) -> int:
    """يعلّم الصفوف المثبتة مسبقًا (seen_before) — يرجع عددها.

    شركتي (mine/diff) = مطابقة مسبقًا (منتهية) · الأخرى = مسجلة
    بانتظار السداد · المسددة تُترك لتعليم paid_before.
    """
    if not gateway:
        return 0
    look = seen_lookup(gateway)
    n = 0
    for r in matched:
        if r.get("paid_before"):
            continue
        key = ((r.get("date") or ""), (r.get("phone") or "").strip(),
               round(float(r.get("amount") or 0), 3),
               (r.get("ref") or "").strip())
        if key in look:
            r["seen_before"] = look[key]
            n += 1
    return n


def op_ids_for(gateway: str, keys: list[tuple]) -> list[int]:
    """ids العمليات الموزَّعة غير المسددة (تجاهل المسدد وغير الموزَّع
    وغير المحصَّل)."""
    ensure_ops_table()
    out = []
    with get_connection() as conn:
        for d, p, a, ref in keys:
            r = conn.execute(
                "SELECT id, COALESCE(paid_transfer_id,0) AS t, "
                "COALESCE(distributed,0) AS d, gw_paid "
                "FROM gateway_ops WHERE gateway=? AND op_date=? AND phone=? "
                "AND ABS(amount-?)<0.0005 AND ref=?",
                (gateway, d, p, float(a), ref)).fetchone()
            if r and not r["t"] and r["d"] and r["gw_paid"] not in (0, 2):
                out.append(int(r["id"]))
    return out


# ---------- الكشف الدفتري الموحد (V4.8) ----------
def _dues_cumulative(date_to: str) -> float:
    rows = other_statement("2000-01-01", date_to)
    return round(float(rows[-1]["cumulative"]), 3) if rows else 0.0


def list_upload_batches() -> list[dict]:
    """الرفعات (ملف + بوابة + يوم رفع): مجاميع للأحدث أولًا.

    day = يوم الرفع من created_at · due_amount = مستحق الأخرى الآن
    (موزَّعة غير مسددة) داخل الرفعة.
    """
    ensure_ops_table()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT o.gateway, COALESCE(o.batch,'') AS batch, "
            "SUBSTR(COALESCE(o.created_at,''),1,10) AS day, "
            "MIN(o.created_at) AS uploaded_at, COUNT(*) AS n_ops, "
            "SUM(COALESCE(o.distributed,0)) AS n_dist, "
            "SUM(CASE WHEN COALESCE(o.paid_transfer_id,0)>0 THEN 1 "
            "ELSE 0 END) AS n_paid, "
            "SUM(CASE WHEN COALESCE(o.distributed,0)=1 "
            "AND COALESCE(o.paid_transfer_id,0)=0 AND o.status='other' "
            "THEN o.amount ELSE 0 END) AS due_amount "
            "FROM gateway_ops o GROUP BY o.gateway, "
            "COALESCE(o.batch,''), SUBSTR(COALESCE(o.created_at,''),1,10) "
            "ORDER BY uploaded_at DESC").fetchall()
    return [{"gateway": r["gateway"], "batch": r["batch"] or "",
             "day": r["day"] or "", "uploaded_at": r["uploaded_at"] or "",
             "n_ops": r["n_ops"] or 0, "n_dist": r["n_dist"] or 0,
             "n_paid": r["n_paid"] or 0,
             "due_amount": round(float(r["due_amount"] or 0), 3)}
            for r in rows]


def upload_batch_ops(gateway: str, batch: str, day: str) -> list[dict]:
    """عمليات رفعة واحدة كاملة بحالاتها (للضغط المزدوج)."""
    ensure_ops_table()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT op_date, ref, customer, phone, amount, status, "
            "COALESCE(distributed,0) AS dist, "
            "COALESCE(paid_transfer_id,0) AS ptid, gw_paid "
            "FROM gateway_ops WHERE gateway=? AND COALESCE(batch,'')=? "
            "AND SUBSTR(COALESCE(created_at,''),1,10)=? "
            "ORDER BY op_date, id", (gateway, batch, day)).fetchall()
    return [{"date": r["op_date"] or "", "ref": r["ref"] or "",
             "customer": r["customer"] or "", "phone": r["phone"] or "",
             "amount": round(float(r["amount"] or 0), 3),
             "status": r["status"] or "unmarked",
             "distributed": bool(r["dist"]), "paid": bool(r["ptid"]),
             "gw_paid": r["gw_paid"]} for r in rows]


def gateway_fee_map() -> dict:
    """عمولة كل بوابة {name: fee_fixed} — تُخصم من مستحق الأخرى (صافي)."""
    return {p["name"]: round(float(p.get("fee_fixed") or 0), 3)
            for p in list_profiles(active_only=False)}


_DUES_ELIGIBLE = ("o.status='other' AND COALESCE(o.distributed,0)=1 "
                  "AND (o.gw_paid IS NULL OR o.gw_paid=1)")


def upload_batches_report() -> list[dict]:
    """تقرير الرفعات الموحد (V4.34) — كل (بوابة + ملف + يوم رفع) = سطر.

    dues = صافي مستحق الأخرى بعد الرسوم (نفس تعريف other_statement
    حرفيًا: other + موزَّعة + تحصيل غير مستبعد) · paid = التوزيعات على
    عمليات الرفعة · remaining = dues − paid (قد يحمل غبار الرسوم السالب
    للمسدد كاملًا — نفس سلوك other_balance، بلا تجميل يكسر المطابقة).
    مرتب من الأحدث (نفس ترتيب list_upload_batches).
    """
    ensure_ops_table()
    ensure_alloc_table()
    fees = gateway_fee_map()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT o.gateway, COALESCE(o.batch,'') AS batch, "
            "SUBSTR(COALESCE(o.created_at,''),1,10) AS day, "
            "MIN(o.created_at) AS uploaded_at, COUNT(*) AS n_ops, "
            "COALESCE(SUM(o.amount),0) AS gross, "
            "SUM(CASE WHEN o.status='mine' AND COALESCE(o.distributed,0)=1 "
            "THEN o.amount ELSE 0 END) AS mine_total, "
            "SUM(CASE WHEN o.status='mine' AND COALESCE(o.distributed,0)=1 "
            "THEN 1 ELSE 0 END) AS n_mine, "
            "SUM(CASE WHEN o.status='other' AND COALESCE(o.distributed,0)=1 "
            "THEN o.amount ELSE 0 END) AS other_gross, "
            "SUM(CASE WHEN COALESCE(o.distributed,0)=0 THEN 1 ELSE 0 END) "
            "AS n_undist, "
            "SUM(CASE WHEN (o.status IS NULL OR o.status='' "
            "OR o.status='unmarked' OR o.status='diff') "
            "AND COALESCE(o.distributed,0)=0 THEN 1 ELSE 0 END) "
            "AS n_unmarked, "
            "COALESCE(SUM(a.pd),0) AS paid "
            "FROM gateway_ops o "
            "LEFT JOIN (SELECT op_id, SUM(amount) AS pd FROM "
            "transfer_op_allocs GROUP BY op_id) a ON a.op_id=o.id "
            "GROUP BY o.gateway, COALESCE(o.batch,''), "
            "SUBSTR(COALESCE(o.created_at,''),1,10) "
            "ORDER BY uploaded_at DESC").fetchall()
        dues = conn.execute(
            "SELECT o.gateway, COALESCE(o.batch,'') AS batch, "
            "SUBSTR(COALESCE(o.created_at,''),1,10) AS day, "
            "COALESCE(SUM(o.amount),0) AS g, COUNT(*) AS n "
            "FROM gateway_ops o WHERE " + _DUES_ELIGIBLE + " "
            "GROUP BY o.gateway, COALESCE(o.batch,''), "
            "SUBSTR(COALESCE(o.created_at,''),1,10)").fetchall()
        names = conn.execute(
            "SELECT DISTINCT o.gateway, COALESCE(o.batch,'') AS batch, "
            "SUBSTR(COALESCE(o.created_at,''),1,10) AS day, "
            "TRIM(COALESCE(o.customer,'')) AS customer FROM gateway_ops o "
            "ORDER BY customer").fetchall()
        mine_names = conn.execute(
            "SELECT DISTINCT o.gateway, COALESCE(o.batch,'') AS batch, "
            "SUBSTR(COALESCE(o.created_at,''),1,10) AS day, "
            "TRIM(COALESCE(o.customer,'')) AS customer FROM gateway_ops o "
            "WHERE o.status='mine' AND COALESCE(o.distributed,0)=1 "
            "ORDER BY customer").fetchall()
    _dues = {(r["gateway"], r["batch"], r["day"]):
             round(float(r["g"]) - int(r["n"])
                   * fees.get(r["gateway"], 0), 3) for r in dues}
    _names: dict = {}
    for r in names:
        if r["customer"]:
            _names.setdefault((r["gateway"], r["batch"], r["day"]),
                              []).append(r["customer"])
    _mine_names: dict = {}
    for r in mine_names:
        if r["customer"]:
            _mine_names.setdefault(
                (r["gateway"], r["batch"], r["day"]), []).append(
                r["customer"])
    out = []
    for r in rows:
        key = (r["gateway"], r["batch"], r["day"])
        _d = _dues.get(key, 0.0)
        _p = round(float(r["paid"] or 0), 3)
        _rem = round(_d - _p, 3)
        if _d <= 0:
            _st = "none"
        elif _rem <= 0.0005:
            _st = "paid"
        elif _p > 0.0005:
            _st = "partial"
        else:
            _st = "unpaid"
        out.append({
            "gateway": r["gateway"], "batch": r["batch"] or "",
            "day": r["day"] or "", "uploaded_at": r["uploaded_at"] or "",
            "n_ops": r["n_ops"] or 0,
            "total_amount": round(float(r["gross"] or 0), 3),
            "mine_total": round(float(r["mine_total"] or 0), 3),
            "n_mine": r["n_mine"] or 0,
            "mine_net": round(float(r["mine_total"] or 0)
                               - int(r["n_mine"] or 0)
                               * fees.get(r["gateway"], 0), 3),
            "other_gross": round(float(r["other_gross"] or 0), 3),
            "n_undistributed": r["n_undist"] or 0,
            "n_unmarked": r["n_unmarked"] or 0,
            "dues": _d, "paid": _p, "remaining": _rem, "status": _st,
            "names": _names.get(key, []),
            "mine_names": _mine_names.get(key, [])})
    return out




def paid_other_items() -> list[dict]:
    """عمليات الأخرى المسددة كاملًا (تبويب المسدد) — الأحدث أولًا.

    [{id, gateway, date, customer, phone, ref, gross, fee, net, paid,
    transfer_id, tdate, method}] — التوزيعات غطت الصافي (V4.49: نفس
    تعريف شاشة المستحق) أو علم قديم (يُفترض مدفوعًا بالكامل) ·
    التسديد = الأحدث تغطية.
    """
    ensure_ops_table()
    ensure_alloc_table()
    ensure_transfer_table()
    fees = gateway_fee_map()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT o.id, o.gateway, o.op_date, "
            "COALESCE(o.customer,'') AS customer, "
            "COALESCE(o.phone,'') AS phone, "
            "COALESCE(o.ref,'') AS ref, o.amount, "
            "COALESCE(a.pd,0) AS pd, COALESCE(a.ltid,0) AS ltid, "
            "COALESCE(o.paid_transfer_id,0) AS pt "
            "FROM gateway_ops o "
            "LEFT JOIN (SELECT op_id, SUM(amount) AS pd, "
            "MAX(transfer_id) AS ltid FROM transfer_op_allocs "
            "GROUP BY op_id) a ON a.op_id=o.id "
            "WHERE o.status='other' AND COALESCE(o.distributed,0)=1 "
            "ORDER BY o.op_date DESC, o.id DESC").fetchall()
        _trs = {t["id"]: {"tdate": t["tdate"], "method": t["method"]}
                for t in conn.execute(
            "SELECT id, tdate, COALESCE(method,'') AS method FROM "
            "other_company_transfers").fetchall()}
    out = []
    for r in rows:
        _g = round(float(r["amount"] or 0), 3)
        _fee = round(fees.get(r["gateway"], 0), 3)
        _net = round(_g - _fee, 3)
        _pd = round(float(r["pd"] or 0), 3)
        _legacy = int(r["pt"] or 0) > 0
        if not ((_pd > 0.0005 and _net - _pd <= 0.0005) or _legacy):
            continue
        if _legacy and _pd <= 0.0005:
            _pd = _net
        _tid = int(r["ltid"] or 0) or int(r["pt"] or 0)
        _t = _trs.get(_tid, {})
        out.append({"id": r["id"], "gateway": r["gateway"] or "",
                    "date": r["op_date"] or "",
                    "customer": r["customer"] or "",
                    "phone": r["phone"] or "", "ref": r["ref"] or "",
                    "gross": _g, "fee": _fee, "net": _net, "paid": _pd,
                    "transfer_id": _tid, "tdate": _t.get("tdate") or "",
                    "method": _t.get("method") or ""})
    return out

def other_ops_status(date_from: str = "", date_to: str = "") -> list[dict]:
    """V4.40: عمليات الأخرى الموزعة بحالتها (تقرير العمليات).

    [{id, date, customer, phone, gateway, ref, gross, fee, net, paid,
    remaining, status, transfer_id, tdate, method, effective_date}] — status:
    paid (مغطى بالكامل أو علم قديم) / partial / due · المتبقي
    صافي-المدفوع (لغة التقارير) · مرتبة بالتاريخ ثم id.
    """
    from datetime import date as _date
    ensure_ops_table()
    ensure_alloc_table()
    ensure_transfer_table()
    fees = gateway_fee_map()
    d1 = date_from or "2000-01-01"
    d2 = date_to or _date.today().isoformat()
    with get_connection() as conn:
        # V8.11: تحقق من وجود عمود effective_date
        _cols = {r[1] for r in conn.execute(
            "PRAGMA table_info(gateway_ops);")}
        _ed_col = (", COALESCE(o.effective_date,'') AS effective_date"
                   if "effective_date" in _cols else ", '' AS effective_date")
        rows = conn.execute(
            "SELECT o.id, o.gateway, o.op_date, "
            "COALESCE(o.customer,'') AS customer, "
            "COALESCE(o.phone,'') AS phone, "
            "COALESCE(o.ref,'') AS ref, o.amount, "
            "COALESCE(a.pd,0) AS pd, COALESCE(a.ltid,0) AS ltid, "
            "COALESCE(o.paid_transfer_id,0) AS pt "
            f"{_ed_col} "
            "FROM gateway_ops o "
            "LEFT JOIN (SELECT op_id, SUM(amount) AS pd, "
            "MAX(transfer_id) AS ltid FROM transfer_op_allocs "
            "GROUP BY op_id) a ON a.op_id=o.id "
            "WHERE o.status='other' AND COALESCE(o.distributed,0)=1 "
            "AND o.op_date BETWEEN ? AND ? "
            "ORDER BY o.op_date, o.id", (d1, d2)).fetchall()
        _trs = {t["id"]: {"tdate": t["tdate"], "method": t["method"],
                           "ref": t["ref"],
                           "voucher_number": t["voucher_number"]}
                for t in conn.execute(
            "SELECT id, tdate, COALESCE(method,'') AS method, "
            "COALESCE(ref,'') AS ref, "
            "COALESCE(voucher_number,'') AS voucher_number FROM "
            "other_company_transfers").fetchall()}
    out = []
    for r in rows:
        _g = round(float(r["amount"] or 0), 3)
        _fee = round(fees.get(r["gateway"], 0), 3)
        _net = round(_g - _fee, 3)
        _pd = round(float(r["pd"] or 0), 3)
        _legacy = int(r["pt"] or 0) > 0
        if (_pd > 0.0005 and _net - _pd <= 0.0005) or _legacy:
            _st = "paid"
            if _legacy and _pd <= 0.0005:
                _pd = _net
        elif _pd > 0.0005:
            _st = "partial"
        else:
            _st = "due"
        _tid = int(r["ltid"] or 0) or int(r["pt"] or 0)
        _t = _trs.get(_tid, {})
        out.append({"id": r["id"], "date": r["op_date"] or "",
                    "customer": r["customer"] or "",
                    "phone": r["phone"] or "",
                    "gateway": r["gateway"] or "",
                    "ref": r["ref"] or "", "gross": _g, "fee": _fee,
                    "net": _net, "paid": _pd,
                    "remaining": round(_net - _pd, 3), "status": _st,
                    "transfer_id": _tid, "tdate": _t.get("tdate") or "",
                    "method": _t.get("method") or "",
                    "transfer_ref": _t.get("ref") or "",
                    "effective_date": (r["effective_date"] or "").strip(),
                    "voucher_number": _t.get("voucher_number") or ""})
    return out


def op_transfer_allocs(op_id: int) -> list[dict]:
    """V4.40: تسديدات عملية واحدة (الأقدم أولًا).

    [{transfer_id, tdate, method, alloc, ref, voucher_number}].
    """
    ensure_alloc_table()
    ensure_transfer_table()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT a.transfer_id, t.tdate, COALESCE(t.method,'') AS method,"
            " a.amount AS alloc, COALESCE(t.ref,'') AS ref, "
            "COALESCE(t.voucher_number,'') AS voucher_number "
            "FROM transfer_op_allocs a "
            "JOIN other_company_transfers t ON t.id=a.transfer_id "
            "WHERE a.op_id=? ORDER BY a.transfer_id",
            (op_id,)).fetchall()
    return [{"transfer_id": r["transfer_id"], "tdate": r["tdate"] or "",
             "method": r["method"] or "", "alloc": round(r["alloc"], 3),
             "ref": r["ref"] or "",
             "voucher_number": r["voucher_number"] or ""} for r in rows]


def expected_collection_date(op_date: str) -> str:
    """V4.41: تاريخ تحصيل البوابة المتوقع من تاريخ العملية.

    القاعدة: ثاني يوم · عمليات الخميس والجمعة والسبت تُحصَّل
    الأحد التالي (عطلة البنوك) · تُحسب لحظيًا بلا تخزين.
    """
    d = date.fromisoformat((op_date or "")[:10])
    if d.weekday() in (3, 4, 5):  # خميس/جمعة/سبت ← الأحد
        delta = 6 - d.weekday()
    else:
        delta = 1
    return (d + timedelta(days=delta)).isoformat()


def collection_status(op_date: str, today: str | None = None,
                      stored_paid: int | None = None,
                      stored_date: str = "") -> tuple[bool, str]:
    """V4.41: (محصَّلة؟، تاريخ العرض) — المخزَّن أولًا ثم المحسوب."""
    exp = expected_collection_date(op_date)
    if stored_paid == 1:
        return True, (stored_date or "")[:10] or exp
    _today = (today or "")[:10] or date.today().isoformat()
    if exp <= _today:
        return True, exp
    return False, exp


def other_daily_settlement(date_from: str, date_to: str) -> dict:
    """كشف التسوية اليومي للأخرى (V4.37) — تجميع بتاريخ التسوية.

    {"days": [{day, ops, n, gross, fees, net, paid, remaining,
    status}], "totals": {...}, "transfers": [...],
    "transfers_total": float} — العمليات: الأخرى الموزَّعة المستحقة
    (ذات تعريف المستحق: other + موزَّعة + تحصيل غير مستبعد) ·
    net = الكامل − عمولة الملف · paid = التوزيعات التراكمية ·
    remaining = net − paid · التحويلات بتواريخها في الفترة.
    """
    ensure_ops_table()
    ensure_alloc_table()
    fees = gateway_fee_map()
    d1 = date_from or "2000-01-01"
    d2 = date_to or date.today().isoformat()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT o.id, o.gateway, o.op_date, "
            "COALESCE(o.customer,'') AS customer, "
            "COALESCE(o.phone,'') AS phone, "
            "COALESCE(o.ref,'') AS ref, o.amount, "
            "COALESCE(a.pd,0) AS pd FROM gateway_ops o "
            "LEFT JOIN (SELECT op_id, SUM(amount) AS pd FROM "
            "transfer_op_allocs GROUP BY op_id) a ON a.op_id=o.id "
            "WHERE " + _DUES_ELIGIBLE + " "
            "AND o.op_date >= ? AND o.op_date <= ? "
            "ORDER BY o.op_date, o.id", (d1, d2)).fetchall()
    days: list[dict] = []
    _by: dict[str, dict] = {}
    for r in rows:
        _day = r["op_date"] or ""
        _g = round(float(r["amount"] or 0), 3)
        _fee = round(fees.get(r["gateway"], 0), 3)
        _net = round(_g - _fee, 3)
        _pd = round(float(r["pd"] or 0), 3)
        _rem = round(_net - _pd, 3)
        _st = ("paid" if _rem <= 0.0005
                else ("partial" if _pd > 0.0005 else "unpaid"))
        _op = {"id": r["id"], "customer": r["customer"] or "",
               "phone": r["phone"] or "", "ref": r["ref"] or "",
               "gateway": r["gateway"] or "", "gross": _g,
               "fee": _fee, "net": _net, "paid": _pd,
               "remaining": _rem, "pay_status": _st}
        _d = _by.get(_day)
        if _d is None:
            _d = {"day": _day, "ops": [], "n": 0, "gross": 0.0,
                  "fees": 0.0, "net": 0.0, "paid": 0.0,
                  "remaining": 0.0, "status": "unpaid"}
            _by[_day] = _d
            days.append(_d)
        _d["ops"].append(_op)
    for _d in days:
        _d["n"] = len(_d["ops"])
        _d["gross"] = round(sum(o["gross"] for o in _d["ops"]), 3)
        _d["fees"] = round(sum(o["fee"] for o in _d["ops"]), 3)
        _d["net"] = round(sum(o["net"] for o in _d["ops"]), 3)
        _d["paid"] = round(sum(o["paid"] for o in _d["ops"]), 3)
        _d["remaining"] = round(_d["net"] - _d["paid"], 3)
        _d["status"] = ("paid" if _d["remaining"] <= 0.0005
                        else ("partial" if _d["paid"] > 0.0005
                              else "unpaid"))
    _t = {"n": sum(d["n"] for d in days),
          "gross": round(sum(d["gross"] for d in days), 3),
          "fees": round(sum(d["fees"] for d in days), 3),
          "net": round(sum(d["net"] for d in days), 3),
          "paid": round(sum(d["paid"] for d in days), 3)}
    _t["remaining"] = round(_t["net"] - _t["paid"], 3)
    _trs = list_other_transfers(d1, d2)
    return {"days": days, "totals": _t, "transfers": _trs,
            "transfers_total": round(
                sum(float(t["amount"] or 0) for t in _trs), 3)}


def op_payments(op_id: int) -> list[dict]:
    """تسديدات غطت العملية (الربط العكسي) — الأقدم أولًا."""
    ensure_alloc_table()
    ensure_transfer_table()
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT a.transfer_id, a.amount AS alloc, t.tdate, "
            "COALESCE(t.ref,'') AS ref, "
            "COALESCE(t.voucher_number,'') AS voucher_number "
            "FROM transfer_op_allocs a "
            "JOIN other_company_transfers t ON t.id=a.transfer_id "
            "WHERE a.op_id=? ORDER BY t.tdate, a.transfer_id",
            (int(op_id),)).fetchall()
    return [{"transfer_id": r["transfer_id"],
             "alloc": round(float(r["alloc"] or 0), 3),
             "tdate": r["tdate"] or "", "ref": r["ref"] or "",
             "voucher_number": r["voucher_number"] or ""}
            for r in rows]


def batch_detail(gateway: str, batch: str, day: str,
                 header: dict | None = None) -> dict:
    """تفاصيل رفعة (V4.34): header (نفس صف التقرير) + ops + totals.

    كل عملية: المستحق (صافي أو صفر) · المسدد (توزيعاتها) · المتبقي
    الدقيق · الحالة (paid بذات قاعدة الاستبعاد من المستحق:
    الإجمالي − المسدد ≤ 0.0005) · تسديداتها (op_payments مضمّنة).
    مفتاح مفقود ← أصفار بلا استثناء.
    header: صف جاهز من upload_batches_report (V4.35: يوفّر إعادة
    حساب التقرير كاملًا عند النداء المتكرر — التصدير/النقر).
    """
    ensure_ops_table()
    ensure_alloc_table()
    if header is not None and (
            header.get("gateway") != (gateway or "")
            or header.get("batch") != (batch or "")
            or header.get("day") != (day or "")):
        header = None
    if header is None:
        header = next(
            (b for b in upload_batches_report()
             if b["gateway"] == (gateway or "")
             and b["batch"] == (batch or "")
             and b["day"] == (day or "")),
            None)
    if header is None:
        _z = {"gateway": gateway or "", "batch": batch or "",
              "day": day or ""}
        return {"header": _z, "ops": [],
                "totals": {"dues": 0.0, "paid": 0.0, "remaining": 0.0}}
    fee = gateway_fee_map().get(gateway or "", 0)
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT o.id, o.op_date, COALESCE(o.customer,'') AS customer, "
            "COALESCE(o.phone,'') AS phone, o.amount, "
            "COALESCE(o.ref,'') AS ref, COALESCE(o.status,'') AS status, "
            "COALESCE(o.distributed,0) AS dist, o.gw_paid, "
            "COALESCE(a.pd,0) AS pd FROM gateway_ops o "
            "LEFT JOIN (SELECT op_id, SUM(amount) AS pd FROM "
            "transfer_op_allocs GROUP BY op_id) a ON a.op_id=o.id "
            "WHERE o.gateway=? AND COALESCE(o.batch,'')=? "
            "AND SUBSTR(COALESCE(o.created_at,''),1,10)=? "
            "ORDER BY o.op_date, o.id",
            (gateway or "", batch or "", day or "")).fetchall()
        pay_rows = conn.execute(
            "SELECT a.op_id, a.transfer_id, a.amount AS alloc, t.tdate, "
            "COALESCE(t.ref,'') AS ref, "
            "COALESCE(t.voucher_number,'') AS voucher_number "
            "FROM transfer_op_allocs a "
            "JOIN other_company_transfers t ON t.id=a.transfer_id "
            "WHERE a.op_id IN (SELECT o.id FROM gateway_ops o "
            "WHERE o.gateway=? AND COALESCE(o.batch,'')=? "
            "AND SUBSTR(COALESCE(o.created_at,''),1,10)=?) "
            "ORDER BY t.tdate, a.transfer_id",
            (gateway or "", batch or "", day or "")).fetchall()
    _pay: dict = {}
    for r in pay_rows:
        _pay.setdefault(r["op_id"], []).append({
            "transfer_id": r["transfer_id"],
            "alloc": round(float(r["alloc"] or 0), 3),
            "tdate": r["tdate"] or "", "ref": r["ref"] or "",
            "voucher_number": r["voucher_number"] or ""})
    ops = []
    for r in rows:
        _gross = round(float(r["amount"] or 0), 3)
        _pd = round(float(r["pd"] or 0), 3)
        _elig = (r["status"] == "other" and r["dist"]
                 and r["gw_paid"] in (None, 1))
        _dues = round(_gross - fee, 3) if _elig else 0.0
        if _gross - _pd <= 0.0005 and _pd > 0:
            _ps = "paid"
        elif _pd > 0.0005:
            _ps = "partial"
        elif _dues > 0:
            _ps = "unpaid"
        else:
            _ps = "na"
        ops.append({
            "id": r["id"], "date": r["op_date"] or "",
            "customer": r["customer"] or "", "phone": r["phone"] or "",
            "amount": _gross, "ref": r["ref"] or "",
            "status": r["status"] or "unmarked",
            "distributed": bool(r["dist"]), "gw_paid": r["gw_paid"],
            "dues": _dues, "paid": _pd,
            "remaining": round(_dues - _pd, 3), "pay_status": _ps,
            "payments": _pay.get(r["id"], [])})
    _td = round(sum(o["dues"] for o in ops), 3)
    _tp = round(sum(o["paid"] for o in ops), 3)
    return {"header": header, "ops": ops,
            "totals": {"dues": _td, "paid": _tp,
                       "remaining": round(_td - _tp, 3)}}


def batch_mine_detail(gateway: str, batch: str, day: str) -> dict:
    """عمليات شركتنا في رفعة + حالة يومها البنكية (نطاق شركتنا).

    {ops: [{id,date,customer,phone,amount,net,ref,gateway,
    day_status,bank_amount}], totals: {n,total,fees,net}} —
    الموزَّع-شركتنا كاملًا · net = amount − عمولة الملف (V4.36).
    مفتاح مفقود ← فارغ بلا استثناء.
    """
    ensure_ops_table()
    _fee = gateway_fee_map().get(gateway or "", 0)
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT id, op_date, COALESCE(customer,'') AS customer, "
            "COALESCE(phone,'') AS phone, amount, COALESCE(ref,'') AS ref "
            "FROM gateway_ops "
            "WHERE gateway=? AND COALESCE(batch,'')=? "
            "AND SUBSTR(COALESCE(created_at,''),1,10)=? "
            "AND status='mine' AND COALESCE(distributed,0)=1 "
            "ORDER BY op_date, id",
            (gateway or "", batch or "", day or "")).fetchall()
    ops = [{"id": r["id"], "date": r["op_date"] or "",
            "customer": r["customer"] or "", "phone": r["phone"] or "",
            "amount": round(float(r["amount"] or 0), 3),
            "net": round(float(r["amount"] or 0) - _fee, 3),
            "ref": r["ref"] or "", "gateway": gateway or "",
            "day_status": "—", "bank_amount": 0.0} for r in rows]
    if ops:
        _dates = sorted({o["date"] for o in ops if o["date"]})
        try:
            _cl = {c["day"]: c for c in daily_clearing(
                _dates[0], _dates[-1], gateway or "")} if _dates else {}
        except Exception:
            _cl = {}
        for o in ops:
            c = _cl.get(o["date"], {})
            o["day_status"] = STATUS_LABEL.get(c.get("status", ""), "—")
            o["bank_amount"] = round(float(c.get("bank_amount") or 0), 3)
    _gross = round(sum(o["amount"] for o in ops), 3)
    _fees = round(_fee * len(ops), 3)
    return {"ops": ops, "totals": {
        "n": len(ops), "total": _gross, "fees": _fees,
        "net": round(_gross - _fees, 3)}}


def unallocated_transfer_amount() -> float:
    """أجزاء التسديدات غير المربوطة بعمليات («دفعات على الحساب»).

    الهوية: Σpaid الرفعات + هذا المبلغ = other_transfers_total().
    """
    ensure_alloc_table()
    ensure_transfer_table()
    with get_connection() as conn:
        _al = conn.execute("SELECT COALESCE(SUM(amount),0) FROM "
                           "transfer_op_allocs").fetchone()[0]
    return round(other_transfers_total() - float(_al or 0), 3)


def other_ledger(date_from: str, date_to: str) -> dict:
    """كشف حساب الأخرى: افتتاحي + مستحق/مسدد برصيد جارٍ."""
    from datetime import date as _date
    op = get_other_opening()
    today = _date.today().isoformat()
    d1 = date_from or "2000-01-01"
    d2 = date_to or today
    try:
        day_before = (_date.fromisoformat(d1) - timedelta(days=1)).isoformat()
    except ValueError:
        day_before = ""
    prior = round(op["amount"]
                  + (_dues_cumulative(day_before) if day_before else 0.0)
                  - (other_transfers_total("", day_before)
                     if day_before else 0.0), 3)
    lines: list[dict] = []
    for r in other_statement(d1, d2):
        lines.append({
            "date": r["day"], "kind": "due",
            "text": (f"مستحق تسوية يوم {r['day']} · {r['gateway']}"
                     + (f" (بنك {r['bank_date']})" if r["bank_date"]
                        else "")),
            "gateway": r["gateway"], "day": r["day"],
            "plus": round(float(r["other_share"]), 3), "minus": 0.0,
            "tid": 0})
    for p in list_other_transfers(d1, d2):
        bits = ["تسديد"]
        if p["ref"]:
            bits.append(p["ref"])
        if p["method"]:
            bits.append(p["method"])
        if p["voucher_number"]:
            bits.append(f"سند {p['voucher_number']}")
        _ops = transfer_ops(int(p["id"])) if p.get("id") else []
        if _ops:
            _nm = [(o.get("customer") or o.get("phone") or o.get("ref")
                    or "—").strip()
                   + f" ({float(o.get('alloc') or 0):,.3f})"
                   for o in _ops]
            bits.append(" + ".join(_nm[:3])
                        + (f" +{len(_nm) - 3}" if len(_nm) > 3 else ""))
        lines.append({"date": p["tdate"], "kind": "pay",
                      "text": " · ".join(bits), "gateway": "",
                      "day": "", "plus": 0.0,
                      "minus": round(float(p["amount"]), 3),
                      "tid": int(p["id"])})
    lines.sort(key=lambda x: (x["date"],
                              0 if x["kind"] == "due" else 1, x["tid"]))
    run = prior
    for x in lines:
        run = round(run + x["plus"] - x["minus"], 3)
        x["balance"] = run
    return {"opening": op, "prior": prior, "lines": lines,
            "total_plus": round(sum(x["plus"] for x in lines), 3),
            "total_minus": round(sum(x["minus"] for x in lines), 3),
            "end": run}


def wipe_counts() -> dict:
    """عدادات المسح الشامل: عمليات/تسديدات/سندات (للعرض قبل التأكيد)."""
    ensure_ops_table()
    ensure_transfer_table()
    with get_connection() as conn:
        n_ops = conn.execute("SELECT COUNT(*) FROM gateway_ops").fetchone()[0]
        rows = conn.execute(
            "SELECT COALESCE(voucher_id,0) AS v FROM "
            "other_company_transfers").fetchall()
    vids = {int(r["v"]) for r in rows if r["v"]}
    return {"ops": int(n_ops), "transfers": len(rows),
            "vouchers": len(vids)}


def wipe_all_gateway_data() -> dict:
    """مسح شامل نهائي: كل عمليات كل البوابات + التسديدات + سنداتها وقيودها.

    تُحفظ نسخة احتياطية من القاعدة تلقائيًا قبل المسح. تُحفظ البيانات
    الرئيسية (ملفات البوابات/الافتتاحي/المستعارة/كشف البنك). يرجع العدادات
    + مسار النسخة.
    """
    ensure_ops_table()
    ensure_transfer_table()
    counts = wipe_counts()
    bdir = os.path.join(_db.DATA_DIR, "backups")
    os.makedirs(bdir, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = os.path.join(bdir, "wipe_%s.sqlite" % stamp)
    shutil.copy2(_db.DB_PATH, backup)
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT DISTINCT COALESCE(voucher_id,0) AS v FROM "
            "other_company_transfers WHERE COALESCE(voucher_id,0)>0"
        ).fetchall()
        vids = {int(r["v"]) for r in rows}
        if vids:
            kids = {int(r[0]) for r in conn.execute(
                "SELECT id FROM vouchers WHERE parent_voucher_id IN (%s)"
                % ",".join("?" * len(vids)), tuple(vids))}
            vids |= kids
            ph = ",".join("?" * len(vids))
            conn.execute(
                "DELETE FROM journal_entries WHERE source_id IN (%s) "
                "AND source_type LIKE 'voucher/_%%' ESCAPE '/'" % ph,
                tuple(vids))
            conn.execute("DELETE FROM vouchers WHERE id IN (%s)" % ph,
                         tuple(vids))
        conn.execute("DELETE FROM transfer_op_allocs")
        conn.execute("DELETE FROM other_company_transfers")
        conn.execute("DELETE FROM gateway_ops")
    counts["backup"] = backup
    return counts
