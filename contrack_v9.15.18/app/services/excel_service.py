# -*- coding: utf-8 -*-
"""تصدير بيانات العملاء مع الأرصدة والمدفوعات إلى ملف Excel احترافي."""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from database.database import get_connection
from services import settings_service


class ExcelError(Exception):
    """خطأ تصدير Excel برسالة عربية."""


_HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
_TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="1E3A5F")
_SUB_FONT = Font(name="Calibri", size=11, color="7B8A9A")
_TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
_TOTAL_FONT = Font(name="Calibri", bold=True, size=12, color="1E3A5F")
_GREEN_FONT = Font(name="Calibri", bold=True, color="1E7D4F")
_RED_FONT = Font(name="Calibri", bold=True, color="B03A2E")
_THIN = Side(style="thin", color="D9E2EC")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_RIGHT = Alignment(horizontal="right", vertical="center")


def _customers_financials() -> list[dict]:
    """بيانات كل عميل مع إجمالي عقوده ومدفوعاته (الفواتير) ورصيده المتبقي."""
    with get_connection() as conn:
        rows = conn.execute("""
            SELECT cu.customer_code, cu.name, cu.phone, cu.email, cu.address,
                   COUNT(DISTINCT c.id)                    AS contracts_count,
                   COALESCE(SUM(c.amount), 0)
                   + COALESCE((SELECT SUM(a.amount_add - a.amount_deduct)
                               FROM addendums a
                               JOIN contracts c2 ON c2.id = a.contract_id
                               WHERE c2.customer_id = cu.id
                                 AND a.status = 'معتمد'
                                 AND c2.status NOT IN ('ملغى')), 0)
                                                           AS total_contracts,
                   COALESCE((SELECT SUM(v.amount) FROM vouchers v
                             JOIN contracts c3 ON c3.id = v.contract_id
                             WHERE c3.customer_id = cu.id
                               AND v.vtype = 'receipt'
                               AND v.cancelled = 0), 0)    AS total_paid,
                   (SELECT COUNT(*) FROM vouchers v
                    JOIN contracts c4 ON c4.id = v.contract_id
                    WHERE c4.customer_id = cu.id
                      AND v.vtype = 'receipt'
                      AND v.cancelled = 0)                 AS invoices_count
            FROM customers cu
            LEFT JOIN contracts c ON c.customer_id = cu.id
                                  AND c.status NOT IN ('ملغى')
            GROUP BY cu.id
            ORDER BY cu.id
        """).fetchall()
    result = []
    for r in rows:
        total = float(r["total_contracts"])
        paid = float(r["total_paid"])
        result.append({
            "code": r["customer_code"], "name": r["name"],
            "phone": r["phone"] or "", "email": r["email"] or "",
            "address": r["address"] or "",
            "contracts_count": r["contracts_count"],
            "invoices_count": r["invoices_count"],
            "total_contracts": total,
            "total_paid": paid,
            "balance": total - paid,
        })
    return result


def export_customers_excel(out_path: str) -> str:
    """تصدير كشف العملاء بالأرصدة إلى Excel. يرجع المسار."""
    try:
        settings = settings_service.get_settings()
        data = _customers_financials()

        wb = Workbook()
        try:
            _company = (getattr(settings_service.get_settings(),
                                'company_name', '') or '').strip()
            wb.properties.creator = _company or settings_service.app_name()
            wb.properties.lastModifiedBy = wb.properties.creator
        except Exception:
            pass
        ws = wb.active
        ws.title = "كشف العملاء"
        ws.sheet_view.rightToLeft = True

        # ===== العنوان =====
        ws.merge_cells("A1:J1")
        ws["A1"] = settings.company_name or "كشف العملاء"
        ws["A1"].font = _TITLE_FONT
        ws["A1"].alignment = _CENTER
        ws.merge_cells("A2:J2")
        ws["A2"] = (f"كشف العملاء والأرصدة · تاريخ الإصدار: "
                    f"{datetime.now().strftime('%Y-%m-%d %H:%M')}")
        ws["A2"].font = _SUB_FONT
        ws["A2"].alignment = _CENTER
        ws.row_dimensions[1].height = 26
        ws.row_dimensions[2].height = 18

        # ===== رأس الجدول =====
        headers = ["رقم العميل", "اسم العميل", "الهاتف", "البريد الإلكتروني",
                   "العنوان", "عدد العقود", "عدد سندات القبض",
                   "إجمالي العقود", "المحصّل (سندات القبض)", "الرصيد المتبقي"]
        header_row = 4
        for col, head in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=head)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = _CENTER
            cell.border = _BORDER
        ws.row_dimensions[header_row].height = 24

        # ===== الصفوف =====
        num_fmt = "#,##0.000"
        row = header_row + 1
        for c in data:
            values = [c["code"], c["name"], c["phone"], c["email"], c["address"],
                      c["contracts_count"], c["invoices_count"],
                      c["total_contracts"], c["total_paid"], c["balance"]]
            for col, val in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.border = _BORDER
                cell.alignment = _CENTER if col != 2 else _RIGHT
                if col in (8, 9, 10):
                    cell.number_format = num_fmt
                if col == 9 and c["total_paid"] > 0:
                    cell.font = _GREEN_FONT
                if col == 10:
                    cell.font = _RED_FONT if c["balance"] > 0 else _GREEN_FONT
            row += 1

        # ===== صف الإجماليات =====
        total_contracts = sum(c["total_contracts"] for c in data)
        total_paid = sum(c["total_paid"] for c in data)
        total_balance = sum(c["balance"] for c in data)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        totals_label = ws.cell(row=row, column=1, value="الإجمالي")
        totals_label.fill = _TOTAL_FILL
        totals_label.font = _TOTAL_FONT
        totals_label.alignment = _CENTER
        for col, val in ((8, total_contracts), (9, total_paid), (10, total_balance)):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = _TOTAL_FILL
            cell.font = _TOTAL_FONT
            cell.alignment = _CENTER
            cell.number_format = num_fmt
            cell.border = _BORDER
        for col in range(1, 8):
            ws.cell(row=row, column=col).border = _BORDER
        ws.row_dimensions[row].height = 22

        # ===== عرض الأعمدة =====
        widths = [12, 26, 14, 24, 26, 11, 12, 16, 16, 16]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # تجميد الرأس + فلتر تلقائي
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = (f"A{header_row}:"
                              f"{get_column_letter(len(headers))}{row - 1}")

        wb.save(out_path)
        return out_path
    except PermissionError as exc:
        raise ExcelError("تعذر الحفظ · الملف مفتوح في برنامج آخر. "
                         "أغلقه وحاول مرة أخرى.") from exc
    except Exception as exc:
        raise ExcelError(f"تعذر إنشاء ملف Excel: {exc}") from exc
