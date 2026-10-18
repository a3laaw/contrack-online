# -*- coding: utf-8 -*-
"""استيراد البيانات القديمة من Excel.

قالب واحد بعدة أوراق: العملاء / العقود / مقاولو الباطن / سندات القبض والصرف.
الاستيراد بترتيب التبعية، مع فحص مسبق (Dry-Run) يعرض الأخطاء قبل الإدخال.
"""
import os
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from database.database import get_connection
from services.customer_service import ValidationError

# ============================================================
# تعريف الأوراق والأعمدة
# ============================================================
SHEETS = {
    "العملاء": {
        "columns": [
            ("اسم العميل *", 28), ("رقم الهاتف *", 16), ("العنوان", 26),
            ("البريد الإلكتروني", 22), ("ملاحظات", 24),
        ],
        "example": ["شركة الخليج التجارية", "99887766",
                    "الكويت - حولي", "info@gulf.com", "عميل قديم منذ 2020"],
    },
    "العقود": {
        "columns": [
            ("اسم العميل *", 28), ("اسم العقد / وصفه", 30),
            ("قيمة العقد *", 16), ("تاريخ العقد *", 16),
            ("تاريخ البداية", 16), ("تاريخ النهاية", 16), ("الحالة", 14),
            ("نوع العقد *", 24),
        ],
        "example": ["شركة الخليج التجارية", "عقد صيانة مبنى الإدارة",
                    "25000.000", "2025-03-01", "2025-03-01", "2026-02-28",
                    "نشط", "تصميم معماري"],
    },
    "مقاولو الباطن": {
        "columns": [
            ("اسم المقاول *", 28), ("رقم الهاتف", 16),
            ("نوع النشاط", 20), ("ملاحظات", 26),
        ],
        "example": ["مؤسسة النور للكهرباء", "66554433", "أعمال كهرباء", ""],
    },
    "الدفعات": {
        "columns": [
            ("اسم العميل *", 28), ("اسم العقد / وصفه", 30),
            ("المبلغ *", 16), ("التاريخ *", 16),
            ("البيان", 26), ("طريقة الدفع", 14),
        ],
        "example": ["شركة الخليج التجارية", "عقد صيانة مبنى الإدارة",
                    "5000.000", "2025-04-15", "الدفعة الأولى", "شيك"],
    },
    "سندات القبض": {
        "columns": [
            ("رقم السند القديم", 18), ("اسم العميل *", 28),
            ("اسم العقد / وصفه", 30), ("المبلغ *", 16),
            ("التاريخ *", 16), ("البيان", 28),
            ("طريقة الدفع", 14), ("ملاحظات", 26),
        ],
        "example": ["RV-OLD-001", "شركة الخليج التجارية",
                    "عقد صيانة مبنى الإدارة", "5000.000", "2025-04-15",
                    "الدفعة الأولى", "شيك", "سند قديم"],
    },
    "عقود الإشراف": {
        "columns": [
            ("اسم العميل *", 28), ("اسم العقد / وصفه", 30),
            ("تاريخ تعهد الإشراف *", 20), ("المدة بالأشهر *", 16),
            ("قيمة الشهر *", 16), ("تاريخ العقد", 16), ("الحالة", 14),
        ],
        "example": ["شركة الخليج التجارية", "إشراف فيلا العارضية",
                    "2025-05-01", "4", "200.000", "2025-04-20", "نشط"],
    },
    "عقود الباطن": {
        "columns": [
            ("اسم المقاول *", 28), ("اسم العميل *", 28),
            ("اسم العقد الرئيسي *", 30), ("قيمة عقد الباطن *", 18),
            ("تاريخ البداية", 16), ("تاريخ النهاية", 16), ("الوصف", 24),
        ],
        "example": ["مؤسسة النور للكهرباء", "شركة الخليج التجارية",
                    "عقد صيانة مبنى الإدارة", "3000.000",
                    "2025-03-15", "2025-09-15", "أعمال الكهرباء"],
    },
    "سندات الصرف": {
        "columns": [
            ("رقم السند القديم", 18), ("حساب المصروف *", 34),
            ("اسم المستفيد (مقاول/مورد) *", 30), ("اسم العميل", 28),
            ("اسم المشروع", 28), ("اسم العقد الرئيسي", 30),
            ("المبلغ *", 16), ("التاريخ *", 16),
            ("البيان", 26), ("طريقة الدفع", 14), ("ملاحظات", 26),
        ],
        "example": ["PV-OLD-001", "مصروفات تشغيلية / بنزين / وقود",
                    "محطة الوقود", "شركة الخليج التجارية", "",
                    "عقد صيانة مبنى الإدارة", "15.000", "2025-05-10",
                    "بنزين زيارة موقع", "تحويل بنكي", "منصرف قديم"],
    },
}

CONTRACT_STATUSES = ["نشط", "معتمد", "متوقف", "منتهي", "مغلق"]
PAY_METHODS = ["نقدًا", "KNET", "شيك", "تحويل بنكي", "أخرى"]


def _name_key(text: str) -> str:
    """مفتاح مطابقة عربي مرن للاستيراد: مسافات/همزات/ياء/ألف مقصورة."""
    s = str(text or "").strip()
    s = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", s)
    s = re.sub(r"[\u064b-\u065f\u0670]", "", s)
    trans = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ى": "ي", "ة": "ه"})
    s = s.translate(trans)
    s = re.sub(r"\s+", " ", s)
    return s.casefold()


def _existing_customer_name_map() -> dict[str, str]:
    try:
        with get_connection() as conn:
            return {_name_key(r["name"]): r["name"] for r in conn.execute(
                "SELECT name FROM customers").fetchall() if (r["name"] or "").strip()}
    except Exception:
        return {}


def _existing_contract_map() -> dict[tuple[str, str], int]:
    try:
        with get_connection() as conn:
            rows = conn.execute("""
                SELECT c.id, cu.name AS customer_name, c.contract_name,
                       c.description, c.contract_number
                FROM contracts c JOIN customers cu ON cu.id = c.customer_id
                WHERE COALESCE(c.status,'') <> 'ملغى'
            """).fetchall()
        result = {}
        for r in rows:
            for label in (r["contract_name"], r["description"], r["contract_number"]):
                if (r["customer_name"] or "").strip() and (label or "").strip():
                    result[(_name_key(r["customer_name"]), _name_key(label))] = int(r["id"])
        return result
    except Exception:
        return {}


def _con_label(contract_key) -> str:
    """اسم العقد المقروء من مفتاح (اسم العميل، اسم العقد)."""
    try:
        if isinstance(contract_key, (tuple, list)) and len(contract_key) > 1:
            return str(contract_key[1] or "")
    except Exception:
        pass
    return ""


def _receipt_overpay(lc, contract_id, amount: float) -> float:
    """كم يتجاوز سند القبض المتبقيَ على العقد؟ (0 = لا تجاوز).

    للتقرير فقط — لا يمنع الإدخال. يستخدم نفس حساب النظام:
    القيمة الحالية للعقد (مع الملحقات وتمديد الإشراف) ناقص المحصّل.
    """
    if not contract_id:
        return 0.0
    try:
        value = lc.contract_current_value(contract_id)
        due_value = value["current"]
        try:
            from services.supervision_service import supervision_status
            sup = supervision_status(contract_id)
            if sup:
                due_value = max(due_value, sup["total_due"])
        except Exception:
            pass
        with get_connection() as conn:
            collected = float(conn.execute(
                "SELECT COALESCE(SUM(amount),0) FROM vouchers "
                "WHERE contract_id=? AND vtype='receipt' AND cancelled=0",
                (contract_id,)).fetchone()[0])
        over = float(amount) - (due_value - collected)
        return over if over > 1e-9 else 0.0
    except Exception:
        return 0.0


# ============================================================
# المرحلة 6 — منع دخول العقود صامتةً تحت «غير محدد»
# ============================================================
def _type_key(name: str) -> str:
    """تطبيع اسم نوع العقد للمقارنة (يتحمّل فروق الهمزة والتاء)."""
    from services.search_utils import normalize
    return normalize(name)


def _known_contract_types() -> dict:
    """أنواع العقود المعتمدة: مطبّع ← الاسم الرسمي (بلا «غير محدد»)."""
    out = {}
    try:
        from services import contract_type_service as cts
        for t in cts.list_types(False):
            name = str(t.get("name") or "").strip()
            if name and name != cts.UNSPECIFIED:
                out[_type_key(name)] = name
    except Exception:
        pass
    return out


def _sheet_has_type_column(wb, sheet_name: str) -> bool:
    """هل عمود «نوع العقد» موجود في موضعه الصحيح بورقة العقود؟

    نتحقق من خلية العنوان رقم 8 تحديدًا (لا من أي خلية في الصف) حتى
    لا تُخدع الدالة بملاحظة جانبية مكتوبة بعد الأعمدة.
    """
    if sheet_name not in wb.sheetnames:
        return False
    ws = wb[sheet_name]
    head = ws.cell(row=1, column=8).value
    return bool(head and "نوع العقد" in str(head))


# ============================================================
# 1) توليد القالب
# ============================================================
def generate_template(out_path: str) -> str:
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    example_font = Font(italic=True, color="9AA7B5", size=11)
    note_font = Font(bold=True, color="B45309", size=10)
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    # ورقة التعليمات أولاً
    ws = wb.active
    ws.title = "اقرأني أولاً"
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 95
    instructions = [
        ("تعليمات تعبئة قالب البيانات القديمة", 16, True, "1E3A5F"),
        ("", 11, False, None),
        ("١. لا تغيّر أسماء الأوراق أو ترتيب الأعمدة أو تحذف صف العناوين.", 12, False, None),
        ("٢. الأعمدة التي تنتهي بنجمة (*) إجبارية · الباقي اختياري.", 12, False, None),
        ("٣. الصف الرمادي في كل ورقة مجرد مثال توضيحي، سيتم تجاهله عند الاستيراد.", 12, False, None),
        ("٤. التواريخ بصيغة: سنة-شهر-يوم مثل 2025-03-01 (أو أي صيغة تاريخ في إكسل).", 12, False, None),
        ("٥. المبالغ أرقام فقط بدون «د.ك» · مثال: 25000.000", 12, False, None),
        ("٦. رقم الهاتف في ورقة «العملاء» مطلوب لأنه رقم تواصل/تعريف أساسي للعميل.", 12, True, "B45309"),
        ("٧. اسم العميل في ورقتي «العقود» و«سندات القبض» يجب أن يطابق تمامًا اسمه في ورقة «العملاء».", 12, True, "B45309"),
        ("٨. «اسم العقد / وصفه» في ورقة «سندات القبض» يجب أن يطابق نفس الاسم في ورقة «العقود».", 12, True, "B45309"),
        ("٩. لو لعميل أكثر من عقد: كرر اسم العميل في صفوف متعددة بورقة العقود.", 12, False, None),
        ("١٠. أرقام العملاء والعقود والسندات سيولّدها النظام تلقائيًا · لا تكتبوها.", 12, False, None),
        ("١١. عقود الإشراف الهندسي لها ورقة خاصة «عقود الإشراف» · تبدأ من تاريخ تعهد الإشراف (بلدية الكويت) وقيمتها = المدة × قيمة الشهر، ولا تُكرر في ورقة «العقود».", 12, True, "7C3AED"),
        ("١٢. دفعات/وارد عقد الإشراف تُسجل في ورقة «سندات القبض» بنفس اسم العقد.", 12, False, None),
        ("١٣. يمكن استخدام ورقة «الدفعات» القديمة أو ورقة «سندات القبض» الجديدة؛ الجديدة تسمح برقم السند القديم وملاحظاته.", 12, True, "1E3A5F"),
        ("١٤. ورقة «عقود الباطن»: اربط المقاول بالعقد الرئيسي وقيمة اتفاقه قبل تسجيل أي صرف له.", 12, False, None),
        ("١٥. ورقة «سندات الصرف»: للصرف لمقاول على عقد اكتب العميل واسم العقد · وللمصروف العام (مورد/كهرباء/رواتب) اترك خانتي العميل والعقد فارغتين.", 12, True, "B45309"),
        ("١٦. عمود «نوع العقد» في ورقة «العقود» إجباري · ويجب أن يكون نوعًا معتمدًا موجودًا في البرنامج مسبقًا.", 12, True, "B45309"),
        ("١٧. لو النوع غير موجود: أنشئه أولًا من زر «رفع أنواع العقود والمراحل والدفعات من Excel» ثم أعد الفحص.", 12, True, "7C3AED"),
        ("١٨. لن يُدخل أي عقد تحت «غير محدد» · أي نوع فارغ أو مجهول يوقف الاستيراد ويُعرض لك الحل بوضوح.", 12, True, "B45309"),
        ("١٩. بعد التعبئة احفظ الملف وسلّمه · وسيُفحص قبل الإدخال ويُعرض أي خطأ بوضوح.", 12, False, None),
    ]
    for i, (text, size, bold, color) in enumerate(instructions, 1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=size,
                         color=color or "243447")
        cell.alignment = Alignment(horizontal="right", vertical="center",
                                   wrap_text=True)
        ws.row_dimensions[i].height = 24 if size <= 12 else 32

    # أوراق البيانات
    for sheet_name, spec in SHEETS.items():
        ws = wb.create_sheet(sheet_name)
        ws.sheet_view.rightToLeft = True
        for col, (head, width) in enumerate(spec["columns"], 1):
            cell = ws.cell(row=1, column=col, value=head)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[1].height = 26
        # صف المثال
        for col, val in enumerate(spec["example"], 1):
            cell = ws.cell(row=2, column=col, value=val)
            cell.font = example_font
            cell.alignment = center
            cell.border = border
        ws.cell(row=2, column=len(spec["columns"]) + 1,
                value="← مثال، سيتم تجاهله").font = note_font
        # حدود لمئة صف جاهزة للتعبئة
        for r in range(3, 103):
            for col in range(1, len(spec["columns"]) + 1):
                ws.cell(row=r, column=col).border = border
                ws.cell(row=r, column=col).alignment = center
        ws.freeze_panes = "A2"

    # قوائم منسدلة للحالة وطريقة الدفع
    ws_contracts = wb["العقود"]
    dv1 = DataValidation(type="list",
                         formula1='"' + ",".join(CONTRACT_STATUSES) + '"',
                         allow_blank=True)
    ws_contracts.add_data_validation(dv1)
    dv1.add("G3:G200")
    # قائمة أنواع العقود المعتمدة (المرحلة 6) — تمنع كتابة نوع مجهول
    _types = sorted(_known_contract_types().values())
    if _types:
        _joined = ",".join(_types)
        # حد إكسل لقائمة مكتوبة مباشرة ~255 حرفًا؛ فوقها نكتفي بالتحقق
        if len(_joined) <= 250:
            dv_type = DataValidation(type="list",
                                     formula1='"' + _joined + '"',
                                     allow_blank=False)
            ws_contracts.add_data_validation(dv_type)
            dv_type.add("H3:H200")
    ws_contracts.cell(row=1, column=9, value="← نوع العقد إجباري ومن "
                      "الأنواع المعتمدة فقط").font = note_font
    for _sheet, _rng in (("الدفعات", "F3:F200"), ("سندات القبض", "G3:G200")):
        if _sheet in wb.sheetnames:
            ws_pay = wb[_sheet]
            dv2 = DataValidation(type="list",
                                 formula1='"' + ",".join(PAY_METHODS) + '"',
                                 allow_blank=True)
            ws_pay.add_data_validation(dv2)
            dv2.add(_rng)

    wb.save(out_path)
    return out_path


def generate_expenses_template(out_path: str) -> str:
    """قالب مستقل لرفع سندات الصرف فقط."""
    wb = Workbook()
    ws = wb.active
    ws.title = "اقرأني أولاً"
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 100
    notes = [
        "قالب رفع سندات الصرف فقط.",
        "املأ ورقة «سندات الصرف» فقط، ويمكن ترك العميل/المشروع/العقد فارغين للمصروف العام.",
        "حساب المصروف يكتب مثل: مصروفات تشغيلية / بنزين / وقود أو اسم الحساب الفرعي مثل: بنزين / وقود.",
        "رقم السند القديم اختياري، وإذا كان موجودًا مسبقًا سيتم تخطيه.",
    ]
    for i, txt in enumerate(notes, 1):
        ws.cell(i, 1, txt).font = Font(bold=(i == 1), size=14 if i == 1 else 12, color="1E3A5F")
        ws.cell(i, 1).alignment = Alignment(horizontal="right", wrap_text=True)
    spec = SHEETS["سندات الصرف"]
    ws = wb.create_sheet("سندات الصرف")
    ws.sheet_view.rightToLeft = True
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    example_font = Font(italic=True, color="64748B")
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, (head, width) in enumerate(spec["columns"], 1):
        cell = ws.cell(row=1, column=col, value=head)
        cell.fill = header_fill; cell.font = header_font; cell.alignment = center; cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width
    for col, val in enumerate(spec["example"], 1):
        cell = ws.cell(row=2, column=col, value=val)
        cell.font = example_font; cell.alignment = center; cell.border = border
    ws.cell(row=2, column=len(spec["columns"]) + 1, value="← مثال، سيتم تجاهله")
    for r in range(3, 303):
        for col in range(1, len(spec["columns"]) + 1):
            ws.cell(row=r, column=col).border = border
            ws.cell(row=r, column=col).alignment = center
    dv = DataValidation(type="list", formula1='"' + ",".join(PAY_METHODS) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add("J3:J300")
    wb.save(out_path)
    return out_path


# ============================================================
# 2) قراءة الملف وتحليله
# ============================================================
def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _cell_date(v) -> str:
    """تطبيع التاريخ لصيغة YYYY-MM-DD · يقبل datetime أو نص."""
    if v is None or v == "":
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    text = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text  # سيُرفض في الفحص إن كان غير صالح


def _cell_amount(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        cleaned = re.sub(r"[^\d.\-]", "", str(v))
        return float(cleaned)
    except ValueError:
        return None


def _is_valid_date(s: str) -> bool:
    try:
        datetime.strptime(s, "%Y-%m-%d")
        return True
    except (ValueError, TypeError):
        return False


def parse_excel(path: str) -> dict:
    """قراءة الملف وإرجاع البيانات الخام + قائمة الأخطاء (فحص مسبق)."""
    if not os.path.exists(path):
        raise ValidationError("الملف غير موجود.")
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as exc:
        raise ValidationError("تعذر فتح الملف، تأكد أنه ملف Excel سليم "
                              "(.xlsx).") from exc

    errors: list[str] = []
    warnings: list[str] = []
    data = {"customers": [], "contracts": [], "subs": [], "payments": [],
            "sup_contracts": []}

    def sheet_rows(name, n_cols):
        if name not in wb.sheetnames:
            warnings.append(f"ورقة «{name}» غير موجودة، سيتم تخطيها.")
            return []
        ws = wb[name]
        rows = []
        example = SHEETS[name]["example"]
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True),
                                start=2):
            vals = [(_cell_str(row[c]) if c < len(row) else "")
                    for c in range(n_cols)]
            marker = _cell_str(row[n_cols]) if len(row) > n_cols else ""
            raw = list(row[:n_cols]) + [None] * (n_cols - len(row))
            # تجاهل صف المثال في القالب حتى لو مسح المستخدم أول خانة وبقيت
            # باقي بيانات المثال أو علامة «مثال، سيتم تجاهله».
            if i == 2:
                same_example = all(
                    (not vals[c]) or vals[c] == _cell_str(example[c])
                    for c in range(min(n_cols, len(example))))
                if "مثال" in marker or same_example:
                    continue
            if not any(vals):
                continue
            rows.append((i, vals, raw))
        return rows

    # ===== العملاء =====
    # يسمح الاستيراد الآن بأن تكون العقود/السندات مرتبطة بعميل موجود مسبقًا
    # في قاعدة البيانات حتى لو لم يكرر المستخدم اسمه في ورقة العملاء.
    existing_customer_aliases = _existing_customer_name_map()
    customer_aliases = dict(existing_customer_aliases)  # norm -> الاسم الرسمي
    seen_names = set()
    seen_name_keys = set()
    for line, vals, _ in sheet_rows("العملاء", 5):
        name, phone, address, email, notes = vals
        if not name:
            errors.append(f"العملاء سطر {line}: اسم العميل فارغ.")
            continue
        if not phone:
            errors.append(f"العملاء سطر {line}: رقم الهاتف مطلوب.")
            continue
        nkey = _name_key(name)
        if nkey in seen_name_keys:
            errors.append(f"العملاء سطر {line}: الاسم «{name}» مكرر في الملف.")
            continue
        seen_name_keys.add(nkey)
        seen_names.add(name)
        customer_aliases[nkey] = existing_customer_aliases.get(nkey, name)
        data["customers"].append({"name": name, "phone": phone,
                                  "address": address, "email": email,
                                  "notes": notes})

    # ===== العقود =====
    # المرحلة 6 — منع الإدخال الصامت تحت «غير محدد»:
    # نوع العقد عمود إجباري، ويجب أن يكون نوعًا معتمدًا موجودًا فعلًا.
    # أي نوع مجهول أو عمود ناقص = خطأ يوقف الاستيراد ويُعرض لك قرار
    # صريح (اختيار نوع موجود / إنشاء النوع / تعديل الملف).
    known_type_names = _known_contract_types()
    unknown_types: dict[str, list[int]] = {}
    contracts_sheet_has_type = _sheet_has_type_column(wb, "العقود")
    if not contracts_sheet_has_type and "العقود" in wb.sheetnames:
        errors.append(
            "ورقة «العقود»: عمود «نوع العقد» غير موجود في الملف · "
            "نزّل القالب الجديد وانقل بياناتك إليه، أو أضف العمود "
            "وحدّد نوع كل عقد · لن تُدخل العقود تحت «غير محدد».")

    contract_keys = set()
    for line, vals, raw in sheet_rows("العقود", 8):
        cname, title, amount_s, cdate, sdate, edate, status, ctype = vals
        amount = _cell_amount(raw[2])
        cdate_n, sdate_n, edate_n = (_cell_date(raw[3]), _cell_date(raw[4]),
                                     _cell_date(raw[5]))
        if not cname:
            errors.append(f"العقود سطر {line}: اسم العميل فارغ.")
            continue
        cname_canonical = customer_aliases.get(_name_key(cname))
        if not cname_canonical:
            errors.append(f"العقود سطر {line}: العميل «{cname}» غير موجود في ورقة العملاء أو قائمة العملاء الحالية.")
            continue
        cname = cname_canonical
        if amount is None or amount <= 0:
            errors.append(f"العقود سطر {line}: قيمة العقد غير صالحة "
                          f"({amount_s}).")
            continue
        if not _is_valid_date(cdate_n):
            errors.append(f"العقود سطر {line}: تاريخ العقد غير صالح "
                          f"({cdate}).")
            continue
        for label, d in (("البداية", sdate_n), ("النهاية", edate_n)):
            if d and not _is_valid_date(d):
                errors.append(f"العقود سطر {line}: تاريخ {label} غير صالح.")
        if status and status not in CONTRACT_STATUSES:
            warnings.append(f"العقود سطر {line}: الحالة «{status}» غير "
                            "معروفة · ستُسجل «نشط».")
            status = "نشط"
        key = (cname, title or f"عقد بقيمة {amount}")
        if key in contract_keys:
            errors.append(f"العقود سطر {line}: عقد مكرر لنفس العميل بنفس "
                          "الاسم · ميّز الاسمين.")
            continue
        contract_keys.add(key)
        # --- التحقق الصريح من نوع العقد (المرحلة 6) ---
        ctype = (ctype or "").strip()
        if contracts_sheet_has_type:
            if not ctype:
                errors.append(
                    f"العقود سطر {line}: «نوع العقد» فارغ · حدّد النوع "
                    "صراحةً · لن يُدخل العقد تحت «غير محدد».")
                continue
            real = known_type_names.get(_type_key(ctype))
            if not real:
                unknown_types.setdefault(ctype, []).append(line)
                continue
            ctype = real
        data["contracts"].append({
            "customer_name": cname, "title": title, "amount": amount,
            "contract_date": cdate_n, "start_date": sdate_n or "",
            "end_date": edate_n or "", "status": status or "نشط",
            "contract_type": ctype})

    # تقرير الأنواع المجهولة بصيغة قرار صريح (المرحلة 6)
    if unknown_types:
        available = ("، ".join(sorted(known_type_names.values()))
                     or "لا يوجد أي نوع معتمد بعد")
        for bad, lines in unknown_types.items():
            shown = "، ".join(str(x) for x in lines[:8])
            more = " وغيرها" if len(lines) > 8 else ""
            errors.append(
                f"العقود: نوع العقد «{bad}» غير معتمد (أسطر: {shown}{more}) · "
                f"اختر أحد الحلول: (١) صحّح الاسم لأحد الأنواع الموجودة: "
                f"{available} · أو (٢) أنشئ النوع أولًا من زر «رفع أنواع "
                f"العقود والمراحل والدفعات من Excel» · أو (٣) أوقف "
                f"الاستيراد وعدّل الملف · لن يُدخل العقد تحت «غير محدد».")

    # ===== مقاولو الباطن =====
    seen_subs = set()
    existing_sub_names = set()
    try:
        with get_connection() as conn:
            existing_sub_names.update(str(r[0] or "") for r in conn.execute("SELECT name FROM subcontractors").fetchall())
    except Exception:
        pass
    seen_subs.update(existing_sub_names)
    file_subs = set()
    for line, vals, _ in sheet_rows("مقاولو الباطن", 4):
        name, phone, activity, notes = vals
        if not name:
            errors.append(f"مقاولو الباطن سطر {line}: الاسم فارغ.")
            continue
        if name in file_subs:
            errors.append(f"مقاولو الباطن سطر {line}: «{name}» مكرر.")
            continue
        file_subs.add(name)
        seen_subs.add(name)
        data["subs"].append({"name": name, "phone": phone,
                             "activity": activity, "notes": notes})

    # ===== عقود الإشراف =====
    for line, vals, raw in sheet_rows("عقود الإشراف", 7):
        cname, title, sdate, months_s, monthly_s, cdate, status = vals
        sdate_n = _cell_date(raw[2])
        monthly = _cell_amount(raw[4])
        cdate_n = _cell_date(raw[5])
        try:
            months = int(float(str(raw[3]).strip() or 0))
        except (ValueError, TypeError):
            months = 0
        if not cname:
            errors.append(f"عقود الإشراف سطر {line}: اسم العميل فارغ.")
            continue
        cname_canonical = customer_aliases.get(_name_key(cname))
        if not cname_canonical:
            errors.append(f"عقود الإشراف سطر {line}: العميل «{cname}» غير موجود في ورقة العملاء أو قائمة العملاء الحالية.")
            continue
        cname = cname_canonical
        if not _is_valid_date(sdate_n):
            errors.append(f"عقود الإشراف سطر {line}: تاريخ تعهد الإشراف "
                          f"غير صالح ({sdate}).")
            continue
        if months <= 0:
            errors.append(f"عقود الإشراف سطر {line}: المدة بالأشهر غير "
                          f"صالحة ({months_s}).")
            continue
        if monthly is None or monthly <= 0:
            errors.append(f"عقود الإشراف سطر {line}: قيمة الشهر غير صالحة "
                          f"({monthly_s}).")
            continue
        if status and status not in CONTRACT_STATUSES:
            warnings.append(f"عقود الإشراف سطر {line}: الحالة «{status}» "
                            "غير معروفة · ستُسجل «نشط».")
            status = "نشط"
        key = (cname, title or f"إشراف {monthly}x{months}")
        if key in contract_keys:
            errors.append(f"عقود الإشراف سطر {line}: عقد بنفس الاسم موجود "
                          "لنفس العميل · ميّز الاسمين.")
            continue
        contract_keys.add(key)
        data["sup_contracts"].append({
            "customer_name": cname, "title": title,
            "sup_start": sdate_n, "months": months, "monthly": monthly,
            "amount": months * monthly,
            "contract_date": cdate_n or sdate_n,
            "status": status or "نشط"})

    # ===== سندات القبض / الدفعات القديمة =====
    contract_totals: dict[tuple, float] = {}
    for c in data["contracts"]:
        key = (c["customer_name"], c["title"] or f"عقد بقيمة {c['amount']}")
        contract_totals[key] = c["amount"]
    # عقود الإشراف قابلة للسداد أيضًا (بحد قيمتها الأساسية عند الاستيراد)
    for c in data["sup_contracts"]:
        key = (c["customer_name"], c["title"] or f"إشراف {c['monthly']}x{c['months']}")
        contract_totals[key] = c["amount"]
    # اسم العميل/العقد الموجود فعليًا في النظام مقبول أيضًا، حتى يمكن استيراد سندات/عقود باطن
    # على عقود قديمة دون تكرار العميل والعقد في ملف Excel.
    try:
        with get_connection() as conn:
            for r in conn.execute("""
                SELECT cu.name AS customer_name, c.contract_name, c.description,
                       c.contract_number, c.amount AS amount
                FROM contracts c JOIN customers cu ON cu.id = c.customer_id
                WHERE COALESCE(c.status,'') <> 'ملغى'
            """).fetchall():
                labels = [r["contract_name"], r["description"], r["contract_number"]]
                for label in labels:
                    key = (r["customer_name"], (label or "").strip())
                    if key[0] and key[1] and key not in contract_totals:
                        contract_totals[key] = float(r["amount"] or 0)
    except Exception:
        pass
    collected: dict[tuple, float] = {}
    seen_receipt_numbers: set[str] = set()
    try:
        with get_connection() as conn:
            existing_receipt_numbers = {str(r[0] or "") for r in conn.execute(
                "SELECT voucher_number FROM vouchers").fetchall()}
    except Exception:
        existing_receipt_numbers = set()

    def parse_receipt_row(sheet_name: str, line: int, vals: list, raw: list,
                          has_old_number: bool = False):
        if has_old_number:
            old_no, cname, ctitle, amount_s, pdate, descr, method, notes = vals
            amount = _cell_amount(raw[3])
            pdate_n = _cell_date(raw[4])
        else:
            old_no = ""
            cname, ctitle, amount_s, pdate, descr, method = vals
            notes = ""
            amount = _cell_amount(raw[2])
            pdate_n = _cell_date(raw[3])
        prefix = f"{sheet_name} سطر {line}"
        if old_no:
            if old_no in seen_receipt_numbers:
                errors.append(f"{prefix}: رقم السند القديم «{old_no}» مكرر داخل الملف.")
                return
            if old_no in existing_receipt_numbers:
                warnings.append(f"{prefix}: رقم السند القديم «{old_no}» موجود مسبقًا في النظام · تم تخطيه.")
                return
            seen_receipt_numbers.add(old_no)
        if not cname:
            errors.append(f"{prefix}: اسم العميل فارغ.")
            return
        cname_canonical = customer_aliases.get(_name_key(cname))
        if not cname_canonical:
            errors.append(f"{prefix}: العميل «{cname}» غير موجود في ورقة العملاء أو قائمة العملاء الحالية.")
            return
        cname = cname_canonical
        if amount is None or amount <= 0:
            errors.append(f"{prefix}: المبلغ غير صالح ({amount_s}).")
            return
        if not _is_valid_date(pdate_n):
            errors.append(f"{prefix}: التاريخ غير صالح ({pdate}).")
            return
        # مطابقة العقد
        matches = [k for k in contract_totals
                   if _name_key(k[0]) == _name_key(cname)
                   and (not ctitle or _name_key(k[1]) == _name_key(ctitle))]
        if ctitle and not matches:
            errors.append(f"{prefix}: العقد «{ctitle}» غير موجود للعميل «{cname}» في ورقة العقود.")
            return
        if not ctitle:
            client_contracts = [k for k in contract_totals if _name_key(k[0]) == _name_key(cname)]
            if len(client_contracts) > 1:
                errors.append(f"{prefix}: للعميل أكثر من عقد · حدد اسم العقد.")
                return
            if not client_contracts:
                errors.append(f"{prefix}: لا يوجد عقد للعميل «{cname}».")
                return
            matches = client_contracts
        key = matches[0]
        collected[key] = collected.get(key, 0.0) + amount
        # تجاوز قيمة العقد تنبيه لا خطأ: النظام يسجل الزيادة كدفعة مقدمة
        # على العقد (نفس سلوك الشاشات العادية)، فلا يتوقف الاستيراد ولا
        # يُحذف السند صامتًا. التفاصيل تظهر في تقرير ما بعد الرفع.
        if collected[key] > contract_totals[key] + 0.001:
            warnings.append(
                f"{prefix}: مجموع سندات قبض عقد «{key[1]}» "
                f"({collected[key]:,.3f}) يتجاوز قيمته "
                f"({contract_totals[key]:,.3f}) · ستُسجَّل الزيادة كدفعة "
                "مقدمة على العقد · راجع قيمة العقد أو تكرار السندات.")
        if method and method not in PAY_METHODS:
            warnings.append(f"{prefix}: طريقة الدفع «{method}» ستُسجل «أخرى».")
            method = "أخرى"
        data["payments"].append({
            "customer_name": cname, "contract_key": key, "amount": amount,
            "date": pdate_n, "description": descr or "سند قبض (مستورد)",
            "method": method or "أخرى", "notes": notes or "",
            "old_number": old_no or ""})

    # الورقة القديمة باقية للتوافق.
    for line, vals, raw in sheet_rows("الدفعات", 6):
        parse_receipt_row("الدفعات", line, vals, raw, has_old_number=False)
    # الورقة الجديدة أوضح لأنها سندات قبض فعلية ويمكن حفظ رقم السند القديم.
    if "سندات القبض" in wb.sheetnames:
        for line, vals, raw in sheet_rows("سندات القبض", 8):
            parse_receipt_row("سندات القبض", line, vals, raw, has_old_number=True)

    # ===== عقود الباطن =====
    data["sub_contracts"] = []
    sub_totals: dict[tuple, float] = {}   # (مقاول، عميل، عقد) -> قيمة
    for line, vals, raw in sheet_rows("عقود الباطن", 7):
        sname, cname, ctitle, amount_s, sdate, edate, descr = vals
        amount = _cell_amount(raw[3])
        sdate_n, edate_n = _cell_date(raw[4]), _cell_date(raw[5])
        if not sname:
            errors.append(f"عقود الباطن سطر {line}: اسم المقاول فارغ.")
            continue
        if sname not in seen_subs:
            errors.append(f"عقود الباطن سطر {line}: المقاول «{sname}» غير "
                          "موجود في ورقة مقاولي الباطن.")
            continue
        cname_canonical = customer_aliases.get(_name_key(cname)) if cname else None
        if not cname or not cname_canonical:
            errors.append(f"عقود الباطن سطر {line}: العميل «{cname}» غير موجود في ورقة العملاء أو قائمة العملاء الحالية.")
            continue
        cname = cname_canonical
        # مطابقة العقد الرئيسي (عادي أو إشراف)
        matches = [k for k in contract_totals
                   if _name_key(k[0]) == _name_key(cname)
                   and (not ctitle or _name_key(k[1]) == _name_key(ctitle))]
        if not matches:
            errors.append(f"عقود الباطن سطر {line}: العقد «{ctitle}» غير "
                          f"موجود للعميل «{cname}».")
            continue
        if amount is None or amount <= 0:
            errors.append(f"عقود الباطن سطر {line}: قيمة عقد الباطن غير "
                          f"صالحة ({amount_s}).")
            continue
        key = matches[0]
        sub_key = (sname, key)
        sub_totals[sub_key] = sub_totals.get(sub_key, 0.0) + amount
        data["sub_contracts"].append({
            "sub_name": sname, "contract_key": key, "amount": amount,
            "start_date": sdate_n or "", "end_date": edate_n or "",
            "description": descr or ""})

    # ===== سندات الصرف =====
    data["expenses"] = []
    sub_paid: dict[tuple, float] = {}
    try:
        from services import expense_account_service as _ea
        _expense_accounts = _ea.list_accounts(active_only=False, posting_only=False)
        expense_account_map = {}
        for a in _expense_accounts:
            for alias in (a.get("full_path"), a.get("name")):
                if alias:
                    expense_account_map[_name_key(alias)] = int(a["id"])
    except Exception:
        expense_account_map = {}
    try:
        with get_connection() as conn:
            existing_payment_numbers = {str(r[0] or "") for r in conn.execute(
                "SELECT voucher_number FROM vouchers WHERE vtype='payment'").fetchall()}
    except Exception:
        existing_payment_numbers = set()
    try:
        with get_connection() as conn:
            project_map = {_name_key(r["name"]): int(r["id"]) for r in conn.execute("SELECT id,name FROM projects").fetchall()}
    except Exception:
        project_map = {}
    seen_payment_numbers: set[str] = set()
    for line, vals, raw in sheet_rows("سندات الصرف", 11):
        old_no, acc_name, pname, cname, project_name, ctitle, amount_s, pdate, descr, method, notes = vals
        amount = _cell_amount(raw[6])
        pdate_n = _cell_date(raw[7])
        prefix = f"سندات الصرف سطر {line}"
        if old_no:
            if old_no in seen_payment_numbers:
                errors.append(f"{prefix}: رقم السند القديم «{old_no}» مكرر داخل الملف.")
                continue
            if old_no in existing_payment_numbers:
                warnings.append(f"{prefix}: رقم السند القديم «{old_no}» موجود مسبقًا في النظام · سيتم تخطيه.")
                data["expenses"].append({"_skip_existing": True})
                seen_payment_numbers.add(old_no)
                continue
            seen_payment_numbers.add(old_no)
        if not pname:
            errors.append(f"{prefix}: اسم المستفيد فارغ.")
            continue
        expense_account_id = expense_account_map.get(_name_key(acc_name))
        if not expense_account_id:
            errors.append(f"{prefix}: حساب المصروف «{acc_name}» غير موجود في إعدادات حسابات المصروفات.")
            continue
        if amount is None or amount <= 0:
            errors.append(f"سندات الصرف سطر {line}: المبلغ غير صالح "
                          f"({amount_s}).")
            continue
        if not _is_valid_date(pdate_n):
            errors.append(f"سندات الصرف سطر {line}: التاريخ غير صالح "
                          f"({pdate}).")
            continue
        if method and method not in PAY_METHODS:
            warnings.append(f"سندات الصرف سطر {line}: طريقة الدفع "
                            f"«{method}» ستُسجل «أخرى».")
            method = "أخرى"
        is_sub = pname in seen_subs
        key = None
        expense_customer_name = ""
        if cname:
            cname_canonical = customer_aliases.get(_name_key(cname))
            if not cname_canonical:
                errors.append(f"{prefix}: العميل «{cname}» غير موجود في العملاء أو قاعدة البيانات.")
                continue
            expense_customer_name = cname_canonical
        project_id = None
        if project_name:
            project_id = project_map.get(_name_key(project_name))
            if not project_id:
                errors.append(f"{prefix}: المشروع «{project_name}» غير موجود.")
                continue
        if cname and ctitle:
            matches = [k for k in contract_totals
                       if _name_key(k[0]) == _name_key(expense_customer_name)
                       and _name_key(k[1]) == _name_key(ctitle)]
            if not matches:
                errors.append(f"{prefix}: العقد «{ctitle}» غير موجود للعميل «{expense_customer_name}».")
                continue
            key = matches[0]
            if is_sub:
                sub_key = (pname, key)
                if sub_key not in sub_totals:
                    warnings.append(f"{prefix}: المقاول «{pname}» غير مربوط بعقد باطن في ملف الرفع؛ سيُسجل الصرف كمنصرف لمورد مرتبط بالعقد.")
                else:
                    sub_paid[sub_key] = sub_paid.get(sub_key, 0.0) + amount
                    if sub_paid[sub_key] > sub_totals[sub_key] + 0.001:
                        errors.append(
                            f"{prefix}: مجموع الصرف للمقاول «{pname}» ({sub_paid[sub_key]:,.3f}) يتجاوز قيمة عقد باطنه ({sub_totals[sub_key]:,.3f}).")
                        sub_paid[sub_key] -= amount
                        continue
        elif ctitle and not cname:
            errors.append(f"{prefix}: عند كتابة اسم العقد يجب كتابة اسم العميل أيضًا.")
            continue
        data["expenses"].append({
            "old_number": old_no, "party_name": pname, "is_sub": is_sub,
            "contract_key": key, "expense_customer_name": expense_customer_name,
            "project_id": project_id, "expense_account_id": expense_account_id,
            "amount": amount, "date": pdate_n,
            "description": descr or ("دفعة مقاول (مستورد)" if is_sub
                                     else "مصروف عام (مستورد)"),
            "method": method or "أخرى", "notes": notes or ""})

    # ---- تنبيه مسبق: سندات قبض تتجاوز قيمة عقدها ----
    # يُعرض في مرحلة الفحص (قبل إدخال أي شيء) كتنبيه لا كخطأ، لأن
    # النظام سيسجل الزيادة كدفعة مقدمة ولن يتوقف الاستيراد.
    try:
        con_value = {}
        for crow in data.get("contracts", []):
            ck = (crow["customer_name"],
                  crow["title"] or f"عقد بقيمة {crow['amount']}")
            con_value[ck] = con_value.get(ck, 0.0) + float(crow["amount"])
        for srow in data.get("sup_contracts", []):
            ck = (srow["customer_name"],
                  srow["title"] or f"إشراف {srow['monthly']}x{srow['months']}")
            con_value[ck] = con_value.get(ck, 0.0) + float(srow["amount"])
        paid_sum = {}
        for prow in data.get("payments", []):
            ck = prow.get("contract_key")
            if ck:
                paid_sum[ck] = paid_sum.get(ck, 0.0) + float(prow["amount"])
        for ck, paid in paid_sum.items():
            value = con_value.get(ck)
            if value is None:
                continue          # عقد موجود مسبقًا — يُفحص وقت الإدخال
            if paid > value + 0.001:
                warnings.append(
                    f"سندات القبض على عقد «{ck[1]}» للعميل «{ck[0]}» "
                    f"مجموعها {paid:,.3f} وقيمة العقد {value:,.3f} "
                    f"(زيادة {paid - value:,.3f}) · "
                    "ستُسجَّل الزيادة كدفعة مقدمة ولن يتوقف الاستيراد · "
                    "راجع قيمة العقد أو تكرار السندات.")
    except Exception:
        pass

    return {"data": data, "errors": errors, "warnings": warnings,
            "counts": {k: len(v) for k, v in data.items()}}


# ============================================================
# 3) التنفيذ الفعلي
# ============================================================
def execute_import(parsed: dict) -> dict:
    """إدخال البيانات المفحوصة · بترتيب التبعية. يرجع إحصائيات."""
    from database.models import Customer, Contract
    from services import customer_service, contract_service
    from services import lifecycle_service as lc

    if parsed["errors"]:
        raise ValidationError("لا يمكن الاستيراد · يوجد أخطاء لم تُصحح.")
    data = parsed["data"]
    stats = {"customers": 0, "contracts": 0, "subs": 0, "payments": 0,
             "sup_contracts": 0, "skipped_customers": 0, "skipped_subs": 0,
             "skipped_contracts": 0, "skipped_payments": 0,
             # سندات تجاوزت المتبقي وسُجّلت كدفعة مقدمة (للمراجعة بعد الرفع)
             "overpaid": []}

    # خرائط الاسم -> المعرف
    existing = {c.name: c.id for c in customer_service.list_customers()}
    existing_norm = {_name_key(name): (name, cid) for name, cid in existing.items()}
    customer_ids: dict[str, int] = dict(existing)
    for row in data["customers"]:
        found = existing_norm.get(_name_key(row["name"]))
        if found:
            official_name, cid = found
            customer_ids[row["name"]] = cid
            customer_ids[official_name] = cid
            stats["skipped_customers"] += 1
            continue
        cid = customer_service.add_customer(Customer(
            customer_code=customer_service.next_customer_code(),
            name=row["name"], phone=row["phone"], address=row["address"],
            email=row["email"],
            notes=(row["notes"] + " " if row["notes"] else "") + "[مستورد]"))
        customer_ids[row["name"]] = cid
        existing_norm[_name_key(row["name"])] = (row["name"], cid)
        stats["customers"] += 1

    # العقود
    contract_ids: dict[tuple, int] = {}
    try:
        with get_connection() as conn:
            for r in conn.execute("""
                SELECT c.id, cu.name AS customer_name, c.contract_name,
                       c.description, c.contract_number
                FROM contracts c JOIN customers cu ON cu.id = c.customer_id
                WHERE COALESCE(c.status,'') <> 'ملغى'
            """).fetchall():
                if (r["customer_name"] or "").strip():
                    for label in (r["contract_name"], r["description"], r["contract_number"]):
                        if (label or "").strip():
                            contract_ids[(r["customer_name"], str(label).strip())] = int(r["id"])
    except Exception:
        pass
    def _existing_contract_for(row: dict, fallback: str) -> int | None:
        cname_key = _name_key(row["customer_name"])
        labels = [row.get("title") or fallback]
        for (cn, label), cid0 in contract_ids.items():
            if _name_key(cn) == cname_key and any(_name_key(label) == _name_key(x) for x in labels if x):
                return cid0
        return None

    for row in data["contracts"]:
        found_con = _existing_contract_for(row, f"عقد بقيمة {row['amount']}")
        if found_con:
            key = (row["customer_name"], row["title"] or f"عقد بقيمة {row['amount']}")
            contract_ids[key] = found_con
            stats["skipped_contracts"] += 1
            continue
        cid = customer_ids[row["customer_name"]]
        con_id = contract_service.add_contract(Contract(
            contract_number=contract_service.next_contract_number(),
            customer_id=cid, contract_date=row["contract_date"],
            description=row["title"] or "عقد مستورد",
            amount=row["amount"], notes="[مستورد]",
            contract_type=row.get("contract_type", "")))
        with get_connection() as conn:
            conn.execute(
                "UPDATE contracts SET start_date=?, end_date=?, status=? WHERE id=?",
                (row["start_date"], row["end_date"], row["status"], con_id))
        lc.log_event(con_id, "إنشاء العقد", "استيراد من البيانات القديمة")
        key = (row["customer_name"],
               row["title"] or f"عقد بقيمة {row['amount']}")
        contract_ids[key] = con_id
        stats["contracts"] += 1

    # عقود الإشراف الهندسي
    from services import supervision_service as sup_srv
    for row in data.get("sup_contracts", []):
        found_con = _existing_contract_for(row, f"إشراف {row['monthly']}x{row['months']}")
        if found_con:
            key = (row["customer_name"], row["title"] or f"إشراف {row['monthly']}x{row['months']}")
            contract_ids[key] = found_con
            stats["skipped_contracts"] += 1
            continue
        cid = customer_ids[row["customer_name"]]
        con_id = contract_service.add_contract(Contract(
            contract_number=contract_service.next_contract_number(),
            customer_id=cid, contract_date=row["contract_date"],
            description=row["title"] or "عقد إشراف مستورد",
            amount=row["amount"], notes="[مستورد]"))
        with get_connection() as conn:
            conn.execute(
                "UPDATE contracts SET status=? WHERE id=?",
                (row["status"], con_id))
        sup_srv.setup_supervision(con_id, row["sup_start"],
                                  row["months"], row["monthly"])
        lc.log_event(con_id, "إنشاء العقد",
                     "استيراد عقد إشراف من البيانات القديمة")
        key = (row["customer_name"],
               row["title"] or f"إشراف {row['monthly']}x{row['months']}")
        contract_ids[key] = con_id
        stats["sup_contracts"] += 1

    # مقاولو الباطن
    existing_subs = {s["name"] for s in lc.list_subcontractors()}
    for row in data["subs"]:
        if row["name"] in existing_subs:
            stats["skipped_subs"] += 1
            continue
        lc.add_subcontractor(lc.next_sub_code(), row["name"],
                             phone=row["phone"], activity=row["activity"],
                             notes=(row["notes"] or "") + " [مستورد]")
        stats["subs"] += 1

    # الدفعات (سندات قبض) · تخطي الموجود بدل إيقاف الرفع
    for row in data["payments"]:
        cid = customer_ids[row["customer_name"]]
        con_id = contract_ids.get(row["contract_key"])
        old_no = (row.get("old_number") or "").strip()
        with get_connection() as conn:
            exists = None
            if old_no:
                exists = conn.execute("SELECT id FROM vouchers WHERE voucher_number=?", (old_no,)).fetchone()
            if not exists:
                exists = conn.execute(
                    "SELECT id FROM vouchers WHERE vtype='receipt' AND cancelled=0 "
                    "AND vdate=? AND party_id=? AND COALESCE(contract_id,0)=COALESCE(?,0) "
                    "AND ABS(amount-?)<0.0005 AND TRIM(COALESCE(description,''))=TRIM(COALESCE(?,'')) LIMIT 1",
                    (row["date"], cid, con_id, row["amount"], row["description"])).fetchone()
        if exists:
            stats["skipped_payments"] += 1
            continue
        # تسجيل التجاوز (إن وُجد) قبل الإدخال حتى يظهر في التقرير النهائي
        over = _receipt_overpay(lc, con_id, row["amount"])
        if over:
            stats["overpaid"].append(
                f"سند قبض بمبلغ {row['amount']:,.3f} بتاريخ "
                f"{row['date']} للعميل «{row['customer_name']}»"
                + (f" على عقد «{_con_label(row.get('contract_key'))}»"
                   if _con_label(row.get("contract_key")) else "")
                + f" تجاوز المتبقي بمقدار {over:,.3f} · "
                "سُجّل كدفعة مقدمة على العقد.")
        # allow_overpay=True: الزيادة تُسجَّل كدفعة مقدمة ولا توقف الاستيراد
        lc.add_voucher(
            "receipt", row["date"], "", row["method"], "",
            row["amount"], row["description"],
            contract_id=con_id, party_type="customer", party_id=cid,
            notes=row.get("notes", ""),
            _voucher_number=row.get("old_number", ""),
            allow_overpay=True)
        stats["payments"] += 1

    # عقود الباطن (ربط المقاول بالعقد الرئيسي)
    sub_ids = {s["name"]: s["id"] for s in lc.list_subcontractors()}
    stats["sub_contracts"] = 0
    for row in data.get("sub_contracts", []):
        con_id = contract_ids.get(row["contract_key"])
        sid = sub_ids.get(row["sub_name"])
        if not con_id or not sid:
            continue
        lc.add_sub_contract(sid, con_id, None, row["amount"],
                            row["start_date"], row["end_date"],
                            row["description"] or "عقد باطن مستورد")
        stats["sub_contracts"] += 1

    # سندات الصرف (دفعات مقاولين + مصروفات عامة)
    stats["expenses"] = 0
    stats["skipped_expenses"] = 0
    for row in data.get("expenses", []):
        if row.get("_skip_existing"):
            stats["skipped_expenses"] += 1
            continue
        pname = row["party_name"]
        sid = sub_ids.get(pname)
        if sid is None:
            # مستفيد غير مسجل (مورد/مصروف عام) · يُسجل كمقاول تلقائيًا
            sid = lc.add_subcontractor(lc.next_sub_code(), pname,
                                       activity="مورد / مصروف عام",
                                       notes="[مستورد]")
            sub_ids[pname] = sid
        con_id = contract_ids.get(row["contract_key"]) \
            if row.get("contract_key") else None
        expense_customer_id = customer_ids.get(row.get("expense_customer_name")) if row.get("expense_customer_name") else None
        # تخطي تكرار سند صرف بلا رقم قديم أيضًا: نفس التاريخ/المستفيد/العقد/المبلغ/البيان.
        with get_connection() as conn:
            dup = conn.execute(
                "SELECT id FROM vouchers WHERE vtype='payment' AND cancelled=0 AND vdate=? "
                "AND TRIM(COALESCE(party,''))=TRIM(COALESCE(?,'')) "
                "AND COALESCE(contract_id,0)=COALESCE(?,0) AND ABS(amount-?)<0.0005 "
                "AND TRIM(COALESCE(description,''))=TRIM(COALESCE(?,'')) LIMIT 1",
                (row["date"], pname, con_id, row["amount"], row["description"])).fetchone()
        if dup:
            stats["skipped_expenses"] += 1
            continue
        # allow_overpay=True: تجاوز حد المقاول لا يوقف الاستيراد
        lc.add_voucher(
            "payment", row["date"], "", row["method"], "",
            row["amount"], row["description"],
            contract_id=con_id, project_id=row.get("project_id"),
            expense_customer_id=expense_customer_id,
            party_type="subcontractor", party_id=sid,
            notes=row.get("notes", ""), _voucher_number=row.get("old_number", ""),
            expense_account_id=row.get("expense_account_id"),
            allow_overpay=True)
        stats["expenses"] += 1

    # بعد رفع دفعة كبيرة على قاعدة شبكة/شير نحدّث إحصاءات SQLite حتى تستخدم الفهارس الجديدة.
    try:
        with get_connection() as conn:
            conn.execute("PRAGMA optimize")
            conn.execute("ANALYZE")
    except Exception:
        pass
    return stats
