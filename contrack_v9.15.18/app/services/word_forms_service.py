# -*- coding: utf-8 -*-
"""خدمة نماذج Word المبسطة — Word First.

تتجنب مسار PDF بالإحداثيات. القالب = ملف .docx يحتوي {{placeholders}}.
"""
from __future__ import annotations

import os
import re
import shutil

from database.database import get_connection
from services.customer_service import ValidationError
from services import form_service as fs


def _norm(s: str) -> str:
    s = (s or "").strip().lower()
    s = re.sub(r"[\s\-]+", "_", s)
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه")
    return re.sub(r"_+", "_", s).strip("_")

SYNONYMS = {
    # customers
    "اسم_العميل": ("system", "customers", "name", "customer"),
    "العميل": ("system", "customers", "name", "customer"),
    "customer_name": ("system", "customers", "name", "customer"),
    "رقم_العميل": ("system", "customers", "customer_code", "customer"),
    "هاتف_العميل": ("system", "customers", "phone", "customer"),
    "تليفون_العميل": ("system", "customers", "phone", "customer"),
    "عنوان_العميل": ("system", "customers", "address", "customer"),
    "محافظه_العميل": ("system", "customers", "address_governorate", "customer"),
    "منطقه_العميل": ("system", "customers", "address_area", "customer"),
    "قطعه_العميل": ("system", "customers", "address_block", "customer"),
    "قسيمه_العميل": ("system", "customers", "address_plot", "customer"),
    "المحافظه": ("system", "customers", "address_governorate", "customer"),
    "المنطقه": ("system", "customers", "address_area", "customer"),
    "القطعه": ("system", "customers", "address_block", "customer"),
    "القسيمه": ("system", "customers", "address_plot", "customer"),
    # contracts
    "رقم_العقد": ("system", "contracts", "contract_number", "contract"),
    "اسم_العقد": ("system", "contracts", "contract_name", "contract"),
    "نوع_العقد": ("system", "contracts", "contract_type", "contract"),
    "تاريخ_العقد": ("system", "contracts", "contract_date", "contract"),
    "مده_التنفيذ": ("system", "contracts", "duration_days", "contract"),
    # financial
    "قيمه_العقد": ("system", "financial", "current_value", "contract"),
    "مبلغ_العقد": ("system", "financial", "current_value", "contract"),
    "القيمه_الحاليه": ("system", "financial", "current_value", "contract"),
    "قيمه_العقد_كتابه": ("system", "financial", "current_value_words", "contract"),
    "مبلغ_العقد_كتابه": ("system", "financial", "current_value_words", "contract"),
    "المبلغ_بالحروف": ("system", "financial", "current_value_words", "contract"),
    "المبلغ_كتابه": ("system", "financial", "current_value_words", "contract"),
    "المحصل": ("system", "financial", "collected", "contract"),
    "اجمالي_المحصل": ("system", "financial", "collected", "contract"),
    "المتبقي": ("system", "financial", "remaining", "contract"),
    "باقي_العقد": ("system", "financial", "remaining", "contract"),
    "جدول_الدفعات": ("system", "financial", "payments_table", "contract"),
    # system
    "تاريخ_اليوم": ("auto", "system", "today", ""),
    "اليوم": ("auto", "system", "today", ""),
    "السنه": ("auto", "system", "year", ""),
    "اسم_المستخدم": ("auto", "system", "user", ""),
    "رقم_الرخصه": ("manual", "", "", ""),
    "رقم_الترخيص": ("manual", "", "", ""),
    "اسم_المالك": ("system", "customers", "name", "customer"),
    "المالك": ("system", "customers", "name", "customer"),
    "هاتف": ("system", "customers", "phone", "customer"),
    "العنوان": ("system", "customers", "address", "customer"),
    "محافظه": ("system", "customers", "address_governorate", "customer"),
    "منطقه": ("system", "customers", "address_area", "customer"),
    "قطعه": ("system", "customers", "address_block", "customer"),
    "قسيمه": ("system", "customers", "address_plot", "customer"),
    "وصف_العقد": ("system", "contracts", "contract_name", "contract"),
    "بيان_العقد": ("system", "contracts", "contract_name", "contract"),
    "تاريخ": ("auto", "system", "today", ""),
    "التاريخ": ("auto", "system", "today", ""),
    "السنة": ("auto", "system", "year", ""),
    "المبلغ": ("system", "financial", "current_value", "contract"),
    "الاجمالي": ("system", "financial", "current_value", "contract"),
    "الباقي": ("system", "financial", "remaining", "contract"),
}
SYNONYMS = {_norm(k): v for k, v in SYNONYMS.items()}


def guess_field(placeholder: str) -> dict:
    key = _norm(placeholder)
    total_aliases = {
        "الاجمالي", "اجمالي", "المجموع", "مجموع", "اجمالي_الدفعات",
        "مجموع_الدفعات", "اجمالي_المدخلات", "مجموع_المبالغ",
    }
    total_words_aliases = {
        "الاجمالي_كتابه", "اجمالي_كتابه", "المجموع_كتابه", "مجموع_كتابه",
        "اجمالي_الدفعات_كتابه", "مجموع_الدفعات_كتابه",
        "المبلغ_بالحروف", "المبلغ_كتابه", "التفقيط", "تفقيط_المبلغ",
    }
    contract_value_aliases = {
        "قيمه_العقد", "مبلغ_العقد", "اجمالي_العقد", "القيمه_الاجماليه",
    }
    contract_words_aliases = {
        "قيمه_العقد_كتابه", "مبلغ_العقد_كتابه", "قيمه_العقد_بالحروف",
        "مبلغ_العقد_بالحروف",
    }
    remaining_aliases = {
        "المتبقي", "متبقي", "باقي_العقد", "الباقي", "المتبقي_من_العقد",
    }
    remaining_words_aliases = {
        "المتبقي_كتابه", "باقي_العقد_كتابه", "الباقي_كتابه",
        "المتبقي_بالحروف", "باقي_العقد_بالحروف",
    }
    if key in contract_value_aliases:
        return {"source_type": "manual", "source_entity": "", "source_key": "",
                "context": "", "field_type": "money", "required": 1,
                "known": True}
    if key in (total_aliases | total_words_aliases | contract_words_aliases |
               remaining_aliases | remaining_words_aliases):
        return {"source_type": "calc", "source_entity": "", "source_key": "",
                "context": "", "field_type": "text", "required": 0,
                "known": True}
    m = re.match(r"^(?:ال)?دفعه_(\d+)_(قيمه|مبلغ|كتابه|بالحروف|نسبه|اسم|عنوان|تاريخ)$", key)
    if m:
        part = m.group(2)
        if part in ("قيمه", "مبلغ"):
            return {"source_type": "manual", "source_entity": "", "source_key": "",
                    "context": "", "field_type": "money", "required": 0,
                    "known": True}
        if part in ("كتابه", "بالحروف"):
            return {"source_type": "calc", "source_entity": "", "source_key": "",
                    "context": "", "field_type": "text", "required": 0,
                    "known": True}
        return {"source_type": "manual", "source_entity": "", "source_key": "",
                "context": "", "field_type": "text", "required": 0,
                "known": True}
    guessed = SYNONYMS.get(key)
    if guessed:
        stype, entity, source_key, context = guessed
        ftype = "text" if source_key.endswith("_words") else "money" if entity == "financial" else "date" if source_key in ("today", "contract_date") else "text"
        return {"source_type": stype, "source_entity": entity,
                "source_key": source_key, "context": context,
                "field_type": ftype, "required": 0, "known": True}
    return {"source_type": "manual", "source_entity": "", "source_key": "",
            "context": "", "field_type": "text", "required": 0, "known": False}


def scan_docx(path: str) -> list[dict]:
    if not path or not os.path.exists(path):
        raise ValidationError("اختر ملف Word صالحًا.")
    placeholders = fs.scan_word_placeholders(path)
    return [{"name": ph, **guess_field(ph)} for ph in placeholders]


def import_docx(name: str, docx_path: str, description: str = "",
                category_id: int | None = None, created_by: str = "") -> dict:
    """استيراد Word دون اشتراط تحويل PDF وقت الرفع."""
    name = (name or "").strip()
    if not name:
        raise ValidationError("اسم النموذج مطلوب.")
    if not docx_path or not os.path.exists(docx_path) or not docx_path.lower().endswith(".docx"):
        raise ValidationError("اختر ملف Word بصيغة .docx")
    try:
        from docx import Document
        Document(docx_path)
    except Exception as exc:
        raise ValidationError("تعذر قراءة ملف Word. تأكد أنه غير تالف وغير مفتوح.") from exc
    fields = scan_docx(docx_path)
    with get_connection() as conn:
        dup = conn.execute("SELECT COUNT(*) FROM form_templates WHERE base_name=?", (name,)).fetchone()[0]
        if dup:
            raise ValidationError("يوجد نموذج بنفس الاسم.")
        cur = conn.execute(
            "INSERT INTO form_templates (base_name, version, description, category_id, status, "
            "pdf_path, pages_count, created_by, updated_by, source_type, word_path) "
            "VALUES (?,1,?,?, 'معتمد', '', 1, ?, ?, 'word', '')",
            (name, description, category_id, created_by, created_by))
        tid = int(cur.lastrowid)
    dest_dir = os.path.join(fs._tpl_dir(), f"tpl-{tid}")
    os.makedirs(dest_dir, exist_ok=True)
    dest = os.path.join(dest_dir, "v1.docx")
    shutil.copy2(docx_path, dest)
    with get_connection() as conn:
        conn.execute("UPDATE form_templates SET word_path=? WHERE id=?", (dest, tid))
    for f in fields:
        fs.add_field(tid, f["name"], f["field_type"], source_type=f["source_type"],
                     source_entity=f["source_entity"], source_key=f["source_key"],
                     context=f["context"], required=f["required"], placeholder=f["name"])
    return {"id": tid, "fields": fields, "placeholders": [f["name"] for f in fields]}



def import_folder(folder: str, category_id: int | None = None,
                  created_by: str = "") -> dict:
    """استيراد كل قوالب Word من مجلد واحد.

    يرجع تقريرًا بعدد الناجح والفاشل حتى لا يتوقف الاستيراد على أول ملف به مشكلة.
    """
    if not folder or not os.path.isdir(folder):
        raise ValidationError("اختر مجلد قوالب Word صالحًا.")
    ok, failed = [], []
    for fn in sorted(os.listdir(folder)):
        if fn.startswith("~$") or not fn.lower().endswith(".docx"):
            continue
        path = os.path.join(folder, fn)
        name = os.path.splitext(fn)[0]
        try:
            ok.append(import_docx(name, path, category_id=category_id,
                                  created_by=created_by))
        except Exception as exc:
            failed.append({"file": fn, "error": str(exc)})
    return {"imported": len(ok), "failed": failed, "items": ok}



def _looks_corrupt_placeholder(text: str) -> bool:
    text = text or ""
    bad = ("<w:", "</w:", "w:r", "w:t", "rsid", "proofErr", "xmlns:")
    return any(x in text for x in bad) or len(text) > 180


def repair_corrupt_template_fields(template_id: int) -> bool:
    """إصلاح القوالب التي استُوردت سابقًا بحقول XML بدل أسماء الحقول.

    لا يمس القوالب السليمة. إذا وجد حقولًا فاسدة، يعيد مسح ملف Word
    بالطريقة الجديدة، ثم يستبدل حقول القالب بأسماء نظيفة مثل: المكتب_الهندسي.
    """
    tpl = fs.get_template(template_id)
    if not tpl or not tpl.get("word_path") or not os.path.exists(tpl["word_path"]):
        return False
    current = fs.list_fields(template_id)
    if not current:
        return False
    if not any(_looks_corrupt_placeholder((f.get("placeholder") or f.get("name") or ""))
               for f in current):
        return False
    fields = scan_docx(tpl["word_path"])
    if not fields:
        return False
    with get_connection() as conn:
        conn.execute("DELETE FROM form_fields WHERE template_id=?", (template_id,))
    for f in fields:
        fs.add_field(template_id, f["name"], f["field_type"],
                     source_type=f["source_type"],
                     source_entity=f["source_entity"],
                     source_key=f["source_key"], context=f["context"],
                     required=f["required"], placeholder=f["name"])
    return True

def analyze_template(template_id: int) -> dict:
    """ملخص جودة ربط قالب Word."""
    fields = template_fields(template_id)
    total = len(fields)
    manual = sum(1 for f in fields if (f.get("source_type") or "manual") == "manual")
    system = sum(1 for f in fields if (f.get("source_type") or "") == "system")
    auto = sum(1 for f in fields if (f.get("source_type") or "") == "auto")
    calc = sum(1 for f in fields if (f.get("source_type") or "") == "calc")
    return {"total": total, "manual": manual, "system": system,
            "auto": auto, "calc": calc,
            "ready_percent": int(((system + auto + calc) / total) * 100) if total else 100}


def update_template_meta(template_id: int, base_name: str, description: str = "",
                         status: str | None = None, updated_by: str = "") -> None:
    fs.update_template_meta(template_id, base_name, description, status, updated_by)


def set_template_status(template_id: int, status: str, updated_by: str = "") -> None:
    fs.set_template_status(template_id, status, updated_by)


def delete_template(template_id: int) -> None:
    fs.delete_template(template_id)

def list_word_templates(search: str = "") -> list[dict]:
    rows = fs.list_templates(search=search, status="")
    return [r for r in rows if (r.get("source_type") or "pdf") == "word"]


def _auto_upgrade_address_field(row: dict) -> dict:
    """تحديث القوالب القديمة التي كان عنوان العميل المفصل فيها يدويًا."""
    current_type = row.get("source_type") or "manual"
    if current_type != "manual" or (row.get("source_entity") or row.get("source_key")):
        return row
    guessed = guess_field(row.get("placeholder") or row.get("name") or "")
    if (guessed.get("source_entity") == "customers" and
            str(guessed.get("source_key") or "").startswith("address_")):
        fs.update_field(row["id"], source_type="system", source_entity="customers",
                        source_key=guessed["source_key"], context="customer",
                        field_type="text", required=row.get("required") or 0)
        row = {**row, **guessed}
    return row


def template_fields(template_id: int) -> list[dict]:
    repair_corrupt_template_fields(template_id)
    rows = [_auto_upgrade_address_field(r) for r in fs.list_fields(template_id)]
    out = []
    for r in rows:
        known = bool(r.get("source_entity") or r.get("source_type") in ("auto", "calc"))
        out.append({**r, "known": known})
    return out


def update_field_mapping(field_id: int, source_type: str, source_entity: str = "",
                         source_key: str = "", context: str = "",
                         field_type: str = "text", required: int = 0) -> None:
    fs.update_field(field_id, source_type=source_type, source_entity=source_entity,
                    source_key=source_key, context=context, field_type=field_type,
                    required=required)


def create_word_document(template_id: int, title: str, values: dict,
                         customer_id: int | None = None,
                         contract_id: int | None = None,
                         by_user: str = "",
                         output_format: str = "word",
                         hide_footer: bool = False) -> dict:
    """توليد مستند من قالب Word.

    hide_footer: إخفاء تذييل المستند في النسخة المولَّدة فقط (القالب
    الأصلي لا يتغيّر) — للطباعة على ورق يحمل تذييل الشركة مسبقًا.
    """
    return fs.create_document(template_id, title, values, customer_id,
                              contract_id, by_user, output_format,
                              hide_footer=hide_footer)


def get_template(template_id: int) -> dict | None:
    return fs.get_template(template_id)
