# -*- coding: utf-8 -*-
"""قوالب Microsoft Word الذكية.

يرفع المستخدم ملف .docx يحتوي حقولاً ديناميكية مثل {{اسم_العميل}} {{رقم_العقد}}،
والنظام يملأها من بيانات العقد/عرض السعر الفعلية ويولّد Word أو PDF جاهزاً —
مع الحفاظ على تنسيق النموذج الأصلي، وإضافة هيدر/فوتر النظام للعقود وعروض الأسعار فقط.
"""
import json
import os
import re
import shutil

from database.database import BASE_DIR, get_connection
from services.customer_service import ValidationError


def _docx():
    """استيراد كسول لمكتبة python-docx مع رسالة عربية واضحة عند غيابها."""
    try:
        from docx import Document
        return Document
    except ImportError as exc:
        raise WordTemplateError(
            "مكتبة قوالب Word غير مثبتة.\n"
            "ثبّتها بالأمر:  pip install python-docx\n"
            "ثم أعد تشغيل البرنامج.") from exc

TEMPLATES_DIR = os.path.join(BASE_DIR, "assets", "word_templates")

# الحقول المدعومة (تظهر للمستخدم كدليل)
SUPPORTED_FIELDS = [
    "اسم_الشركة", "عنوان_الشركة", "هاتف_الشركة", "بريد_الشركة",
    "اسم_العميل", "رقم_العميل", "هاتف_العميل", "عنوان_العميل",
    "رقم_العقد", "اسم_العقد", "تاريخ_العقد", "تاريخ_التوقيع",
    "تاريخ_البداية", "تاريخ_النهاية", "مدة_العقد",
    "قيمة_العقد", "القيمة_الأصلية", "قيمة_العقد_كتابة", "المبلغ_بالحروف", "العملة",
    "اسم_المشروع", "رقم_المشروع", "وصف_العقد", "نوع_العقد",
    "جدول_البنود", "جدول_الدفعات", "دفعة_1_قيمة", "دفعة_1_كتابة",
    "دفعة_1_نسبة", "دفعة_1_اسم", "دفعة_2_قيمة", "دفعة_2_كتابة",
    "اليوم", "التاريخ_الحالي", "يوم_الأسبوع", "اسم_اليوم",
]

_PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


_AR_WEEKDAYS = ("الاثنين", "الثلاثاء", "الأربعاء", "الخميس",
                "الجمعة", "السبت", "الأحد")


def _weekday_name(day=None) -> str:
    """اسم اليوم بالعربية (اليوم افتراضيًا)."""
    from datetime import date as _date
    day = day or _date.today()
    try:
        return _AR_WEEKDAYS[day.weekday() % 7]
    except Exception:
        return ""


class WordTemplateError(Exception):
    """خطأ قوالب Word برسالة عربية."""


# ============================================================
# إدارة القوالب المرفوعة
# ============================================================
def upload_template(src_path: str, template_kind: str = "contract") -> str:
    """نسخ قالب Word مرفوع إلى مجلد القوالب. يرجع المسار الجديد."""
    if not os.path.exists(src_path):
        raise WordTemplateError("الملف غير موجود.")
    if not src_path.lower().endswith(".docx"):
        raise WordTemplateError(
            "يجب أن يكون الملف بصيغة Word الحديثة (.docx).\n"
            "إذا كان ملفك .doc قديمًا افتحه في Word واحفظه كـ .docx أولاً.")
    # التحقق أنه ملف docx سليم
    Document = _docx()
    try:
        Document(src_path)
    except Exception as exc:
        raise WordTemplateError("تعذر قراءة الملف، تأكد أنه ملف Word سليم.") from exc
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    base = os.path.splitext(os.path.basename(src_path))[0]
    safe = re.sub(r'[\\/:*?"<>|]', "_", base)[:50]
    dest = os.path.join(TEMPLATES_DIR, f"{template_kind}__{safe}.docx")
    counter = 1
    while os.path.exists(dest):
        dest = os.path.join(TEMPLATES_DIR,
                            f"{template_kind}__{safe}_{counter}.docx")
        counter += 1
    shutil.copy2(src_path, dest)
    return dest


def list_templates(template_kind: str | None = None) -> list[dict]:
    """قائمة القوالب المرفوعة."""
    if not os.path.isdir(TEMPLATES_DIR):
        return []
    result = []
    for name in sorted(os.listdir(TEMPLATES_DIR)):
        if not name.lower().endswith(".docx"):
            continue
        kind = name.split("__")[0] if "__" in name else "contract"
        if template_kind and kind != template_kind:
            continue
        # الاسم المعروض: آخر جزء بعد بادئات النوع
        display = name.split("__")[-1].rsplit(".docx", 1)[0].replace("_", " ")
        result.append({"path": os.path.join(TEMPLATES_DIR, name),
                       "name": display, "kind": kind})
    return result


def delete_template(path: str) -> None:
    real = os.path.realpath(path)
    if not real.startswith(os.path.realpath(TEMPLATES_DIR)):
        raise WordTemplateError("الملف خارج مجلد القوالب.")
    os.remove(real)


def scan_placeholders(template_path: str) -> list[str]:
    """استخراج الحقول الديناميكية الموجودة في القالب."""
    doc = _docx()(template_path)
    found = set()

    def scan_text(text):
        for m in _PLACEHOLDER_RE.finditer(text):
            found.add(m.group(1).strip())

    for p in doc.paragraphs:
        scan_text(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    scan_text(p.text)
    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                scan_text(p.text)
    return sorted(found)


# ============================================================
# ربط حقول القوالب الرسمية (العقود/عروض الأسعار)
# نفس فكرة مركز النماذج العام لكن بدون تحويل القالب الرسمي إلى سجل نموذج عام.
# ============================================================
def _official_maps_path() -> str:
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    return os.path.join(TEMPLATES_DIR, "_official_field_maps.json")


def _official_map_key(template_path: str) -> str:
    try:
        return os.path.basename(os.path.realpath(template_path))
    except Exception:
        return os.path.basename(template_path or "")


def _load_official_maps() -> dict:
    path = _official_maps_path()
    if not os.path.exists(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _save_official_maps(data: dict) -> None:
    path = _official_maps_path()
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data or {}, f, ensure_ascii=False, indent=2)
    os.replace(tmp, path)


def official_template_fields(template_path: str) -> list[dict]:
    """حقول قالب رسمي مع ربط قابل للتعديل مثل باقي النماذج."""
    from services import word_forms_service as wfs
    names = scan_placeholders(template_path)
    data = _load_official_maps()
    saved = data.get(_official_map_key(template_path), {})
    out = []
    for name in names:
        guessed = wfs.guess_field(name)
        field = {"name": name, "placeholder": name, **guessed}
        if isinstance(saved.get(name), dict):
            field.update(saved[name])
        if ((field.get("source_type") or "manual") == "manual" and
                not (field.get("source_entity") or field.get("source_key")) and
                guessed.get("source_entity") == "customers" and
                str(guessed.get("source_key") or "").startswith("address_")):
            field.update(guessed)
        field.setdefault("field_type", "text")
        field.setdefault("source_type", "manual")
        field.setdefault("source_entity", "")
        field.setdefault("source_key", "")
        field.setdefault("context", "")
        field.setdefault("required", 0)
        out.append(field)
    return out


def update_official_field_mapping(template_path: str, placeholder: str, *,
                                  source_type: str = "manual",
                                  source_entity: str = "", source_key: str = "",
                                  context: str = "", field_type: str = "text",
                                  required: int = 0) -> None:
    if not placeholder:
        raise WordTemplateError("اسم الحقل غير صحيح.")
    data = _load_official_maps()
    key = _official_map_key(template_path)
    data.setdefault(key, {})[placeholder] = {
        "source_type": source_type or "manual",
        "source_entity": source_entity or "",
        "source_key": source_key or "",
        "context": context or "",
        "field_type": field_type or "text",
        "required": int(required or 0),
    }
    _save_official_maps(data)


def official_field_values(template_path: str, obj_id: int | None, *,
                          is_quote: bool = False) -> dict[str, str]:
    """قيم شاشة تعبئة القالب الرسمي حسب الربط المحفوظ، مع قابلية التعديل اليدوي."""
    from services import form_service as fs
    base: dict[str, str] = {}
    ctx: dict = {}
    if obj_id:
        try:
            base = quotation_field_values(int(obj_id)) if is_quote else contract_field_values(int(obj_id))
        except Exception:
            base = {}
        if not is_quote:
            try:
                ctx = fs.build_context(contract_id=int(obj_id))
            except Exception:
                ctx = {}
    result = {}
    for f in official_template_fields(template_path):
        ph = f.get("placeholder") or f.get("name") or ""
        st = f.get("source_type") or "manual"
        value = ""
        if st in ("system", "auto") and (f.get("source_entity") or ""):
            value = fs.resolve_field(f, ctx) if ctx or f.get("source_entity") in ("settings", "system") else ""
        elif st == "calc":
            value = base.get(ph, "")
        elif st != "manual":
            value = base.get(ph, "")
        else:
            # الحقول اليدوية تبقى فارغة إلا إذا كان اسمها مطابقًا لقيمة معروفة؛
            # تظهر القيمة للمراجعة ويمكن تعديلها، ولا تُعتبر ربطًا إجباريًا.
            value = base.get(ph, "")
        result[ph] = str(value or "")
    return result


# ============================================================
# بناء بيانات الحقول من العقد
# ============================================================
def _amount_in_words(value: float, currency: str) -> str:
    """تفقيط مبلغ العقد/العرض بدقة 3 خانات للدينار الكويتي."""
    try:
        from services.tafqeet import amount_to_words
        return amount_to_words(float(value or 0), currency or "د.ك")
    except Exception:
        return f"{float(value or 0):,.3f} {currency}"


def _payment_percent_from_notes(notes: str) -> str:
    m = re.search(r"نسبة\s+([0-9]+(?:\.[0-9]+)?)%", notes or "")
    return (m.group(1) + "%") if m else ""


def _contract_payment_values(contract_id: int, contract_type: str,
                             amount: float, currency: str) -> dict[str, str]:
    """حقول دفعات العقد: قيمة/نسبة/تفقيط وجدول كامل.

    الأولوية لخطة السداد الفعلية `payment_plan` إن وُلدت عند إنشاء العقد.
    وإن لم توجد بعد، نحسبها من قالب دفعات نوع العقد بنفس منطق المعاينة.
    """
    rows = []
    with get_connection() as conn:
        plan = conn.execute(
            "SELECT title, due_date, amount, sort_order, notes FROM payment_plan "
            "WHERE contract_id=? ORDER BY sort_order, id",
            (contract_id,)).fetchall()
    if plan:
        for p in plan:
            title = re.sub(r"^\s*\d+\.\s*", "", p["title"] or "دفعة")
            amount_f = float(p["amount"] or 0)
            rows.append({
                "title": title,
                "amount": amount_f,
                "amount_text": f"{amount_f:,.3f}",
                "words": _amount_in_words(amount_f, currency),
                "pct": _payment_percent_from_notes(p["notes"] or ""),
                "due": p["due_date"] or "",
            })
    else:
        try:
            from services import workflow_service as wf
            preview = wf.type_payment_preview(contract_type, amount)
            for p in preview.get("rows", []):
                amount_f = float(p.get("amount") or 0)
                rows.append({
                    "title": p.get("title") or "دفعة",
                    "amount": amount_f,
                    "amount_text": f"{amount_f:,.3f}" if amount_f else "",
                    "words": _amount_in_words(amount_f, currency) if amount_f else "",
                    "pct": (f"{float(p.get('pct') or 0):g}%" if float(p.get("pct") or 0) else ""),
                    "due": p.get("due_text") or "",
                })
        except Exception:
            rows = []
    values = {}
    table_lines = []
    for idx, p in enumerate(rows, 1):
        values[f"دفعة_{idx}_اسم"] = p["title"]
        values[f"دفعة_{idx}_قيمة"] = p["amount_text"]
        values[f"دفعة_{idx}_كتابة"] = p["words"]
        values[f"دفعة_{idx}_نسبة"] = p["pct"]
        values[f"دفعة_{idx}_تاريخ"] = p["due"]
        values[f"الدفعة_{idx}_اسم"] = p["title"]
        values[f"الدفعة_{idx}_قيمة"] = p["amount_text"]
        values[f"الدفعة_{idx}_كتابة"] = p["words"]
        values[f"الدفعة_{idx}_نسبة"] = p["pct"]
        values[f"الدفعة_{idx}_تاريخ"] = p["due"]
        table_lines.append(
            f"{idx}. {p['title']}: {p['amount_text']} {currency}"
            + (f" ({p['pct']})" if p["pct"] else "")
            + (f" — {p['due']}" if p["due"] else ""))
    values["جدول_الدفعات"] = "\n".join(table_lines) or "—"
    values["عدد_الدفعات"] = str(len(rows))
    return values


def contract_field_values(contract_id: int) -> dict[str, str]:
    """كل قيم الحقول الديناميكية لعقد معين."""
    from datetime import date
    from services import settings_service
    from services import lifecycle_service as lc

    dash = lc.contract_dashboard(contract_id)
    contract = dash["contract"]
    settings = settings_service.get_settings()
    tl = dash["timeline"]
    currency = contract.get("currency") or settings.currency
    payment_values = _contract_payment_values(
        contract_id, contract.get("contract_type") or "", float(dash.get("current_value") or contract.get("amount") or 0), currency)

    with get_connection() as conn:
        customer = conn.execute(
            "SELECT * FROM customers WHERE id=?",
            (contract["customer_id"],)).fetchone()
        project = None
        if contract.get("project_id"):
            project = conn.execute(
                "SELECT * FROM projects WHERE id=?",
                (contract["project_id"],)).fetchone()
        boq = conn.execute(
            "SELECT description, qty, unit, unit_price, qty*unit_price AS total "
            "FROM boq_items WHERE contract_id=? "
            "AND COALESCE(deleted_addendum_id,0)=0 ORDER BY id",
            (contract_id,)).fetchall()

    boq_lines = "\n".join(
        f"- {b['description']}: {b['qty']:,.3f} {b['unit'] or ''} × "
        f"{b['unit_price']:,.3f} = {b['total']:,.3f} {currency}"
        for b in boq) or "—"

    return {
        "اسم_الشركة": settings.company_name,
        "عنوان_الشركة": settings.company_address,
        "هاتف_الشركة": settings.company_phone,
        "بريد_الشركة": settings.company_email,
        "اسم_العميل": customer["name"] if customer else "",
        "رقم_العميل": customer["customer_code"] if customer else "",
        "هاتف_العميل": (customer["phone"] or "") if customer else "",
        "عنوان_العميل": (customer["address"] or "") if customer else "",
        "محافظة_العميل": (customer["address_governorate"] or "") if customer else "",
        "منطقة_العميل": (customer["address_area"] or "") if customer else "",
        "قطعة_العميل": (customer["address_block"] or "") if customer else "",
        "قسيمة_العميل": (customer["address_plot"] or "") if customer else "",
        "المحافظة": (customer["address_governorate"] or "") if customer else "",
        "المنطقة": (customer["address_area"] or "") if customer else "",
        "القطعة": (customer["address_block"] or "") if customer else "",
        "القسيمة": (customer["address_plot"] or "") if customer else "",
        "محافظة": (customer["address_governorate"] or "") if customer else "",
        "منطقة": (customer["address_area"] or "") if customer else "",
        "قطعة": (customer["address_block"] or "") if customer else "",
        "قسيمة": (customer["address_plot"] or "") if customer else "",
        "رقم_العقد": contract["contract_number"],
        "اسم_العقد": contract.get("contract_name") or "",
        "تاريخ_العقد": contract["contract_date"],
        "تاريخ_التوقيع": contract.get("sign_date") or contract["contract_date"],
        "تاريخ_البداية": contract.get("start_date") or "",
        "تاريخ_النهاية": contract.get("end_date") or "",
        "مدة_العقد": str(tl["duration"]) if tl["duration"] else "",
        "قيمة_العقد": f"{dash['current_value']:,.3f}",
        "القيمة_الأصلية": f"{dash['original_value']:,.3f}",
        "قيمة_العقد_كتابة": _amount_in_words(dash["current_value"], currency),
        "المبلغ_بالحروف": _amount_in_words(dash["current_value"], currency),
        "العملة": currency,
        "اسم_المشروع": project["name"] if project else "",
        "رقم_المشروع": project["project_code"] if project else "",
        "وصف_العقد": contract.get("description") or "",
        "نوع_العقد": contract.get("contract_type") or "",
        "جدول_البنود": boq_lines,
        "اليوم": date.today().isoformat(),
        "التاريخ_الحالي": date.today().isoformat(),
        "يوم_الأسبوع": _weekday_name(),
        "اسم_اليوم": _weekday_name(),
        **payment_values,
    }


# ============================================================
# التعبئة · مع الحفاظ على التنسيق
# ============================================================
def _replace_in_paragraph(paragraph, values: dict[str, str]) -> None:
    """استبدال الحقول في فقرة مع الحفاظ على التنسيق.

    المشكلة المعروفة: Word قد يقسم {{الحقل}} على عدة runs.
    الحل: العمل على نص الفقرة كاملاً ثم إعادة توزيعه على الـ runs.
    """
    full_text = "".join(run.text for run in paragraph.runs)
    if "{{" not in full_text:
        return

    def substitute(match):
        key = match.group(1).strip()
        return values.get(key, match.group(0))

    new_text = _PLACEHOLDER_RE.sub(substitute, full_text)
    if new_text == full_text:
        return
    # وضع النص الجديد في أول run والحفاظ على تنسيقه، وإفراغ الباقي
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""



def _stamp_document(doc) -> None:
    """بيانات المستند باسم الشركة · بدل الاسم الافتراضي للمكتبة."""
    try:
        from services import settings_service as _ss
        settings = _ss.get_settings()
        company = (getattr(settings, "company_name", "") or "").strip()
        props = doc.core_properties
        props.author = company or _ss.app_name()
        props.last_modified_by = props.author
        props.company = company
        props.comments = ""
    except Exception:
        pass


def _clear_container(container) -> None:
    """تفريغ هيدر/فوتر Word بشكل آمن قبل وضع هيدر/فوتر النظام."""
    try:
        for p in container.paragraphs:
            for run in p.runs:
                run.text = ""
        for table in list(container.tables):
            try:
                table._element.getparent().remove(table._element)
            except Exception:
                pass
    except Exception:
        pass


def _add_image_to_container(container, image_path: str, width) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    if not image_path or not os.path.exists(image_path):
        return
    _clear_container(container)
    p = container.paragraphs[0] if container.paragraphs else container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(image_path, width=width)


def apply_system_letterhead(doc, *, enabled: bool = True) -> None:
    """إضافة هيدر/فوتر النظام إلى Word للعقود وعروض الأسعار فقط.

    لا تُستخدم هذه الدالة في مركز النماذج العام، حتى لا يتم فرض الهيدر/الفوتر
    على كل نموذج، وإنما فقط مستندات العقود وعروض الأسعار المطلوبة رسميًا.
    """
    if not enabled:
        return
    try:
        from services import settings_service as _ss
        settings = _ss.get_settings()
        header = (getattr(settings, "header_image", "") or "").strip()
        footer = (getattr(settings, "footer_image", "") or "").strip()
        if not header and not footer:
            return
        for section in doc.sections:
            width = section.page_width - section.left_margin - section.right_margin
            if header:
                _add_image_to_container(section.header, header, width)
            if footer:
                _add_image_to_container(section.footer, footer, width)
    except Exception:
        # لا نوقف توليد العقد إذا كانت صورة الهيدر/الفوتر غير صالحة.
        pass


def word_to_pdf(docx_path: str, out_pdf: str) -> str:
    """تحويل Word إلى PDF عبر LibreOffice عند توفره."""
    from services.form_service import word_to_pdf as _convert
    return _convert(docx_path, out_pdf)


def fill_template(template_path: str, contract_id: int, out_path: str) -> str:
    """تعبئة قالب Word ببيانات العقد وحفظه في out_path."""
    if not os.path.exists(template_path):
        raise WordTemplateError("ملف القالب غير موجود.")
    values = contract_field_values(contract_id)
    Document = _docx()
    try:
        doc = Document(template_path)
    except Exception as exc:
        raise WordTemplateError("تعذر فتح القالب.") from exc

    for p in doc.paragraphs:
        _replace_in_paragraph(p, values)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p, values)
    # الهيدر والفوتر في ملف Word نفسه
    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                _replace_in_paragraph(p, values)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            _replace_in_paragraph(p, values)
    apply_system_letterhead(doc, enabled=True)
    _stamp_document(doc)
    try:
        doc.save(out_path)
    except PermissionError as exc:
        raise WordTemplateError(
            "تعذر الحفظ · الملف مفتوح في برنامج آخر.") from exc
    return out_path


def quotation_field_values(quotation_id: int) -> dict[str, str]:
    """كل قيم الحقول الديناميكية لعرض سعر معين."""
    from datetime import date
    from services import settings_service
    from services import lifecycle_service as lc

    quote = lc.get_quotation(quotation_id)
    if not quote:
        raise WordTemplateError("عرض السعر غير موجود.")
    settings = settings_service.get_settings()
    currency = quote.get("currency") or settings.currency
    with get_connection() as conn:
        customer = conn.execute(
            "SELECT * FROM customers WHERE id=?",
            (quote.get("customer_id"),)).fetchone() if quote.get("customer_id") else None
    items = quote.get("items") or []
    item_lines = "\n".join(
        f"- {i['description']}: {float(i.get('qty') or 1):,.3f} {i.get('unit') or ''} × "
        f"{float(i.get('unit_price') or 0):,.3f} = {float(i.get('total') or 0):,.3f} {currency}"
        for i in items) or "—"
    return {
        "اسم_الشركة": settings.company_name,
        "عنوان_الشركة": settings.company_address,
        "هاتف_الشركة": settings.company_phone,
        "بريد_الشركة": settings.company_email,
        "اسم_العميل": quote.get("customer_name") or "",
        "رقم_العميل": (customer["customer_code"] or "") if customer else "",
        "هاتف_العميل": (customer["phone"] or "") if customer else "",
        "عنوان_العميل": (customer["address"] or "") if customer else "",
        "محافظة_العميل": (customer["address_governorate"] or "") if customer else "",
        "منطقة_العميل": (customer["address_area"] or "") if customer else "",
        "قطعة_العميل": (customer["address_block"] or "") if customer else "",
        "قسيمة_العميل": (customer["address_plot"] or "") if customer else "",
        "المحافظة": (customer["address_governorate"] or "") if customer else "",
        "المنطقة": (customer["address_area"] or "") if customer else "",
        "القطعة": (customer["address_block"] or "") if customer else "",
        "القسيمة": (customer["address_plot"] or "") if customer else "",
        "محافظة": (customer["address_governorate"] or "") if customer else "",
        "منطقة": (customer["address_area"] or "") if customer else "",
        "قطعة": (customer["address_block"] or "") if customer else "",
        "قسيمة": (customer["address_plot"] or "") if customer else "",
        "رقم_عرض_السعر": quote.get("quote_number") or "",
        "تاريخ_عرض_السعر": quote.get("qdate") or "",
        "اسم_المشروع": quote.get("project_name") or "",
        "جدول_البنود": item_lines,
        "الإجمالي_قبل_الخصم": f"{float(quote.get('subtotal') or 0):,.3f}",
        "الخصم": f"{float(quote.get('discount') or 0):,.3f}",
        "الضريبة": f"{float(quote.get('tax') or 0):,.3f}",
        "قيمة_العقد": f"{float(quote.get('total') or 0):,.3f}",
        "الإجمالي": f"{float(quote.get('total') or 0):,.3f}",
        "قيمة_العقد_كتابة": _amount_in_words(float(quote.get('total') or 0), currency),
        "الإجمالي_كتابة": _amount_in_words(float(quote.get('total') or 0), currency),
        "المبلغ_بالحروف": _amount_in_words(float(quote.get('total') or 0), currency),
        "العملة": currency,
        "مدة_العقد": str(quote.get("execution_period") or ""),
        "شروط_الدفع": quote.get("payment_terms") or "",
        "شروط_التنفيذ": quote.get("execution_terms") or "",
        "ملاحظات": quote.get("notes") or "",
        "اليوم": date.today().isoformat(),
        "التاريخ_الحالي": date.today().isoformat(),
        "يوم_الأسبوع": _weekday_name(),
        "اسم_اليوم": _weekday_name(),
    }


def fill_template_with_values(template_path: str, values: dict[str, str], out_path: str,
                              *, letterhead: bool = True) -> str:
    """تعبئة قالب Word بقيم يدوية/معدلة من شاشة الإدخال.

    تُستخدم لقوالب العقود/العروض الرسمية عندما يريد المستخدم نفس نظام
    النماذج العامة: كل حقل يظهر قابلًا للتعديل حتى لو كانت له قيمة تلقائية.
    """
    if not os.path.exists(template_path):
        raise WordTemplateError("ملف القالب غير موجود.")
    Document = _docx()
    try:
        doc = Document(template_path)
    except Exception as exc:
        raise WordTemplateError("تعذر فتح القالب.") from exc
    values = {str(k): str(v or "") for k, v in (values or {}).items()}
    for p in doc.paragraphs:
        _replace_in_paragraph(p, values)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p, values)
    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                _replace_in_paragraph(p, values)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            _replace_in_paragraph(p, values)
    apply_system_letterhead(doc, enabled=letterhead)
    _stamp_document(doc)
    try:
        doc.save(out_path)
    except PermissionError as exc:
        raise WordTemplateError("تعذر الحفظ · الملف مفتوح في برنامج آخر.") from exc
    return out_path


def fill_quotation_template(template_path: str, quotation_id: int, out_path: str) -> str:
    """تعبئة قالب Word ببيانات عرض السعر مع هيدر/فوتر النظام."""
    if not os.path.exists(template_path):
        raise WordTemplateError("ملف القالب غير موجود.")
    values = quotation_field_values(quotation_id)
    Document = _docx()
    try:
        doc = Document(template_path)
    except Exception as exc:
        raise WordTemplateError("تعذر فتح قالب عرض السعر.") from exc
    for p in doc.paragraphs:
        _replace_in_paragraph(p, values)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p, values)
    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                _replace_in_paragraph(p, values)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            _replace_in_paragraph(p, values)
    apply_system_letterhead(doc, enabled=True)
    _stamp_document(doc)
    try:
        doc.save(out_path)
    except PermissionError as exc:
        raise WordTemplateError("تعذر الحفظ · الملف مفتوح في برنامج آخر.") from exc
    return out_path


def create_sample_template(kind: str = "contract") -> str:
    """إنشاء قالب Word نموذجي جاهز يوضح طريقة استخدام الحقول."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    Document = _docx()
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    doc = Document()

    def add_p(text, size=12, bold=False, center=False):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # RTL
        from docx.oxml.ns import qn
        pPr = p._p.get_or_add_pPr()
        bidi = pPr.makeelement(qn('w:bidi'), {})
        pPr.append(bidi)
        return p

    if kind == "contract":
        add_p("{{اسم_الشركة}}", 16, bold=True, center=True)
        add_p("{{عنوان_الشركة}} · هاتف: {{هاتف_الشركة}}", 10, center=True)
        add_p("", 8)
        add_p("عقد اتفاق رقم {{رقم_العقد}}", 15, bold=True, center=True)
        add_p("", 8)
        add_p("إنه في يوم {{التاريخ_الحالي}} تم الاتفاق بين كل من:", 12)
        add_p("الطرف الأول: {{اسم_الشركة}}", 12, bold=True)
        add_p("الطرف الثاني: {{اسم_العميل}} · هاتف: {{هاتف_العميل}}", 12, bold=True)
        add_p("العنوان: {{محافظة_العميل}} - {{منطقة_العميل}} - قطعة {{قطعة_العميل}} - قسيمة {{قسيمة_العميل}}", 11)
        add_p("", 8)
        add_p("المشروع: {{اسم_المشروع}}", 12)
        add_p("قيمة العقد: {{قيمة_العقد}} {{العملة}} ({{قيمة_العقد_كتابة}})", 12)
        add_p("مدة العقد: {{مدة_العقد}} يوم · من {{تاريخ_البداية}} إلى {{تاريخ_النهاية}}", 12)
        add_p("", 8)
        add_p("نطاق العمل:", 13, bold=True)
        add_p("{{وصف_العقد}}", 12)
        add_p("{{جدول_البنود}}", 11)
        add_p("", 20)
        add_p("توقيع الطرف الأول: ______________          توقيع الطرف الثاني: ______________", 12)
        name = "نموذج_عقد_جاهز"
    else:
        add_p("{{اسم_الشركة}}", 16, bold=True, center=True)
        add_p("{{عنوان_الشركة}} · هاتف: {{هاتف_الشركة}}", 10, center=True)
        add_p("", 8)
        add_p("عرض سعر رقم {{رقم_عرض_السعر}}", 15, bold=True, center=True)
        add_p("", 8)
        add_p("التاريخ: {{تاريخ_عرض_السعر}}", 12)
        add_p("السادة / {{اسم_العميل}} المحترمين", 12, bold=True)
        add_p("العنوان: {{محافظة_العميل}} - {{منطقة_العميل}} - قطعة {{قطعة_العميل}} - قسيمة {{قسيمة_العميل}}", 11)
        add_p("تحية طيبة وبعد،", 12)
        add_p("يسرنا تقديم عرض سعرنا للأعمال التالية الخاصة بـ {{اسم_المشروع}}:", 12)
        add_p("{{جدول_البنود}}", 11)
        add_p("", 8)
        add_p("الإجمالي: {{قيمة_العقد}} {{العملة}}", 13, bold=True)
        add_p("مدة التنفيذ: {{مدة_العقد}} يوم", 12)
        add_p("", 16)
        add_p("وتفضلوا بقبول فائق الاحترام،", 12)
        add_p("{{اسم_الشركة}}", 12, bold=True)
        name = "نموذج_عرض_سعر_جاهز"

    _stamp_document(doc)
    dest = os.path.join(TEMPLATES_DIR, f"{kind}__{name}.docx")
    doc.save(dest)
    return dest


# ============================================================
# ربط القوالب بأنواع العقود · قالب لكل نوع عقد
# ============================================================
def type_kind(type_name: str) -> str:
    """اسم «نوع القالب» المخزَّن لكل نوع عقد.

    القوالب تُحفظ باسم `<kind>__<اسم الملف>.docx`؛ لذلك تُستبعد الشرطة
    السفلية وحروف المسارات من اسم النوع.
    """
    safe = re.sub(r'[\\/:*?"<>|_]', "-", (type_name or "").strip())
    return f"contract-{safe}" if safe else "contract"


def templates_for_contract(contract_id: int) -> list[dict]:
    """قوالب هذا العقد: قوالب نوعه أولًا ثم القوالب العامة.

    كل عنصر: {path, name, kind, scope} حيث scope = «قالب النوع» أو «عام».
    """
    with get_connection() as conn:
        row = conn.execute(
            "SELECT COALESCE(contract_type,'') AS t FROM contracts WHERE id=?",
            (contract_id,)).fetchone()
    ctype = (row["t"] if row else "") or ""
    result = []
    if ctype:
        for t in list_templates(type_kind(ctype)):
            result.append({**t, "scope": f"قالب «{ctype}»"})
    for t in list_templates("contract"):
        result.append({**t, "scope": "قالب عام"})
    return result
