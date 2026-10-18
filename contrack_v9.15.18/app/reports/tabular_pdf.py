# -*- coding: utf-8 -*-
"""مولّد تقارير جدولية موحّد · PDF أفقي + Excel · لكل تقارير مركز التقارير.

يستقبل: عنوان التقرير + رؤوس الأعمدة + الصفوف + صف إجمالي اختياري.
"""
from datetime import datetime

from services import settings_service


def _require_pdf_dependencies():
    """تحميل مكتبات PDF وقت التصدير فقط حتى لا يتعطل Excel عند نقص reportlab."""
    try:
        from reportlab.lib import colors as _c
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas as rl_canvas
        from reports.pdf_utils import (
            apply_pdf_metadata, ar, register_fonts, hex_color, light_tint,
            PDFError, AR_FONT, AR_FONT_BOLD, PRIMARY, ACCENT, BORDER, GREEN, RED,
            safe_image, draw_header_image, draw_footer_image,
            footer_image_height,
        )
    except ModuleNotFoundError as exc:
        missing = exc.name or "reportlab"
        raise RuntimeError(
            f"مكتبة {missing} غير مثبتة. لتصدير PDF شغّل INSTALL.bat أو ثبّت المتطلبات من requirements.txt."
        ) from exc
    return {
        "_c": _c, "A4": A4, "landscape": landscape, "mm": mm,
        "rl_canvas": rl_canvas, "apply_pdf_metadata": apply_pdf_metadata,
        "ar": ar, "register_fonts": register_fonts, "hex_color": hex_color,
        "light_tint": light_tint, "PDFError": PDFError,
        "AR_FONT": AR_FONT, "AR_FONT_BOLD": AR_FONT_BOLD,
        "PRIMARY": PRIMARY, "ACCENT": ACCENT, "BORDER": BORDER,
        "GREEN": GREEN, "RED": RED, "safe_image": safe_image,
        "draw_header_image": draw_header_image,
        "draw_footer_image": draw_footer_image,
        "footer_image_height": footer_image_height,
    }


def generate_table_pdf(title: str, heads: list[str], col_widths: list[float],
                       rows: list[list], out_path: str,
                       totals_row: list | None = None,
                       footer_note: str = "",
                       color_cols: dict | None = None,
                       orientation: str = "landscape",
                       row_kinds: list | None = None,
                       right_align_cols: tuple | list = (),
                       number_cols: list | tuple | None = None,
                       accent_rows: dict | None = None,
                       keep_accents_in_economy: bool = False) -> str:
    """توليد PDF جدولي.

    color_cols: {فهرس_عمود: "green"|"red"|"auto"} · auto: أحمر إن > 0 وإلا أخضر.
    col_widths بالمليمتر · تُطبَّع تلقائياً لتملأ عرض الصفحة.
    orientation: "landscape" أفقي (الافتراضي) أو "portrait" رأسي.
    row_kinds: [(نوع الصف، عدد الأبناء)] — صف "customer" يُبرز بخط عريض
               وخلفية مميزة، وتُطبع كل الصفوف مفتوحة دائمًا.
    right_align_cols: أعمدة نصية تُحاذى لليمين بدل الوسط.
    accent_rows: {نص الخلية الثانية: لون hex} — صف مميز (إجمالي/رصيد)
               بخلفية ملوّنة وخط عريض أكبر (مرآة المعاينة).
    keep_accents_in_economy: إبقاء تلوين الصفوف المميزة حتى في الوضع
               الاقتصادي (استثناء التقرير المالي — V3.19).
    """
    pdf = _require_pdf_dependencies()
    _c = pdf["_c"]
    mm = pdf["mm"]
    if str(orientation).lower().startswith("port"):
        PAGE_W, PAGE_H = pdf["A4"]              # رأسي
    else:
        PAGE_W, PAGE_H = pdf["landscape"](pdf["A4"])   # أفقي
    MARGIN = 12 * mm
    ROW_H = 10 * mm          # صف أعلى ليتسع للخط الأكبر
    # أحجام الخطوط (أكبر من السابق لوضوح الطباعة)
    # ===== معيار موحّد لخطوط كل التقارير =====
    # ثابت في كل تقرير مهما اختلفت أعمدته أو ألوانه، حتى تخرج
    # المطبوعات بهوية بصرية واحدة (كان البولد يُشتق من تلوين
    # العمود فيختلف شكل كل تقرير عن الآخر).
    F_CELL = 10.5      # نص الخلية العادي
    F_HEAD = 11.0      # رأس الجدول
    F_TOTAL = 11.5     # صف الإجمالي
    if str(orientation).lower().startswith("port"):
        # الرأسي أضيق — نصغّر قليلًا حتى لا تتزاحم الأعمدة
        F_CELL, F_HEAD, F_TOTAL = 9.0, 9.5, 10.0
        ROW_H = 9 * mm
    rl_canvas = pdf["rl_canvas"]
    apply_pdf_metadata = pdf["apply_pdf_metadata"]
    ar = pdf["ar"]
    register_fonts = pdf["register_fonts"]
    hex_color = pdf["hex_color"]
    light_tint = pdf["light_tint"]
    PDFError = pdf["PDFError"]
    AR_FONT = pdf["AR_FONT"]
    AR_FONT_BOLD = pdf["AR_FONT_BOLD"]
    PRIMARY = pdf["PRIMARY"]
    ACCENT = pdf["ACCENT"]
    GREEN = pdf["GREEN"]
    RED = pdf["RED"]
    try:
        register_fonts()
        settings = settings_service.get_settings()
        economy = bool(getattr(settings, "print_economy", False))
        primary = hex_color(settings.invoice_primary, PRIMARY)
        accent = hex_color(settings.invoice_accent, ACCENT)
        light_bg = light_tint(settings.invoice_primary)
        # خلفية صف العميل: أغمق قليلًا من تظليل الصفوف العادي ليبرز
        cust_bg = _c.HexColor("#DCE7F5")
        border_c = _c.HexColor("#9FB3C8")   # خطوط جدول أوضح
        # لون نص الأشرطة (الترويسة ورأس الجدول وصف الإجمالي):
        # أبيض فوق الخلفية الملونة · أسود في الوضع الاقتصادي لأن
        # الخلفية تصبح بيضاء (كان أسود فيخرج الهيدر أسود داكنًا).
        bar_text = _c.white
        if economy:
            primary = _c.white          # شريط أبيض بدل الأسود
            accent = _c.white
            light_bg = _c.white
            cust_bg = _c.white
            border_c = _c.black
            bar_text = _c.black         # نص أسود فوق الأبيض
            color_cols = {}  # بلا تلوين أخضر/أحمر · أسود فقط
        color_cols = color_cols or {}
        _num_cols = (set(number_cols) if number_cols is not None else None)
        # حماية: بعض التقارير تمرّر None صراحةً
        right_align_cols = tuple(right_align_cols or ())

        right = PAGE_W - MARGIN
        left = MARGIN
        avail = right - left
        scale = avail / (sum(col_widths) * mm)
        cols = [w * mm * scale for w in col_widths]
        total_w = sum(cols)

        # مراجع تُقرأ داخل الكلاس (يُعرَّف قبل تحميل الصور أدناه)
        _footer_img_ref = [None]
        _draw_footer_ref = [None]

        class _NumberedCanvas(rl_canvas.Canvas):
            def __init__(self, *args, **kwargs):
                super().__init__(*args, **kwargs)
                self._saved_page_states = []
            def showPage(self):
                self._saved_page_states.append(dict(self.__dict__))
                self._startPage()
            def save(self):
                total = len(self._saved_page_states) + 1
                self._saved_page_states.append(dict(self.__dict__))
                for state in self._saved_page_states:
                    self.__dict__.update(state)
                    # صورة الفوتر المرفوعة على كل صفحة (توحيدًا مع
                    # الفواتير والسندات وكشوف الحساب)
                    _fh = 0.0
                    if _footer_img_ref[0]:
                        try:
                            _fh = _draw_footer_ref[0](self, PAGE_W,
                                                      _footer_img_ref[0]) or 0.0
                        except Exception:
                            pass
                    self.setFillColor(_c.black)
                    self.setFont(AR_FONT, 8)
                    # V4.33: الرقم فوق شريط الفوتر لا عليه (8مم فقط بلا صورة).
                    _num_y = (_fh + 3.5 * mm) if _fh else 8 * mm
                    self.drawCentredString(PAGE_W / 2, _num_y,
                                           ar(f"صفحة رقم {self._pageNumber} من {total}"))
                    super().showPage()
                super().save()

        c = _NumberedCanvas(out_path, pagesize=(PAGE_W, PAGE_H))
        apply_pdf_metadata(c, title)

        _safe_image = pdf["safe_image"]
        _draw_header_image = pdf["draw_header_image"]
        _draw_footer_image = pdf["draw_footer_image"]
        header_img = _safe_image(getattr(settings, "header_image", "") or "")
        footer_img = _safe_image(getattr(settings, "footer_image", "") or "")
        _footer_img_ref[0] = footer_img
        _draw_footer_ref[0] = _draw_footer_image
        # V4.33: حجز شريط الفوتر + رقم الصفحة — المحتوى يتوقف فوقهما.
        _footer_h = 0.0
        if footer_img:
            try:
                _footer_h = pdf["footer_image_height"](PAGE_W, footer_img) or 0.0
            except Exception:
                _footer_h = 0.0
        _band = _footer_h + 8 * mm if _footer_h else 0.0

        def header():
            """ترويسة موحّدة مع بقية مطبوعات النظام.

            تُستخدم صورة الترويسة المرفوعة في الإعدادات إن وُجدت (كما في
            الفواتير والسندات)، وإلا يُرسم الشريط النصي القديم فلا تظهر
            الصفحة بترويسة فارغة.
            """
            if header_img:
                h = _draw_header_image(c, PAGE_W, PAGE_H, header_img)
                y0 = PAGE_H - h - 4 * mm
                c.setFillColor(_c.black)
                c.setFont(AR_FONT_BOLD, 12)
                c.drawRightString(right, y0, ar(title))
                c.setFont(AR_FONT, 9)
                c.drawString(left, y0,
                             ar("تاريخ الإصدار: "
                                + datetime.now().strftime('%Y-%m-%d %H:%M')))
                return y0 - 8 * mm
            c.setFillColor(primary)
            c.rect(0, PAGE_H - 26 * mm, PAGE_W, 26 * mm, stroke=0, fill=1)
            if economy:     # إطار سفلي يفصل الترويسة البيضاء عن الجدول
                c.setStrokeColor(_c.black)
                c.setLineWidth(1)
                c.line(0, PAGE_H - 26 * mm, PAGE_W, PAGE_H - 26 * mm)
            c.setFillColor(bar_text)
            c.setFont(AR_FONT_BOLD, 14)
            c.drawRightString(right, PAGE_H - 11 * mm, ar(settings.company_name))
            c.setFont(AR_FONT_BOLD, 11)
            c.drawRightString(right, PAGE_H - 20 * mm, ar(title))
            c.setFont(AR_FONT, 9)
            c.drawString(left, PAGE_H - 20 * mm,
                         ar(f"تاريخ الإصدار: "
                            f"{datetime.now().strftime('%Y-%m-%d %H:%M')}"))
            return PAGE_H - 34 * mm

        def _fit_font(text, font, size, width, floor=6.5):
            """أصغر حجم خط يجعل النص يدخل في عرض العمود (بحد أدنى)."""
            while size > floor and c.stringWidth(text, font, size) > width:
                size -= 0.25
            return size

        def _wrap2(text, font, size, width, floor):
            """V3.17: لفّ النص على سطرين ضمن العرض — (س1، س2، الحجم).

            يجرّب الحجم الحالي ثم أصغر تدريجًا حتى floor. يعيد None
            إن تعذّر (كلمة أطول من العمود أو يلزم سطر ثالث) فيُستخدم
            مسار التصغير/القطع القديم.
            """
            words = [wd for wd in str(text).split(" ") if wd]
            if len(words) < 2:
                return None
            sz = size
            while sz >= floor - 0.01:
                l1, l2, fail = [], [], False
                for wd in words:
                    c1 = l1 + [wd]
                    if c.stringWidth(ar(" ".join(c1)), font, sz) <= width:
                        l1 = c1
                    else:
                        c2 = l2 + [wd]
                        if c.stringWidth(ar(" ".join(c2)),
                                         font, sz) <= width:
                            l2 = c2
                        else:
                            fail = True
                            break
                if not fail and l2:
                    while len(l1) > len(l2) + 1:
                        nl2 = [l1[-1]] + l2
                        if c.stringWidth(ar(" ".join(nl2)),
                                         font, sz) > width:
                            break
                        l2, l1 = nl2, l1[:-1]
                    return (" ".join(l1), " ".join(l2), sz)
                if not fail:
                    return None
                sz -= 1.0
            return None

        def _split_head(head, w):
            """تقسيم عنوان طويل إلى سطرين عند أنسب مسافة."""
            words = str(head).split()
            if len(words) < 2:
                return [head]
            best, diff = None, None
            for i in range(1, len(words)):
                a = " ".join(words[:i])
                b = " ".join(words[i:])
                d = abs(c.stringWidth(ar(a), AR_FONT_BOLD, F_HEAD)
                        - c.stringWidth(ar(b), AR_FONT_BOLD, F_HEAD))
                if diff is None or d < diff:
                    best, diff = (a, b), d
            return list(best)

        HEAD_H = ROW_H * 1.35        # رأس الجدول أعلى ليتسع لسطرين

        def table_header(y):
            c.setFillColor(primary)
            c.rect(right - total_w, y - HEAD_H, total_w, HEAD_H,
                   stroke=0, fill=1)
            if economy:     # حدود سوداء لرأس الجدول الأبيض
                c.setStrokeColor(_c.black)
                c.rect(right - total_w, y - HEAD_H, total_w, HEAD_H,
                       stroke=1, fill=0)
            c.setFillColor(bar_text)
            x = right
            for w, head in zip(cols, heads):
                avail = w - 2.5 * mm
                atext = ar(str(head))
                if c.stringWidth(atext, AR_FONT_BOLD, F_HEAD) <= avail:
                    c.setFont(AR_FONT_BOLD, F_HEAD)
                    c.drawCentredString(x - w / 2,
                                        y - HEAD_H / 2 - 1.2 * mm, atext)
                else:
                    parts = _split_head(head, w)
                    if len(parts) == 2:
                        fs = min(_fit_font(ar(parts[0]), AR_FONT_BOLD,
                                           F_HEAD, avail),
                                 _fit_font(ar(parts[1]), AR_FONT_BOLD,
                                           F_HEAD, avail))
                        c.setFont(AR_FONT_BOLD, fs)
                        c.drawCentredString(x - w / 2,
                                            y - HEAD_H / 2 + 0.7 * mm,
                                            ar(parts[0]))
                        c.drawCentredString(x - w / 2,
                                            y - HEAD_H / 2 - 3.1 * mm,
                                            ar(parts[1]))
                    else:
                        fs = _fit_font(atext, AR_FONT_BOLD, F_HEAD, avail)
                        c.setFont(AR_FONT_BOLD, fs)
                        c.drawCentredString(x - w / 2,
                                            y - HEAD_H / 2 - 1.2 * mm, atext)
                x -= w
            return y - HEAD_H

        # V3.17: ملاءمة تلقائية — قياس أعرض محتوى لكل عمود ثم سدّ
        # العجز من فائض الأعمدة الأخرى (المجموع ثابت = عرض الصفحة).
        _NJ = len(cols)
        _needs = [0.0] * _NJ
        # V3.18: العمود المرن فيه نص يُلفّ (مسافة) فيتحمّل التقليص،
        # أما الجامد (تاريخ/هاتف/رقم — بلا مسافات) فلا بديل له عن العرض.
        _flex = [False] * _NJ
        _numc = [False] * _NJ
        for _j, _head in enumerate(heads):
            if _j >= _NJ:
                break
            _parts = _split_head(_head, cols[_j])
            _hw = max([c.stringWidth(ar(str(_p)), AR_FONT_BOLD, F_HEAD)
                       for _p in _parts] or [0.0])
            _needs[_j] = max(_needs[_j], _hw)
        for _i, _row in enumerate(rows):
            try:
                _kind = row_kinds[_i][0] if row_kinds and _i < len(row_kinds) else ""
            except Exception:
                _kind = ""
            _acc = (len(_row) > 1 and str(_row[1]) in (accent_rows or {}))
            for _j, _val in enumerate(_row):
                if _j >= _NJ:
                    break
                _isn = isinstance(_val, (int, float)) and (
                    _num_cols is None or _j in _num_cols)
                _tx = (f"{_val:,.3f}" if _isn
                       else str(_val if _val not in (None, "") else "-"))
                if _acc:
                    _f, _s = AR_FONT_BOLD, F_CELL + 1.0
                elif _isn or _kind == "customer":
                    _f, _s = AR_FONT_BOLD, F_CELL
                else:
                    _f, _s = AR_FONT, F_CELL
                _needs[_j] = max(
                    _needs[_j], c.stringWidth(ar(_tx), _f, _s))
                if any(_ch.isspace() for _ch in _tx):
                    _flex[_j] = True
                if _isn:
                    _numc[_j] = True
        if totals_row:
            for _j, _val in enumerate(totals_row):
                if _j >= _NJ or _val == "":
                    continue
                _isn = isinstance(_val, (int, float)) and (
                    _num_cols is None or _j in _num_cols)
                _tx = f"{_val:,.3f}" if _isn else str(_val)
                _needs[_j] = max(
                    _needs[_j], c.stringWidth(ar(_tx), AR_FONT_BOLD, F_TOTAL))
        _needs = [_n + 3 * mm for _n in _needs]
        # هامش أمان: العمود المانح يحتفظ باحتياجه + 1.5مم فلا يلتف
        # نصٌ كان يدخل سطرًا واحدًا قبل إعادة التوزيع.
        _SLACK = 1.5 * mm
        # V3.18: هدف مزدوج — الجامد (تاريخ/هاتف/رقم) يحتاج عرضه
        # كاملًا، والمرن يكتفي بالنصف (يلتف على سطرين) عند الضرورة.
        _half = [(_n / 2.0 + 2 * mm) if _flex[_j] else _n
                 for _j, _n in enumerate(_needs)]
        _cols = list(cols)

        def _redist(_rigid_only, _deep):
            # جولة توزيع: نسدّ عجز الجامد فقط أو المرن فقط.
            # _deep: المرن يمنح حتى هدف النصف لا الكامل.
            for _pass in range(_NJ * 2):
                # V3.18: داخل الجامد — غير الرقمي (تاريخ/هاتف) أولًا
                # لأن أرضية تصغيره أعلى (6.5 مقابل 5.5 للرقمي).
                _def = sorted(
                    ((_needs[_j] - _cols[_j], _j) for _j in range(_NJ)
                     if _needs[_j] > _cols[_j] + 0.5 * mm
                     and (not _flex[_j]) == _rigid_only),
                    key=lambda _t: ((1 if _numc[_t[1]] else 0)
                                    if _rigid_only else 0, -_t[0]))
                if not _def:
                    break
                _moved = False
                for _d, _j in _def:
                    for _s, _k in sorted(
                            ((_cols[_x] - (_half[_x]
                                           if _deep and _flex[_x]
                                           else _needs[_x] + _SLACK), _x)
                             for _x in range(_NJ) if _x != _j),
                            key=lambda _t: (1 if _flex[_t[1]] else 0,
                                            -_t[0])):
                        if _s <= 0.5 * mm:
                            continue
                        _give = min(_s, _needs[_j] - _cols[_j])
                        if _give <= 0.01 * mm:
                            continue
                        _cols[_k] -= _give
                        _cols[_j] += _give
                        _moved = True
                        if _cols[_j] >= _needs[_j] - 0.01 * mm:
                            break
                if not _moved:
                    break

        _redist(True, False)    # 1) الجامد من الفائض العادي
        _redist(True, True)     # 2) الجامد الباقي من أنصاف المرن
        _redist(False, True)    # 3) المرن بما تبقّى
        cols = _cols
        total_w = sum(cols)

        y = header()
        y = table_header(y)
        _accents = accent_rows or {}
        for i, row in enumerate(rows):
            if y - ROW_H < max(26 * mm, _band):
                c.showPage()
                y = header()
                y = table_header(y)
            kind = ""
            if row_kinds and i < len(row_kinds):
                try:
                    kind = row_kinds[i][0]
                except Exception:
                    kind = ""
            is_cust = (kind == "customer")
            acc_fill = None
            if len(row) > 1 and str(row[1]) in _accents:
                try:
                    acc_fill = _c.HexColor(_accents[str(row[1])])
                except Exception:
                    acc_fill = None
            # V3.19: استثناء التقرير المالي — يحتفظ بالتلوين في الاقتصادي.
            if acc_fill is not None and (not economy
                                         or keep_accents_in_economy):
                # V3.16: صف مميز — خلفية ملوّنة (إجمالي/رصيد/صافي)
                c.setFillColor(acc_fill)
                c.rect(right - total_w, y - ROW_H, total_w, ROW_H,
                       stroke=0, fill=1)
            elif is_cust and not economy:
                # صف العميل بارز: خلفية أوضح تفصل كل مجموعة عن التي قبلها
                c.setFillColor(cust_bg)
                c.rect(right - total_w, y - ROW_H, total_w, ROW_H,
                       stroke=0, fill=1)
            elif i % 2 == 0:
                c.setFillColor(light_bg)
                c.rect(right - total_w, y - ROW_H, total_w, ROW_H,
                       stroke=0, fill=1)
            c.setStrokeColor(border_c)
            c.setLineWidth(1.1 if is_cust else 0.5)
            c.rect(right - total_w, y - ROW_H, total_w, ROW_H, stroke=1, fill=0)
            c.setLineWidth(1)
            x = right
            for j, (w, val) in enumerate(zip(cols, row)):
                # الأرقام تُنسّق بثلاث خانات عشرية فقط إن كان العمود
                # ماليًا (number_cols). بدون ذلك كان عمود «م» التسلسلي
                # يظهر «1.000» فيبدو مبلغًا لا ترقيمًا.
                is_num = isinstance(val, (int, float)) and (
                    _num_cols is None or j in _num_cols)
                text = f"{val:,.3f}" if is_num else str(
                    val if val not in (None, "") else "-")
                mode = color_cols.get(j)
                if mode == "green" and is_num and val:
                    c.setFillColor(GREEN)
                elif mode == "red" and is_num and val:
                    c.setFillColor(RED)
                elif mode == "auto" and is_num:
                    c.setFillColor(RED if val > 0.001 else GREEN)
                elif mode == "profit" and is_num:
                    c.setFillColor(GREEN if val >= -0.001 else RED)
                else:
                    c.setFillColor(_c.black)
                # قاعدة البولد الموحّدة في كل التقارير:
                #   · الأعمدة الرقمية (المالية) عريضة لتبرز
                #   · صف العميل/المجموعة عريض كاملًا
                #   · الصف المميز عريض كاملًا بخط أكبر
                #   · ما عدا ذلك عادي — بلا اعتماد على التلوين
                bold = is_num or is_cust or acc_fill is not None
                font = AR_FONT_BOLD if bold else AR_FONT
                fsize = F_CELL + 1.0 if acc_fill is not None else F_CELL
                c.setFont(font, fsize)
                atext = ar(text)
                avail = w - 3 * mm
                if c.stringWidth(atext, font, fsize) > avail:
                    if is_num:
                        # الأرقام لا تُقص أبدًا — تُصغَّر حتى تدخل كاملة
                        fsize = _fit_font(atext, font, fsize, avail, 5.5)
                        c.setFont(font, fsize)
                    else:
                        # V3.17: النص الطويل: تصغير لسطر واحد أولًا،
                        # فإن تعذّر يُلفّ على سطرين داخل نفس الصف،
                        # والقطع عند حد الكلمة هو الملاذ الأخير فقط.
                        _f0 = fsize
                        fsize = _fit_font(atext, font, fsize, avail,
                                          max(6.5, F_CELL - 2.0))
                        c.setFont(font, fsize)
                        if c.stringWidth(atext, font, fsize) > avail:
                            _wr = _wrap2(text, font, _f0, avail, 6.5)
                            if _wr is not None:
                                _ln1, _ln2, _fs2 = _wr
                                _top = y - 3.0 * mm
                                _bot = y - ROW_H + 2.4 * mm
                                _fs2 = min(
                                    _fs2, ((_top - _bot) / mm) / 0.42)
                                c.setFont(font, _fs2)
                                if j in right_align_cols:
                                    c.drawRightString(x - 2.5 * mm, _top,
                                                      ar(_ln1))
                                    c.drawRightString(x - 2.5 * mm, _bot,
                                                      ar(_ln2))
                                else:
                                    c.drawCentredString(x - w / 2, _top,
                                                        ar(_ln1))
                                    c.drawCentredString(x - w / 2, _bot,
                                                        ar(_ln2))
                                x -= w
                                continue
                            # الملاذ الأخير: تصغير للحد الأدنى ثم قطع
                            # عند حد الكلمة بعلامة حذف صريحة.
                            fsize = _fit_font(atext, font, fsize, avail,
                                              6.5)
                            c.setFont(font, fsize)
                        c.setFont(font, fsize)
                        if c.stringWidth(atext, font, fsize) > avail:
                            while (c.stringWidth(atext + "…", font, fsize)
                                   > avail and len(atext) > 2):
                                atext = atext[:-1]
                            # V3.17: القطع عند حد الكلمة بدل منتصفها.
                            if " " in atext:
                                _wcut = atext.rsplit(" ", 1)[0]
                                if len(_wcut) >= 2:
                                    atext = _wcut
                            atext += "…"
                if j in right_align_cols:
                    c.drawRightString(x - 2.5 * mm,
                                      y - ROW_H + 3.4 * mm, atext)
                else:
                    c.drawCentredString(x - w / 2,
                                        y - ROW_H + 3.4 * mm, atext)
                x -= w
            y -= ROW_H

        if totals_row:
            if y - ROW_H < max(28 * mm, _band):
                c.showPage()
                y = header()
            c.setFillColor(accent)
            c.rect(right - total_w, y - ROW_H, total_w, ROW_H, stroke=0, fill=1)
            if economy:     # صف الإجمالي الأبيض يحتاج إطارًا ليُرى
                c.setStrokeColor(_c.black)
                c.setLineWidth(1.2)
                c.rect(right - total_w, y - ROW_H, total_w, ROW_H,
                       stroke=1, fill=0)
                c.setLineWidth(1)
            c.setFillColor(bar_text)
            c.setFont(AR_FONT_BOLD, F_TOTAL)
            x = right
            for j, (w, val) in enumerate(zip(cols, totals_row)):
                if val != "":
                    _isn = isinstance(val, (int, float)) and (
                        _num_cols is None or j in _num_cols)
                    text = f"{val:,.3f}" if _isn else str(val)
                    atext = ar(text)
                    fs = _fit_font(atext, AR_FONT_BOLD, F_TOTAL,
                                   w - 2.5 * mm, 5.5)
                    c.setFont(AR_FONT_BOLD, fs)
                    c.drawCentredString(x - w / 2,
                                        y - ROW_H + 3.4 * mm, atext)
                x -= w
            y -= ROW_H

        if footer_note:
            # الملاحظة تُقسَّم على أسطر تدخل في عرض الصفحة (لا تخرج عنها)
            y -= 7 * mm
            c.setFillColor(_c.black)
            fs = 10.5 if not str(orientation).lower().startswith("port") else 9
            c.setFont(AR_FONT_BOLD, fs)
            avail = right - left
            parts = [x.strip() for x in str(footer_note).split("|")
                     if x.strip()]
            lines, cur = [], ""
            for part in parts:
                cand = (cur + "   |   " + part) if cur else part
                if c.stringWidth(ar(cand), AR_FONT_BOLD, fs) <= avail:
                    cur = cand
                else:
                    if cur:
                        lines.append(cur)
                    cur = part
            if cur:
                lines.append(cur)
            for ln in lines:
                if y < max(16 * mm, _band):
                    c.showPage()
                    y = header()
                    c.setFillColor(_c.black)
                    c.setFont(AR_FONT_BOLD, fs)
                c.drawRightString(right, y, ar(ln))
                y -= 6 * mm

        c.save()
        return out_path
    except PDFError:
        raise
    except Exception as exc:
        raise PDFError(f"تعذر إنشاء التقرير: {exc}") from exc


def generate_table_excel(title: str, heads: list[str], rows: list[list],
                         out_path: str, totals_row: list | None = None,
                         number_cols: list[int] | None = None,
                         accent_rows: dict | None = None) -> str:
    """توليد نفس التقرير كملف Excel أنيق RTL."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    settings = settings_service.get_settings()
    number_cols = number_cols or []

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    title_font = Font(name="Calibri", bold=True, size=16, color="1E3A5F")
    sub_font = Font(name="Calibri", size=11, color="7B8A9A")
    total_fill = PatternFill("solid", fgColor="DDEBF7")
    total_font = Font(name="Calibri", bold=True, size=12, color="1E3A5F")
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    wb = Workbook()
    try:
        _company = (getattr(settings_service.get_settings(),
                            'company_name', '') or '').strip()
        wb.properties.creator = _company or settings_service.app_name()
        wb.properties.lastModifiedBy = wb.properties.creator
    except Exception:
        pass
    ws = wb.active
    ws.title = title[:28] or "تقرير"
    ws.sheet_view.rightToLeft = True

    ncols = len(heads)
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = settings.company_name or title
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = (f"{title} · تاريخ الإصدار: "
                f"{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    ws["A2"].font = sub_font
    ws["A2"].alignment = center

    hr = 4
    for col, head in enumerate(heads, start=1):
        cell = ws.cell(row=hr, column=col, value=head)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[hr].height = 24

    r = hr + 1
    _accents = accent_rows or {}
    for row in rows:
        _acc_fill = None
        if len(row) > 1 and str(row[1]) in _accents:
            try:
                _acc_fill = PatternFill(
                    "solid", fgColor=str(_accents[str(row[1])]).lstrip("#"))
            except Exception:
                _acc_fill = None
        for col, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.border = border
            cell.alignment = center
            if (col - 1) in number_cols:
                cell.number_format = "#,##0.000"
            if _acc_fill is not None:
                # V3.16: صف مميز — تعبئة ملوّنة وخط عريض
                cell.fill = _acc_fill
                cell.font = Font(name="Calibri", bold=True, size=12)
        r += 1

    if totals_row:
        for col, val in enumerate(totals_row, start=1):
            cell = ws.cell(row=r, column=col,
                           value=val if val != "" else None)
            cell.fill = total_fill
            cell.font = total_font
            cell.border = border
            cell.alignment = center
            if (col - 1) in number_cols:
                cell.number_format = "#,##0.000"

    for col in range(1, ncols + 1):
        width = max(14, len(str(heads[col - 1])) + 6)
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(out_path)
    return out_path


def generate_report_workbook(title: str, sheets: list[dict],
                             out_path: str) -> str:
    """تقرير Excel متعدد الأوراق بنفس هوية generate_table_excel.

    sheets = [{name, heads, rows, number_cols, bold_rows, widths,
    subtitle}]: bold_rows فهارس صفوف البيانات العريضة (0-based) ·
    widths عروض أعمدة اختيارية {index: width}.
    """
    from openpyxl import Workbook
    from openpyxl.styles import (Alignment, Border, Font, PatternFill,
                                 Side)
    from openpyxl.utils import get_column_letter

    settings = settings_service.get_settings()

    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    title_font = Font(name="Calibri", bold=True, size=16, color="1E3A5F")
    sub_font = Font(name="Calibri", size=11, color="7B8A9A")
    total_fill = PatternFill("solid", fgColor="DDEBF7")
    total_font = Font(name="Calibri", bold=True, size=12, color="1E3A5F")
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center",
                       wrap_text=True)

    wb = Workbook()
    try:
        _company = (getattr(settings_service.get_settings(),
                            "company_name", "") or "").strip()
        wb.properties.creator = _company or settings_service.app_name()
        wb.properties.lastModifiedBy = wb.properties.creator
    except Exception:
        pass
    first = True
    for sh in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = (sh["name"] or "ورقة")[:31]
        ws.sheet_view.rightToLeft = True
        heads, rows = sh["heads"], sh["rows"]
        number_cols = sh.get("number_cols") or []
        bold_rows = set(sh.get("bold_rows") or [])
        ncols = len(heads)
        last = get_column_letter(ncols)
        ws.merge_cells("A1:" + last + "1")
        ws["A1"] = settings.company_name or title
        ws["A1"].font = title_font
        ws["A1"].alignment = center
        ws.merge_cells("A2:" + last + "2")
        ws["A2"] = (sh.get("subtitle") or title)
        ws["A2"].value = str(ws["A2"].value) + " · تاريخ الإصدار: " \
            + datetime.now().strftime("%Y-%m-%d %H:%M")
        ws["A2"].font = sub_font
        ws["A2"].alignment = center
        hr = 4
        for col, head in enumerate(heads, start=1):
            cell = ws.cell(row=hr, column=col, value=head)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border
        ws.row_dimensions[hr].height = 24
        r = hr + 1
        for idx, row in enumerate(rows):
            for col, val in enumerate(row, start=1):
                cell = ws.cell(row=r, column=col,
                               value=val if val != "" else None)
                cell.border = border
                cell.alignment = center
                if (col - 1) in number_cols:
                    cell.number_format = "#,##0.000"
                if idx in bold_rows:
                    cell.fill = total_fill
                    cell.font = total_font
            r += 1
        _widths = sh.get("widths") or {}
        for col in range(1, ncols + 1):
            width = _widths.get(col - 1,
                                max(14, len(str(heads[col - 1])) + 6))
            ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(out_path)
    return out_path
