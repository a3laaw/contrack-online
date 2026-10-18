# -*- coding: utf-8 -*-
"""نماذج Excel للعهد والتصفيات."""


def generate_blank_custody_settlement_excel(out_path: str) -> str:
    """ينشئ نموذج Excel فارغًا يُسلّم للموظف لتسجيل الوارد والمنصرف.

    النموذج لا يعتمد على قاعدة البيانات: الموظف يكتب الحركات، والملف يحسب
    الإجماليات والرصيد تلقائيًا بصيغ Excel.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            f"مكتبة {exc.name or 'openpyxl'} غير مثبتة. شغّل INSTALL.bat أو ثبّت المتطلبات."
        ) from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "نموذج تصفية العهدة"
    ws.sheet_view.rightToLeft = True

    primary = "1E3A5F"
    light = "EAF2F8"
    green = "DCFCE7"
    red = "FEE2E2"
    border_color = "94A3B8"
    thin = Side(style="thin", color=border_color)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:H1")
    ws["A1"] = "نموذج تصفية عهدة - وارد / منصرف"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=primary)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    info = [
        ("A2", "اسم الموظف / المستلم"), ("B2", ""),
        ("C2", "حساب العهدة"), ("D2", ""),
        ("E2", "تاريخ التصفية"), ("F2", ""),
        ("G2", "رصيد افتتاحي"), ("H2", 0),
    ]
    for cell, value in info:
        ws[cell] = value
        ws[cell].border = border
        ws[cell].alignment = Alignment(horizontal="center" if cell[0] in "ACEG" else "right", vertical="center")
        if cell[0] in "ACEG":
            ws[cell].font = Font(bold=True)
            ws[cell].fill = PatternFill("solid", fgColor=light)
    ws["H2"].number_format = '#,##0.000'

    ws.merge_cells("A3:B3"); ws["A3"] = "إجمالي الوارد"; ws["C3"] = "=SUM(D6:D105)+H2"
    ws.merge_cells("D3:E3"); ws["D3"] = "إجمالي المنصرف"; ws["F3"] = "=SUM(E6:E105)"
    ws["G3"] = "الرصيد النهائي"; ws["H3"] = "=C3-F3"
    for cell in ("A3", "D3", "G3"):
        ws[cell].font = Font(bold=True)
        ws[cell].fill = PatternFill("solid", fgColor=light)
        ws[cell].alignment = Alignment(horizontal="center")
    for cell in ("C3", "F3", "H3"):
        ws[cell].font = Font(bold=True)
        ws[cell].number_format = '#,##0.000'
        ws[cell].alignment = Alignment(horizontal="center")
        ws[cell].border = border
    ws["C3"].fill = PatternFill("solid", fgColor=green)
    ws["F3"].fill = PatternFill("solid", fgColor=red)
    ws["H3"].fill = PatternFill("solid", fgColor="DBEAFE")

    ws.merge_cells("A4:H4")
    ws["A4"] = "يكتب الموظف الوارد في عمود وارد، والمصاريف في عمود منصرف. الرصيد يحسب تلقائيًا: الرصيد السابق + الوارد - المنصرف."
    ws["A4"].font = Font(bold=True, color="475569")
    ws["A4"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws["A4"].fill = PatternFill("solid", fgColor="F8FAFC")

    headers = ["م", "التاريخ", "البيان", "وارد", "منصرف", "الرصيد", "العميل / المورد", "ملاحظات / رقم مرفق"]
    start = 5
    for col, head in enumerate(headers, 1):
        c = ws.cell(start, col, head)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=primary)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    for r in range(6, 106):
        ws.cell(r, 1, r - 5)
        # الرصيد التراكمي: رصيد افتتاحي + الوارد حتى هذا السطر - المنصرف حتى هذا السطر.
        ws.cell(r, 6, f"=$H$2+SUM($D$6:D{r})-SUM($E$6:E{r})")
        for col in range(1, 9):
            c = ws.cell(r, col)
            c.border = border
            c.alignment = Alignment(horizontal="right" if col in (3, 7, 8) else "center", vertical="center", wrap_text=True)
            if col in (4, 5, 6):
                c.number_format = '#,##0.000'
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="FFFFFF")
            else:
                c.fill = PatternFill("solid", fgColor="F8FAFC")

    widths = {"A": 8, "B": 14, "C": 46, "D": 14, "E": 14, "F": 14, "G": 24, "H": 28}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for r in range(6, 106):
        ws.row_dimensions[r].height = 28
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = "A5:H105"

    # صفحة مختصرة للطباعة.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    for col in range(1, 9):
        ws.cell(2, col).border = border
        ws.cell(3, col).border = border
        ws.cell(4, col).border = border

    wb.save(out_path)
    return out_path
